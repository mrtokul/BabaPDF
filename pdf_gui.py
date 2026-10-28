import os
import sys
import ctypes
import json
import hashlib
import html
import re
import math
import uuid
import shutil
import secrets
import subprocess
import tempfile
import time
import threading
import traceback
from datetime import datetime
import tkinter as tk
from tkinter import filedialog, messagebox, ttk, simpledialog

try:
    import pymupdf as fitz  # PyMuPDF (yeni ad)
except ImportError:
    import fitz  # PyMuPDF (eski ad)
import pdfplumber

try:
    from tkinterdnd2 import DND_FILES, TkinterDnD
    DND_AVAILABLE = True
except ImportError:
    DND_AVAILABLE = False

APP_NAME = "Baba PDF"
APP_VERSION = "1.12.0"
APP_DATE = "28.10.2026"
UNDO_MAX_STEPS = 15
UNDO_MAX_BYTES = 400 * 1024 * 1024
ZOOM_MIN, ZOOM_MAX = 0.3, 5.0

# Pastel renk paletleri (açık ve koyu tema)
THEMES = {
    "light": dict(
        bg="#F8F5FC", panel="#EEE8F8", text="#3B3556", muted="#7A7396", entry="#FFFFFF",
        accent="#B9AAF0", accent_hover="#A795E8", accent_dark="#7C6FD6",
        lavender="#E4DDFA", mint="#D3F0E2", peach="#FFE0CF", sky="#D6EAFA", rose="#FADCE4", lemon="#FFF3C9",
        canvas="#E7E9F4", shadow="#C9CCE0", thumb_bg="#DCDFEE", card_border="#D8D0EE",
    ),
    "dark": dict(
        bg="#25232E", panel="#2F2C3B", text="#EDE9F7", muted="#A9A3C0", entry="#3A3748",
        accent="#6F63A8", accent_hover="#8174C2", accent_dark="#B9AAF0",
        lavender="#3B3560", mint="#2F4A3F", peach="#54392F", sky="#2E4457", rose="#573843", lemon="#5A5230",
        canvas="#2B2A33", shadow="#1B1A20", thumb_bg="#201F27", card_border="#463F66",
    ),
}
FONT_UI = "Segoe UI"


def sweep_old_temp_files(max_age_seconds=600):
    """Önceki oturumlardan kalan geçici çalışma kopyalarını (baba_pdf_<hex>.pdf) siler."""
    pattern = re.compile(r"baba_pdf_[0-9a-f]{32}\.pdf$")
    folder = tempfile.gettempdir()
    try:
        names = os.listdir(folder)
    except OSError:
        return
    now = time.time()
    for name in names:
        if pattern.match(name):
            path = os.path.join(folder, name)
            try:
                if now - os.path.getmtime(path) > max_age_seconds:
                    os.remove(path)
            except OSError:
                pass


def reveal_in_folder(path):
    """Dosyayı Windows Gezgini'nde gösterir (diğer sistemlerde klasörü açar)."""
    try:
        if sys.platform == "win32":
            subprocess.Popen(["explorer", "/select,", os.path.normpath(path)])
        elif sys.platform == "darwin":
            subprocess.Popen(["open", "-R", path])
        else:
            subprocess.Popen(["xdg-open", os.path.dirname(path) or "."])
    except OSError:
        pass


def update_manifest_url():
    """Güncelleme bildirimi adresi: BABA_PDF_UPDATE_URL ortam değişkeni veya AppData altındaki BabaPDF klasöründeki guncelleme_adresi.txt. Varsayılan yok."""
    url = os.environ.get("BABA_PDF_UPDATE_URL", "").strip()
    if not url:
        try:
            path = os.path.join(os.environ.get("APPDATA") or tempfile.gettempdir(), "BabaPDF", "guncelleme_adresi.txt")
            with open(path, encoding="utf-8") as handle:
                url = handle.read().strip()
        except OSError:
            url = ""
    return url


def parse_version(text):
    parts = []
    for piece in str(text).strip().lstrip("vV").split("."):
        digits = "".join(ch for ch in piece if ch.isdigit())
        parts.append(int(digits) if digits else 0)
    return tuple(parts)


def check_for_update(url, timeout=8):
    """https adresindeki {"version": "...", "url": "...", "notes": "..."} bildirimini okur; yalnızca GET, veri göndermez.
    (yeni_surum_var_mi, bildirim) döndürür."""
    import urllib.request
    if not url.lower().startswith("https://"):
        raise ValueError("Güncelleme adresi https:// ile başlamalıdır.")
    request = urllib.request.Request(url, headers={"User-Agent": f"{APP_NAME}/{APP_VERSION}"})
    with urllib.request.urlopen(request, timeout=timeout) as response:  # noqa: S310
        manifest = json.loads(response.read(65536).decode("utf-8"))
    return parse_version(manifest.get("version", "0")) > parse_version(APP_VERSION), manifest


def pdf_default_app_registered():
    """Baba PDF'in kurulumla Windows'a PDF işleyicisi olarak kaydedilip kaydedilmediğini döndürür."""
    try:
        import winreg
        with winreg.OpenKey(winreg.HKEY_CURRENT_USER, r"Software\RegisteredApplications") as key:
            winreg.QueryValueEx(key, "BabaPDF")
        return True
    except OSError:
        return False


def register_pdf_handler(exe_path):
    """HKCU altında Baba PDF'i .pdf için aday işleyici olarak kaydeder (yönetici hakkı gerekmez, sistem ayarı değişmez).
    Varsayılan seçimi Windows'ta kullanıcı yapar; kurulum betiği (Inno Setup) aynı anahtarları yazar."""
    import winreg
    prog_id = "BabaPDF.Document"
    base = r"Software\Classes"
    with winreg.CreateKey(winreg.HKEY_CURRENT_USER, rf"{base}\{prog_id}") as key:
        winreg.SetValueEx(key, "", 0, winreg.REG_SZ, "PDF Belgesi (Baba PDF)")
    with winreg.CreateKey(winreg.HKEY_CURRENT_USER, rf"{base}\{prog_id}\DefaultIcon") as key:
        winreg.SetValueEx(key, "", 0, winreg.REG_SZ, f'"{exe_path}",0')
    with winreg.CreateKey(winreg.HKEY_CURRENT_USER, rf"{base}\{prog_id}\shell\open\command") as key:
        winreg.SetValueEx(key, "", 0, winreg.REG_SZ, f'"{exe_path}" "%1"')
    with winreg.CreateKey(winreg.HKEY_CURRENT_USER, rf"{base}\.pdf\OpenWithProgids") as key:
        winreg.SetValueEx(key, prog_id, 0, winreg.REG_NONE, b"")
    epub_id = "BabaPDF.Epub"
    with winreg.CreateKey(winreg.HKEY_CURRENT_USER, rf"{base}\{epub_id}") as key:
        winreg.SetValueEx(key, "", 0, winreg.REG_SZ, "EPUB Kitabı (Baba PDF)")
    with winreg.CreateKey(winreg.HKEY_CURRENT_USER, rf"{base}\{epub_id}\DefaultIcon") as key:
        winreg.SetValueEx(key, "", 0, winreg.REG_SZ, f'"{exe_path}",0')
    with winreg.CreateKey(winreg.HKEY_CURRENT_USER, rf"{base}\{epub_id}\shell\open\command") as key:
        winreg.SetValueEx(key, "", 0, winreg.REG_SZ, f'"{exe_path}" "%1"')
    with winreg.CreateKey(winreg.HKEY_CURRENT_USER, rf"{base}\.epub\OpenWithProgids") as key:
        winreg.SetValueEx(key, epub_id, 0, winreg.REG_NONE, b"")
    book_id = "BabaPDF.Book"  # MOBI, AZW3, IBA, RTF: yalnızca "birlikte aç" adayı; varsayılanı kullanıcı seçer
    with winreg.CreateKey(winreg.HKEY_CURRENT_USER, rf"{base}\{book_id}") as key:
        winreg.SetValueEx(key, "", 0, winreg.REG_SZ, "E-Kitap / Belge (Baba PDF)")
    with winreg.CreateKey(winreg.HKEY_CURRENT_USER, rf"{base}\{book_id}\DefaultIcon") as key:
        winreg.SetValueEx(key, "", 0, winreg.REG_SZ, f'"{exe_path}",0')
    with winreg.CreateKey(winreg.HKEY_CURRENT_USER, rf"{base}\{book_id}\shell\open\command") as key:
        winreg.SetValueEx(key, "", 0, winreg.REG_SZ, f'"{exe_path}" "%1"')
    for ext in BOOK_EXTENSIONS:
        if ext != ".epub":
            with winreg.CreateKey(winreg.HKEY_CURRENT_USER, rf"{base}\{ext}\OpenWithProgids") as key:
                winreg.SetValueEx(key, book_id, 0, winreg.REG_NONE, b"")
    cap = r"Software\BabaPDF\Capabilities"
    with winreg.CreateKey(winreg.HKEY_CURRENT_USER, cap) as key:
        winreg.SetValueEx(key, "ApplicationName", 0, winreg.REG_SZ, APP_NAME)
        winreg.SetValueEx(key, "ApplicationDescription", 0, winreg.REG_SZ, "PDF ve e-kitap görüntüleme, PDF düzenleme")
    with winreg.CreateKey(winreg.HKEY_CURRENT_USER, cap + r"\FileAssociations") as key:
        winreg.SetValueEx(key, ".pdf", 0, winreg.REG_SZ, prog_id)
        winreg.SetValueEx(key, ".epub", 0, winreg.REG_SZ, epub_id)
        for ext in BOOK_EXTENSIONS:
            if ext != ".epub":
                winreg.SetValueEx(key, ext, 0, winreg.REG_SZ, book_id)
    with winreg.CreateKey(winreg.HKEY_CURRENT_USER, r"Software\RegisteredApplications") as key:
        winreg.SetValueEx(key, "BabaPDF", 0, winreg.REG_SZ, cap)


# --- Tek kopya koruması: açılış sürerken yinelenen başlatmaları (tek/çift/art arda tıklama) yok say ---
STARTUP_GUARD_DEFAULT = 20.0  # ölçüm yokken açılış süresi tahmini (sn)
LAST_STARTUP_SECONDS = None


def _process_start_epoch(pid):
    """Windows'ta sürecin oluşturulma zamanı (epoch sn); süreç yoksa veya okunamazsa None."""
    try:
        kernel = ctypes.windll.kernel32
        kernel.OpenProcess.restype = ctypes.c_void_p
        handle = kernel.OpenProcess(0x1000, False, int(pid))  # PROCESS_QUERY_LIMITED_INFORMATION
        if not handle:
            return None
        try:
            created, exited, kernel_t, user_t = (ctypes.c_ulonglong() for _ in range(4))
            if not kernel.GetProcessTimes(ctypes.c_void_p(handle), ctypes.byref(created), ctypes.byref(exited),
                                          ctypes.byref(kernel_t), ctypes.byref(user_t)):
                return None
            if exited.value:  # süreç bitmiş
                return None
            return (created.value - 116444736000000000) / 1e7
        finally:
            kernel.CloseHandle(ctypes.c_void_p(handle))
    except Exception:  # noqa: BLE001
        return None


def _process_image(pid):
    """Sürecin EXE yolu (küçük harf); okunamazsa None."""
    try:
        kernel = ctypes.windll.kernel32
        kernel.OpenProcess.restype = ctypes.c_void_p
        handle = kernel.OpenProcess(0x1000, False, int(pid))
        if not handle:
            return None
        try:
            buf = ctypes.create_unicode_buffer(1024)
            size = ctypes.c_ulong(len(buf))
            if kernel.QueryFullProcessImageNameW(ctypes.c_void_p(handle), 0, buf, ctypes.byref(size)):
                return buf.value.lower()
            return None
        finally:
            kernel.CloseHandle(ctypes.c_void_p(handle))
    except Exception:  # noqa: BLE001
        return None


def _same_image(pid_a, pid_b):
    """İki süreç aynı EXE dosyasından mı çalışıyor (tek dosya EXE'de önyükleyici = üst süreç)."""
    a, b = _process_image(pid_a), _process_image(pid_b)
    return bool(a) and a == b


def launch_click_epoch():
    """Kullanıcının programı başlattığı an: tek dosya EXE'de üst (önyükleyici) sürecin oluşturulma zamanı."""
    own = _process_start_epoch(os.getpid()) or time.time()
    if getattr(sys, "frozen", False):
        parent = _process_start_epoch(os.getppid())
        if parent and parent <= own and _same_image(os.getppid(), os.getpid()):
            return parent
    return own


def _startup_record_path():
    return os.path.join(os.path.dirname(error_log_path()), "baslatma.json")


class _StartupLock:
    """Kayıt okuma-yazmayı süreçler arasında sıraya sokan adlandırılmış mutex."""

    def __enter__(self):
        self.handle = None
        try:
            kernel = ctypes.windll.kernel32
            kernel.CreateMutexW.restype = ctypes.c_void_p
            self.handle = kernel.CreateMutexW(None, False, "Local\\BabaPDFBaslatmaKayitMutex")
            if self.handle:
                kernel.WaitForSingleObject(ctypes.c_void_p(self.handle), 5000)
        except Exception:  # noqa: BLE001
            self.handle = None
        return self

    def __exit__(self, *_):
        if self.handle:
            kernel = ctypes.windll.kernel32
            kernel.ReleaseMutex(ctypes.c_void_p(self.handle))
            kernel.CloseHandle(ctypes.c_void_p(self.handle))
        return False


def _read_startup_record():
    try:
        with open(_startup_record_path(), encoding="utf-8") as handle:
            data = json.load(handle)
        return data if isinstance(data, dict) else {}
    except (OSError, ValueError):
        return {}


def _write_startup_record(record):
    try:
        with open(_startup_record_path(), "w", encoding="utf-8") as handle:
            json.dump(record, handle)
    except OSError:
        pass


def startup_guard_seconds(durations):
    """Ölçülmüş açılış sürelerinin ortancası × 1,5 + 2 sn; ölçüm yoksa varsayılan."""
    values = sorted(float(v) for v in durations if isinstance(v, (int, float)) and v > 0)
    if not values:
        return STARTUP_GUARD_DEFAULT
    return values[len(values) // 2] * 1.5 + 2.0


def claim_startup(arg="", click=None, now=None, pid=None):
    """Bu başlatma yeni bir kopya açabilir mi? Başka bir kopyanın açılışı sürerken (tıklama anı onun
    [tıklama, hazır] aralığındaysa) False döner; açılış bittikten sonraki başlatmalara izin verilir."""
    click = launch_click_epoch() if click is None else click
    pid = os.getpid() if pid is None else pid
    with _StartupLock():
        record = _read_startup_record()
        durations = record.get("sureler", [])
        other = record.get("pid")
        if other and other != pid and record.get("click"):
            ready = record.get("ready")
            end = ready if ready else record["click"] + startup_guard_seconds(durations)
            alive = bool(ready) or _process_start_epoch(other) is not None
            same_file = (not arg) or arg == record.get("arg", "")
            if alive and same_file and record["click"] - 2.0 <= click <= end:
                return False
        _write_startup_record({"click": click, "ready": None, "pid": pid, "arg": arg, "sureler": durations})
    return True


def mark_startup_ready(now=None):
    """Pencere göründü: açılış süresini (tıklamadan hazıra) kaydeder ve sonraki başlatmalara izin verir."""
    global LAST_STARTUP_SECONDS
    now = time.time() if now is None else now
    with _StartupLock():
        record = _read_startup_record()
        if record.get("pid") != os.getpid():
            return None
        duration = max(0.0, now - float(record.get("click", now)))
        durations = (list(record.get("sureler", [])) + [round(duration, 2)])[-5:]
        record.update({"ready": now, "sureler": durations})
        _write_startup_record(record)
    LAST_STARTUP_SECONDS = duration
    return duration


def error_log_path():
    base = os.environ.get("LOCALAPPDATA") or os.environ.get("APPDATA") or tempfile.gettempdir()
    folder = os.path.join(base, "BabaPDF")
    os.makedirs(folder, exist_ok=True)
    return os.path.join(folder, "baba_pdf_error.log")


def log_exception(context, exc=None):
    try:
        with open(error_log_path(), "a", encoding="utf-8") as log:
            log.write(f"\n[{datetime.now().isoformat(timespec='seconds')}] {context}\n")
            if exc is not None:
                log.write(f"{type(exc).__name__}: {exc}\n")
            log.write(traceback.format_exc())
            log.write("\n")
    except Exception:
        pass


def show_runtime_error(parent, context, exc):
    log_exception(context, exc)
    messagebox.showerror(
        "Uygulama hatası",
        f"{context}:\n{exc}\n\nAyrıntılar kaydedildi:\n{error_log_path()}",
        parent=parent,
    )


def app_data_dir():
    base = os.environ.get("APPDATA") or os.path.expanduser("~")
    folder = os.path.join(base, "BabaPDF")
    os.makedirs(folder, exist_ok=True)
    return folder


def recent_files_path():
    return os.path.join(app_data_dir(), "recent_files.json")


def load_recent_files():
    try:
        with open(recent_files_path(), "r", encoding="utf-8") as source:
            items = json.load(source)
        return [path for path in items if isinstance(path, str) and os.path.exists(path)][:10]
    except (OSError, ValueError, TypeError):
        return []


def clear_recent_files():
    """Son açılanlar listesini boşaltır (dosyalara ve okuma konumlarına dokunmaz)."""
    try:
        with open(recent_files_path(), "w", encoding="utf-8") as target:
            json.dump([], target)
        return True
    except OSError:
        return False


def add_recent_file(path):
    items = [path] + [item for item in load_recent_files() if item != path]
    try:
        with open(recent_files_path(), "w", encoding="utf-8") as target:
            json.dump(items[:10], target, ensure_ascii=False, indent=2)
    except OSError:
        pass


def settings_path():
    return os.path.join(app_data_dir(), "settings.json")


def load_settings():
    try:
        with open(settings_path(), "r", encoding="utf-8") as source:
            data = json.load(source)
        return data if isinstance(data, dict) else {}
    except (OSError, ValueError):
        return {}


def save_setting(key, value):
    data = load_settings()
    data[key] = value
    try:
        with open(settings_path(), "w", encoding="utf-8") as target:
            json.dump(data, target, ensure_ascii=False, indent=2)
    except OSError:
        pass


def recovery_path(pdf_path):
    # hash() her çalıştırmada değişir; kurtarma dosyası yeniden başlatınca bulunabilsin diye kalıcı özet kullan.
    safe_name = hashlib.sha1(os.path.abspath(pdf_path).lower().encode("utf-8")).hexdigest()[:16]
    return os.path.join(app_data_dir(), f"recovery_{safe_name}.pdf")


def find_recovery_for(path):
    candidate = recovery_path(path)
    if not os.path.exists(candidate):
        return None
    try:
        if os.path.getmtime(candidate) > os.path.getmtime(path):
            return candidate
    except OSError:
        pass
    return None


def resource_path(relative_path):
    """PyInstaller kaynak dosyalarının yolunu döndürür."""
    try:
        base_path = sys._MEIPASS
    except Exception:
        base_path = os.path.dirname(os.path.abspath(__file__))
    return os.path.join(base_path, relative_path)



def configure_tesseract():
    """Tesseract'ı Windows'taki yaygın kurulum yollarında ve PATH içinde arar; bulamazsa None döner."""
    try:
        import pytesseract
    except ImportError:
        return None
    candidates = [
        os.environ.get("TESSERACT_PATH"),
        os.path.join(os.environ.get("PROGRAMFILES", ""), "Tesseract-OCR", "tesseract.exe"),
        os.path.join(os.environ.get("PROGRAMFILES(X86)", ""), "Tesseract-OCR", "tesseract.exe"),
        os.path.join(os.path.dirname(sys.executable), "tesseract.exe"),
        resource_path(os.path.join("tesseract", "tesseract.exe")),
    ]
    for candidate in candidates:
        if candidate and os.path.exists(candidate):
            pytesseract.pytesseract.tesseract_cmd = candidate
            bundled_data = os.path.join(os.path.dirname(candidate), "tessdata")
            if os.path.isdir(bundled_data):
                os.environ["TESSDATA_PREFIX"] = bundled_data
            return candidate
    # Hiçbir bilinen yolda yoksa PATH'e bak (varsayılan komut adı "tesseract").
    return shutil.which(pytesseract.pytesseract.tesseract_cmd)


def load_ocr():
    """pytesseract ve Pillow'u yükler, Tesseract'ı doğrular. Sorun varsa RuntimeError fırlatır."""
    try:
        import pytesseract
        from PIL import Image
    except ImportError as exc:
        raise RuntimeError("OCR için pytesseract ve Pillow kurulmalıdır.") from exc
    tesseract_cmd = configure_tesseract()
    if not tesseract_cmd or not os.path.exists(tesseract_cmd):
        raise RuntimeError(
            "Tesseract bulunamadı. https://github.com/tesseract-ocr/tesseract kurulmalı "
            "veya TESSERACT_PATH ayarlanmalıdır."
        )
    return pytesseract, Image


def pick_ocr_langs(pytesseract, wanted):
    """İstenen diller arasından kurulu olanları 'tur+eng' biçiminde döndürür."""
    try:
        available = set(pytesseract.get_languages(config=""))
    except Exception:
        return "+".join(wanted)
    chosen = [lang for lang in wanted if lang in available]
    if not chosen:
        raise RuntimeError(
            f"Tesseract dil verisi bulunamadı ({', '.join(wanted)}). Dil paketlerini (tessdata) kurun."
        )
    return "+".join(chosen)



def unicode_font_file():
    """Türkçe karakterleri (ğ ş ı İ) çizebilen bir sistem yazı tipi dosyası bulur; yoksa None."""
    windir = os.environ.get("WINDIR", r"C:\Windows")
    for name in ("arial.ttf", "segoeui.ttf", "calibri.ttf", "tahoma.ttf"):
        candidate = os.path.join(windir, "Fonts", name)
        if os.path.exists(candidate):
            return candidate
    return None


def turkish_fold(text):
    """Büyük/küçük harf ve Türkçe i/ı ayrımını yok sayan arama anahtarı."""
    return text.replace("İ", "i").replace("I", "ı").lower().replace("ı", "i")


def insert_text_tr(page, point, text, fontsize=11, color=(0, 0, 0), angle=0, opacity=None, rotate=0):
    """Sayfaya Türkçe karakterleri koruyarak metin ekler (yazı tipi bulunamazsa yerleşik Helvetica).
    angle: derece (saat yönünün tersi); insert_text'in rotate değeri yalnızca 90'ın katlarını kabul eder.
    rotate: 90'ın katı; döndürülmüş sayfada metnin okunur durması için sayfanın dönüşü verilir."""
    font_file = unicode_font_file()
    kwargs = {"fontname": "babatr", "fontfile": font_file} if font_file else {}
    if angle:
        origin = fitz.Point(point)
        kwargs["morph"] = (origin, fitz.Matrix(angle))
    elif rotate:
        kwargs["rotate"] = rotate
    if opacity is not None:
        kwargs["fill_opacity"] = opacity
    page.insert_text(point, text, fontsize=fontsize, color=color, overlay=True, **kwargs)


def text_width_tr(text, fontsize):
    """insert_text_tr ile çizilecek metnin yaklaşık genişliği (nokta)."""
    font_file = unicode_font_file()
    font = fitz.Font(fontfile=font_file) if font_file else fitz.Font("helv")
    return font.text_length(text, fontsize=fontsize)


# --- Yazı ekleme: tam biçim denetimi (yazı tipi, boyut, kalın, italik, altı/üstü çizili, renk, hizalama, açı) ---
TEXT_FONTS = {  # aile: (normal, kalın, italik, kalın italik) Windows yazı tipi dosya adları (.ttf)
    "Arial": ("arial", "arialbd", "ariali", "arialbi"),
    "Times New Roman": ("times", "timesbd", "timesi", "timesbi"),
    "Courier New": ("cour", "courbd", "couri", "courbi"),
    "Calibri": ("calibri", "calibrib", "calibrii", "calibriz"),
    "Segoe UI": ("segoeui", "segoeuib", "segoeuii", "segoeuiz"),
    "Tahoma": ("tahoma", "tahomabd", "tahoma", "tahomabd"),
    "Verdana": ("verdana", "verdanab", "verdanai", "verdanaz"),
    "Georgia": ("georgia", "georgiab", "georgiai", "georgiaz"),
}
TEXT_FONT_DIR = None  # sınama için: yazı tipi klasörünü değiştirir


def _font_dir():
    return TEXT_FONT_DIR or os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts")


def text_font_families():
    """Bu bilgisayarda bulunan (normal biçimi var olan) yazı tipi aileleri."""
    found = [name for name, files in TEXT_FONTS.items() if os.path.exists(os.path.join(_font_dir(), files[0] + ".ttf"))]
    return found or ["Arial"]


def text_font_file(family, bold=False, italic=False):
    """(dosya, sahte_kalın) döndürür. İstenen biçim yoksa normal dosya + kalınsa sahte kalın; hiçbiri yoksa genel yazı tipi."""
    files = TEXT_FONTS.get(family, TEXT_FONTS["Arial"])
    wanted = (1 if bold else 0) + (2 if italic else 0)
    folder = _font_dir()
    path = os.path.join(folder, files[wanted] + ".ttf")
    if os.path.exists(path):
        return path, False
    if bold and not italic:  # kalın dosya yok: normal + sahte kalın
        normal = os.path.join(folder, files[0] + ".ttf")
        if os.path.exists(normal):
            return normal, True
    if italic and bold:  # kalın italik yoksa italik + sahte kalın
        ital = os.path.join(folder, files[2] + ".ttf")
        if os.path.exists(ital):
            return ital, True
    normal = os.path.join(folder, files[0] + ".ttf")
    if os.path.exists(normal):
        return normal, bold
    return unicode_font_file(), bold


def default_text_style():
    return {"family": "Arial", "size": 12.0, "bold": False, "italic": False, "underline": False, "strike": False,
            "color": (0.1, 0.1, 0.1), "opacity": 1.0, "align": "left", "spacing": 1.2, "angle": 0.0}


def text_block_layout(text, style):
    """Yazı bloğunun ölçüsü: (satırlar, satır genişlikleri, genişlik, yükseklik, taban çizgisi yükseltileri, yazı tipi dosyası, sahte_kalın)."""
    path, fake_bold = text_font_file(style["family"], style["bold"], style["italic"])
    font = fitz.Font(fontfile=path) if path else fitz.Font("helv")
    size = float(style["size"])
    lines = text.split("\n") or [""]
    widths = [font.text_length(line, fontsize=size) for line in lines]
    line_h = size * float(style["spacing"])
    asc, desc = font.ascender, font.descender
    pad = (line_h - size * (asc - desc)) / 2
    bases = [i * line_h + pad + size * asc for i in range(len(lines))]
    return lines, widths, max(widths + [1.0]), line_h * len(lines), bases, path, fake_bold


def text_block_bbox(text, style, origin=(0, 0)):
    """Döndürme (açı) uygulanmış bloğun görünen sınır kutusu (sayfa/görünüm uzayında)."""
    _l, _w, width, height, _b, _p, _f = text_block_layout(text, style)
    ox, oy = origin
    rot = fitz.Matrix(-float(style["angle"]))  # görünürde saat yönünün tersi
    corners = [fitz.Point(ox + dx, oy + dy) for dx, dy in ((0, 0), (width, 0), (width, height), (0, height))]
    anchor = fitz.Point(ox, oy)
    pts = [(c - anchor) * rot + anchor for c in corners]
    return fitz.Rect(min(p.x for p in pts), min(p.y for p in pts), max(p.x for p in pts), max(p.y for p in pts))


def draw_text_block(page, origin, text, style):
    """Metin bloğunu sayfaya gerçek (seçilebilir) yazı olarak çizer. origin: bloğun sol üstü, GÖRÜNEN (döndürülmüş)
    sayfa uzayında; açı bu noktanın çevresinde, saat yönünün tersine uygulanır. Döndürülmüş sayfalarda da okunur kalır."""
    lines, widths, width, _height, bases, path, fake_bold = text_block_layout(text, style)
    size = float(style["size"])
    to_page = page.derotation_matrix
    origin = fitz.Point(origin)
    fix = origin * to_page
    angle = float(style["angle"])
    morph = (fix, fitz.Matrix(angle)) if angle else None
    kwargs = {"fontname": "babayz%d" % (1 + (hash_stable(path or "") % 997)), "fontfile": path} if path else {}
    opacity = float(style["opacity"])
    if opacity < 1:
        kwargs["fill_opacity"] = opacity
    if fake_bold:
        kwargs.update(render_mode=2, border_width=max(0.2, size * 0.035))
    if morph:
        kwargs["morph"] = morph
    if page.rotation:
        kwargs["rotate"] = page.rotation
    color = tuple(style["color"])
    line_w = max(0.5, size / 16.0)
    for i, line in enumerate(lines):
        if not line:
            continue
        shift = {"center": (width - widths[i]) / 2, "right": width - widths[i]}.get(style["align"], 0.0)
        x0 = origin.x + shift
        y = origin.y + bases[i]
        page.insert_text(fitz.Point(x0, y) * to_page, line, fontsize=size, color=color, overlay=True, **kwargs)
        for flag, offset in (("underline", size * 0.13), ("strike", -size * 0.28)):
            if style[flag]:
                p1 = fitz.Point(x0, y + offset) * to_page
                p2 = fitz.Point(x0 + widths[i], y + offset) * to_page
                page.draw_line(p1, p2, color=color, width=line_w, stroke_opacity=opacity,
                               morph=morph, overlay=True)


def hash_stable(text):
    return int(hashlib.md5(text.encode("utf-8")).hexdigest()[:6], 16)





# --- Office dışa aktarma yardımcıları (Word / Excel / PowerPoint) --------------------------------------
_OF_FONTS = (("arial", "Arial"), ("helvetica", "Arial"), ("liberationsans", "Arial"), ("times", "Times New Roman"),
             ("liberationserif", "Times New Roman"), ("courier", "Courier New"), ("liberationmono", "Courier New"),
             ("calibri", "Calibri"), ("carlito", "Calibri"), ("cambria", "Cambria"), ("caladea", "Cambria"),
             ("segoeui", "Segoe UI"), ("tahoma", "Tahoma"), ("verdana", "Verdana"), ("georgia", "Georgia"),
             ("trebuchet", "Trebuchet MS"), ("consolas", "Consolas"), ("garamond", "Garamond"),
             ("palatino", "Palatino Linotype"), ("bookantiqua", "Book Antiqua"), ("comicsans", "Comic Sans MS"),
             ("dejavusans", "DejaVu Sans"), ("dejavuserif", "DejaVu Serif"), ("opensans", "Open Sans"),
             ("roboto", "Roboto"), ("lato", "Lato"), ("montserrat", "Montserrat"))
_OF_LIGATURES = {"ﬁ": "fi", "ﬂ": "fl", "ﬀ": "ff", "ﬃ": "ffi", "ﬄ": "ffl", "ﬅ": "st", "ﬆ": "st"}
_OF_BULLET = re.compile(r"^\s*([•▪◦●○■□‣⁃·\-–—*]|\d{1,3}[.)]|[a-zA-Z][.)])\s+")


def office_font_name(raw):
    """PDF'teki (alt kümeli) yazı tipi adını Office'te bulunan bir ada çevirir."""
    base = re.sub(r"^[A-Z]{6}\+", "", raw or "").strip()
    key = re.sub(r"[^a-z]", "", base.lower().split(",")[0].split("-")[0])
    if "symbol" in key or "dingbats" in key or "wingding" in key:
        return "Arial"  # madde imi yazı tipleri Office'te yok: imi "•" yapıldı
    for prefix, name in _OF_FONTS:
        if key.startswith(prefix):
            return name
    name = base.split(",")[0].split("-")[0]
    name = re.sub(r"(PSMT|MT|PS)$", "", name)
    name = re.sub(r"(?<=[a-z])(?=[A-Z])", " ", name).strip()
    return name or "Arial"


def _of_text(value):
    text = PDFProcessor._clean_text(value)
    for lig, plain in _OF_LIGATURES.items():
        text = text.replace(lig, plain)
    for pua in ("\uf0b7", "\uf0a7", "\uf0d8", "\uf0fc", "\uf076", "\uf02d"):
        text = text.replace(pua, "•")
    return text.replace("­", "")


def _of_hex(color):
    return "%02X%02X%02X" % ((color >> 16) & 255, (color >> 8) & 255, color & 255)


def _of_marks(page):
    """Sayfadaki ince yatay çizgiler (altı/üstü çizgi) ve dolgulu dikdörtgenler (hücre/zemin rengi)."""
    hlines, fills = [], []
    try:
        drawings = page.get_drawings()
    except Exception:  # noqa: BLE001
        return hlines, fills
    for d in drawings:
        for item in d.get("items", ()):
            if item[0] == "l":
                p1, p2 = item[1], item[2]
                if abs(p1.y - p2.y) < 1.2 and abs(p2.x - p1.x) > 3:
                    hlines.append((min(p1.x, p2.x), max(p1.x, p2.x), (p1.y + p2.y) / 2))
            elif item[0] == "re":
                r = item[1]
                if r.height < 1.6 and r.width > 3:
                    hlines.append((r.x0, r.x1, (r.y0 + r.y1) / 2))
        fill = d.get("fill")
        if fill and d.get("rect") is not None and d["rect"].width > 4 and d["rect"].height > 4:
            rgb = tuple(int(round(c * 255)) for c in fill[:3])
            if rgb != (255, 255, 255):
                fills.append((fitz.Rect(d["rect"]), rgb))
    return hlines, fills


def _of_span(span, hlines):
    """PDF span'ından Office biçim sözlüğü: yazı tipi, boyut, kalın, italik, renk, altı/üstü çizili, üst simge."""
    name = span.get("font", "")
    low = name.lower()
    size = float(span.get("size", 11))
    flags = span.get("flags", 0)
    x0, _y0, x1, _y1 = span["bbox"]
    base = span.get("origin", (x0, span["bbox"][3]))[1]
    width = max(x1 - x0, 1.0)
    underline = strike = False
    for lx0, lx1, ly in hlines:
        if lx0 >= x0 - 4 and lx1 <= x1 + 4 and (lx1 - lx0) >= 0.8 * width:
            if base + 0.4 <= ly <= base + 0.28 * size + 1.5:
                underline = True
            elif base - 0.5 * size - 1.5 <= ly <= base - 0.2 * size + 1.5:
                strike = True
    return {"font": office_font_name(name), "size": size, "bold": bool(flags & 16) or "bold" in low or "black" in low,
            "italic": bool(flags & 2) or "italic" in low or "oblique" in low, "color": span.get("color", 0),
            "underline": underline, "strike": strike, "sup": bool(flags & 1)}


def _of_xy_cut(items):
    """Okuma sırası: özyinelemeli XY-kesimi (önce yatay boşluklar, sonra dikey boşluklar = sütunlar)."""
    if len(items) <= 1:
        return list(items)

    def split(axis, gap):
        order = sorted(items, key=lambda it: it["bbox"][axis])
        groups, end = [[order[0]]], order[0]["bbox"][axis + 2]
        for it in order[1:]:
            if it["bbox"][axis] >= end + gap:
                groups.append([it])
                end = it["bbox"][axis + 2]
            else:
                groups[-1].append(it)
                end = max(end, it["bbox"][axis + 2])
        return groups

    groups = split(1, 0.0)
    if len(groups) == 1:
        groups = split(0, 8.0)
    if len(groups) > 1:
        return [x for g in groups for x in _of_xy_cut(g)]
    return sorted(items, key=lambda it: (it["bbox"][1], it["bbox"][0]))


def _of_align(bbox, left, right, nlines=2, justified_hint=False):
    """Bir bloğun hizası: left / center / right / justify (kapsayıcının sol-sağ sınırlarına göre)."""
    x0, x1 = bbox[0], bbox[2]
    width = max(right - left, 1.0)
    if justified_hint and (x1 - x0) > 0.6 * width and abs(x0 - left) < 6 and abs(right - x1) < 6:
        return "justify"
    lo, ro = x0 - left, right - x1
    if lo > 6 and ro > 6 and abs(lo - ro) < max(4.0, 0.02 * width):
        return "center"
    if ro < 8 and lo > max(0.25 * width, ro + 6):
        return "right"
    return "left"


def _of_block_infos(block, hlines):
    """Bloğu bilgi sözlüklerine çevirir: aynı satırdaki parçalar (madde imi + metin) birleştirilir,
    madde imiyle başlayan her satır ayrı paragraf olur."""
    rows = []
    for line in sorted(block.get("lines", ()), key=lambda ln: (ln["bbox"][1], ln["bbox"][0])):
        if not any(sp.get("text", "").strip() for sp in line.get("spans", ())):
            continue
        for row in rows:
            top, bottom = max(row["bbox"][1], line["bbox"][1]), min(row["bbox"][3], line["bbox"][3])
            if bottom - top > 0.6 * min(row["bbox"][3] - row["bbox"][1], line["bbox"][3] - line["bbox"][1]):
                first, second = (row, line) if row["bbox"][0] <= line["bbox"][0] else (line, row)
                gap = second["bbox"][0] - first["bbox"][2]
                spans = list(first["spans"])
                if gap > 1.0 and spans and not spans[-1]["text"].endswith(" "):
                    spans.append(dict(spans[-1], text=" "))
                spans += list(second["spans"])
                row["spans"] = spans
                row["bbox"] = (first["bbox"][0], min(row["bbox"][1], line["bbox"][1]),
                               max(row["bbox"][2], line["bbox"][2]), max(row["bbox"][3], line["bbox"][3]))
                break
        else:
            rows.append({"bbox": tuple(line["bbox"]), "spans": list(line["spans"])})
    groups = []
    for row in rows:
        text = _of_text("".join(sp["text"] for sp in row["spans"]))
        if groups and not _OF_BULLET.match(text):
            groups[-1].append(row)
        else:
            groups.append([row])
    result = []
    for group in groups:
        bbox = (min(r["bbox"][0] for r in group), min(r["bbox"][1] for r in group),
                max(r["bbox"][2] for r in group), max(r["bbox"][3] for r in group))
        info = _of_block_info1({"lines": group, "bbox": bbox}, hlines)
        if info:
            result.append(info)
    return result


def _of_block_info1(block, hlines):
    """Blok -> satırlar (span biçimleriyle), birleştirme kararları, satır aralığı ve genişlik bilgisi."""
    lines = []
    for line in block.get("lines", ()):
        spans = [(_of_text(sp["text"]), _of_span(sp, hlines)) for sp in line.get("spans", ()) if sp.get("text")]
        if not any(t.strip() for t, _s in spans):
            continue
        lines.append({"bbox": line["bbox"], "spans": spans})
    if not lines:
        return None
    pitch = None
    if len(lines) > 1:
        gaps = [lines[i + 1]["bbox"][1] - lines[i]["bbox"][1] for i in range(len(lines) - 1)]
        gaps = [g for g in gaps if g > 1]
        if gaps:
            pitch = sorted(gaps)[len(gaps) // 2]
    sizes = [s["size"] for ln in lines for _t, s in ln["spans"]]
    size = max(sizes) if sizes else 11.0
    if pitch is None:
        pitch = max(lines[0]["bbox"][3] - lines[0]["bbox"][1], size * 1.15)
    widths = [ln["bbox"][2] - ln["bbox"][0] for ln in lines]
    longest = max(widths)
    justified = len(lines) >= 3 and sum(abs(ln["bbox"][2] - block["bbox"][2]) < 2.0 for ln in lines[:-1]) >= 0.7 * (len(lines) - 1)
    return {"lines": lines, "pitch": pitch, "size": size, "longest": longest, "justified": justified,
            "bbox": tuple(block["bbox"])}


def _of_join(info):
    """Satırları paragraf metnine çevirir: (metin, biçim) parçaları + satır sonu kararları (None = satır sonu ekle)."""
    parts = []
    lines = info["lines"]
    for i, ln in enumerate(lines):
        spans = list(ln["spans"])
        last = i == len(lines) - 1
        if not last and spans:
            text, st = spans[-1]
            nxt = lines[i + 1]["spans"][0][0] if lines[i + 1]["spans"] else ""
            width = ln["bbox"][2] - ln["bbox"][0]
            if text.rstrip().endswith("-") and nxt[:1].islower() and len(text.rstrip()) > 2 and text.rstrip()[-2].isalpha():
                spans[-1] = (text.rstrip()[:-1], st)  # satır sonu tire: kelimeyi birleştir
                glue = ""
            elif width < 0.8 * info["longest"] or len(lines) == 1:
                glue = None  # kısa satır: gerçek satır sonu
            else:
                glue = " "
            parts.extend(spans)
            if glue is None:
                parts.append((None, None))
            elif glue:
                parts.append((glue, spans[-1][1]))
        else:
            parts.extend(spans)
    return parts


def _of_cluster_images(doc, page, tables_bbox, min_area=1500.0):
    """Vektör çizim kümelerini (grafik, şekil, ayraç) resim olarak alır; tablo ve ince çizgiler hariç."""
    out = []
    try:
        clusters = page.cluster_drawings()
    except Exception:  # noqa: BLE001
        return out
    page_area = page.rect.width * page.rect.height
    for rect in clusters:
        rect = fitz.Rect(rect) & page.rect
        if rect.is_empty or rect.width < 12 or rect.height < 12 or rect.width * rect.height < min_area:
            continue
        if rect.width * rect.height > 0.8 * page_area:
            continue  # tüm sayfa zemini
        if any((rect & tb).get_area() > 0.5 * rect.get_area() for tb in tables_bbox):
            continue
        pix = page.get_pixmap(matrix=fitz.Matrix(2.5, 2.5), clip=rect, alpha=False)
        out.append((rect, pix.tobytes("png")))
    return out


def _of_page_items(doc, page, with_tables=True):
    """Sayfanın öğeleri okuma sırasıyla: metin blokları, tablolar, resimler, çizim kümeleri; ayrıca kenar boşlukları."""
    hlines, fills = _of_marks(page)
    tables = []
    if with_tables:
        try:
            tables = [t for t in page.find_tables().tables if t.row_count >= 1 and t.col_count >= 2 and len(t.cells) >= 2]
        except Exception:  # noqa: BLE001
            tables = []
    tbboxes = [fitz.Rect(t.bbox) for t in tables]
    clusters = _of_cluster_images(doc, page, tbboxes)
    cboxes = [r for r, _d in clusters]
    items = []
    flags = fitz.TEXT_PRESERVE_WHITESPACE | fitz.TEXT_MEDIABOX_CLIP
    for block in page.get_text("dict", flags=flags)["blocks"]:
        if block.get("type") != 0:
            continue
        center = fitz.Point((block["bbox"][0] + block["bbox"][2]) / 2, (block["bbox"][1] + block["bbox"][3]) / 2)
        if any(center in tb for tb in tbboxes) or any(center in cb for cb in cboxes):
            continue  # tablonun / şeklin içindeki yazı o öğeyle birlikte gider
        for info in _of_block_infos(block, hlines):
            items.append({"kind": "text", "bbox": fitz.Rect(info["bbox"]), "info": info})
    try:
        drawn = [fitz.Rect(d["rect"]) for d in page.get_drawings()]
    except Exception:  # noqa: BLE001
        drawn = []
    for table in tables:
        tb = fitz.Rect(table.bbox)
        ruled = any(d.x1 >= tb.x0 - 2 and d.x0 <= tb.x1 + 2 and d.y1 >= tb.y0 - 2 and d.y0 <= tb.y1 + 2 for d in drawn)
        items.append({"kind": "table", "bbox": tb, "table": table, "ruled": ruled})
    page_area = page.rect.width * page.rect.height
    for bbox, data in PDFProcessor._page_images(doc, page):
        if bbox.width * bbox.height > 0.85 * page_area:
            continue  # sayfa zemini resmi
        if any((bbox & cb).get_area() > 0.5 * bbox.get_area() for cb in cboxes):
            continue
        items.append({"kind": "image", "bbox": fitz.Rect(bbox), "data": data})
    for rect, data in clusters:
        items.append({"kind": "image", "bbox": rect, "data": data})
    items = _of_xy_cut(items)
    if items:
        left = max(14.0, min(it["bbox"].x0 for it in items))
        right = max(14.0, page.rect.width - max(it["bbox"].x1 for it in items))
        top = max(14.0, min(it["bbox"].y0 for it in items))
        bottom = max(14.0, page.rect.height - max(it["bbox"].y1 for it in items))
    else:
        left = right = top = bottom = 36.0
    margins = tuple(min(v, 0.45 * d) for v, d in zip((left, right, top, bottom),
                                                      (page.rect.width, page.rect.width, page.rect.height, page.rect.height)))
    return items, margins, hlines, fills


def _of_grid(table, tol=1.5):
    """Tablo hücrelerinden satır/sütun sınırlarını ve birleşik hücre aralıklarını çıkarır."""
    def edges(values):
        out = []
        for v in sorted(values):
            if not out or v - out[-1] > tol:
                out.append(v)
        return out
    cells = [fitz.Rect(c) for c in table.cells if c]
    xs = edges([c.x0 for c in cells] + [c.x1 for c in cells])
    ys = edges([c.y0 for c in cells] + [c.y1 for c in cells])

    def idx(seq, v):
        return min(range(len(seq)), key=lambda i: abs(seq[i] - v))
    spans = []
    for c in cells:
        c0, c1 = idx(xs, c.x0), idx(xs, c.x1)
        r0, r1 = idx(ys, c.y0), idx(ys, c.y1)
        if c1 > c0 and r1 > r0:
            spans.append((r0, c0, r1 - 1, c1 - 1, c))
    return xs, ys, spans


def _of_cell_fill(rect, fills):
    best = None
    for frect, rgb in fills:
        if (frect & rect).get_area() >= 0.7 * rect.get_area() and frect.get_area() <= 1.6 * rect.get_area() + 50:
            best = rgb
    return best


def _of_cell_blocks(page, rect, hlines):
    """Hücrenin içindeki yazı blokları (bilgi sözlükleri)."""
    flags = fitz.TEXT_PRESERVE_WHITESPACE | fitz.TEXT_MEDIABOX_CLIP
    out = []
    for block in page.get_text("dict", clip=rect, flags=flags)["blocks"]:
        if block.get("type") == 0:
            out.extend(_of_block_infos(block, hlines))
    return out


def _of_number(text):
    """Metin sayıya benziyorsa (Türkçe biçim öncelikli) (değer, sayı_biçimi) döndürür; değilse None."""
    s = text.strip().replace(" ", "").replace(" ", "")
    if not s or len(s) > 24:
        return None
    pct = False
    if s.startswith("%") and len(s) > 1:
        s, pct = s[1:], True
    elif s.endswith("%") and len(s) > 1:
        s, pct = s[:-1], True
    patterns = ((r"^[-+]?\d{1,3}(,\d{3})+\.\d+$", lambda v: v.replace(",", "")),
                (r"^[-+]?\d{1,3}(\.\d{3})+(,\d+)?$", lambda v: v.replace(".", "").replace(",", ".")),
                (r"^[-+]?\d+,\d+$", lambda v: v.replace(",", ".")),
                (r"^[-+]?\d+\.\d+$", lambda v: v),
                (r"^[-+]?\d+$", lambda v: v))
    for pattern, conv in patterns:
        if re.match(pattern, s):
            if re.match(r"^[-+]?0\d", s) or len(re.sub(r"\D", "", s)) > 15:
                return None  # 0532… gibi kimlik/telefon dizileri metin kalır
            try:
                value = float(conv(s))
            except ValueError:
                return None
            decimals = len(re.split(r"[.,]", s)[-1]) if re.search(r"[.,]\d{1,2}$", s) else 0
            grouped = bool(re.search(r"\d[.,]\d{3}(?!\d)", s)) and not re.match(r"^[-+]?\d+\.\d{1,2}$", s)
            if pct:
                return value / 100.0, "0.00%" if decimals else "0%"
            fmt = ("#,##0." + "0" * decimals) if decimals else ("#,##0" if grouped else "General")
            if value == int(value) and not decimals:
                value = int(value)
            return value, fmt
    return None


def _of_columns(items, left, right):
    """Yan yana duran metin sütunlarını (gazete/makale düzeni) tek 'cols' öğesinde toplar; öğe sırası korunur."""
    width = max(right - left, 1.0)
    texts = [it for it in items if it["kind"] == "text" and it["bbox"].width < 0.62 * width]
    groups = []
    for it in sorted(texts, key=lambda i: i["bbox"].x0):
        for g in groups:
            if abs(g[0]["bbox"].x0 - it["bbox"].x0) < 8:
                g.append(it)
                break
        else:
            groups.append([it])
    runs = []  # her sütun adayını dikey boşluklara göre parçalara ayır; yeterince yüksek parçalar aday
    for g in groups:
        g = sorted(g, key=lambda i: i["bbox"].y0)
        run = [g[0]]
        for it in g[1:]:
            if it["bbox"].y0 - run[-1]["bbox"].y1 > 28:
                runs.append(run)
                run = []
            run.append(it)
        runs.append(run)
    runs = [r for r in runs if r[-1]["bbox"].y1 - r[0]["bbox"].y0 > 60]
    cols = None
    for r in runs:
        y0, y1 = r[0]["bbox"].y0, r[-1]["bbox"].y1
        mates = [r]
        for o in runs:
            if o is r or abs(o[0]["bbox"].x0 - r[0]["bbox"].x0) < 20:
                continue
            oy0, oy1 = o[0]["bbox"].y0, o[-1]["bbox"].y1
            if min(y1, oy1) - max(y0, oy0) > 0.5 * min(y1 - y0, oy1 - oy0) and all(
                    abs(m[0]["bbox"].x0 - o[0]["bbox"].x0) >= 20 for m in mates):
                mates.append(o)
        if len(mates) >= 2:
            cols = mates
            break
    if not cols:
        return items
    extent = [(min(i["bbox"].y0 for i in g), max(i["bbox"].y1 for i in g)) for g in cols]
    cols.sort(key=lambda g: g[0]["bbox"].x0)
    for a, b in zip(cols, cols[1:]):
        if max(i["bbox"].x1 for i in a) > b[0]["bbox"].x0 + 2:
            return items  # sütunlar karışıyor
    members = {id(i) for g in cols for i in g}
    ymin = min(e[0] for e in extent)
    ymax = max(e[1] for e in extent)
    out, placed = [], False
    for it in items:
        if id(it) in members:
            if not placed:
                xs = [g[0]["bbox"].x0 for g in cols]
                rights = [max(i["bbox"].x1 for i in g) for g in cols]
                out.append({"kind": "cols", "bbox": fitz.Rect(xs[0], ymin, rights[-1], ymax), "cols": cols})
                placed = True
            continue
        out.append(it)
    return out


def _of_col_aligns(page, spans, hlines):
    """Tablo sütunlarının ortak hizası: sütundaki yazıların sağ (veya orta) kenarları çakışıyorsa o hiza."""
    cols = {}
    for r0, c0, r1, c1, rect in spans:
        if c1 > c0:
            continue
        blocks = _of_cell_blocks(page, fitz.Rect(rect.x0 + 0.5, rect.y0 + 0.5, rect.x1 - 0.5, rect.y1 - 0.5), hlines)
        if blocks:
            cols.setdefault(c0, []).append((min(b["bbox"][0] for b in blocks), max(b["bbox"][2] for b in blocks)))
    result = {}
    for c, lst in cols.items():
        if len(lst) < 2:
            continue
        x0s, x1s = [a for a, _b in lst], [b for _a, b in lst]
        spread0, spread1 = max(x0s) - min(x0s), max(x1s) - min(x1s)
        mids = [(a + b) / 2 for a, b in lst]
        if spread1 < 2.5 and spread0 > 3:
            result[c] = "right"
        elif max(mids) - min(mids) < 2.5 and spread0 > 3:
            result[c] = "center"
        elif spread0 < 2.5:
            result[c] = "left"
    return result


class PDFProcessor:
    def __init__(self, pdf_path):
        self.pdf_path = pdf_path
        if not os.path.exists(pdf_path):
            raise FileNotFoundError(f"Dosya bulunamadı: {pdf_path}")

    def extract_text(self, output_txt_path):
        text = ""
        with pdfplumber.open(self.pdf_path) as pdf:
            for i, page in enumerate(pdf.pages):
                page_text = page.extract_text()
                if page_text:
                    text += f"--- Sayfa {i+1} ---\n" + page_text + "\n\n"

        with open(output_txt_path, "w", encoding="utf-8") as f:
            f.write(text)

    def convert_to_images(self, output_folder):
        os.makedirs(output_folder, exist_ok=True)
        with fitz.open(self.pdf_path) as doc:
            for i, page in enumerate(doc):
                pix = page.get_pixmap(dpi=150)
                image_path = os.path.join(output_folder, f"sayfa_{i+1}.png")
                pix.save(image_path)

    @staticmethod
    def parse_page_range(value, total):
        pages = set()
        for part in value.replace(';', ',').split(','):
            part = part.strip()
            if not part:
                continue
            if '-' in part:
                left, right = [int(x.strip()) for x in part.split('-', 1)]
                pages.update(range(max(1, left), min(total, right) + 1))
            else:
                page = int(part)
                if 1 <= page <= total:
                    pages.add(page)
        return sorted(pages)

    @staticmethod
    def extract_pages(pdf_path, page_numbers, output_path):
        with fitz.open(pdf_path) as source, fitz.open() as output:
            for number in page_numbers:
                output.insert_pdf(source, from_page=number - 1, to_page=number - 1)
            output.save(output_path, garbage=4, deflate=True, clean=True)

    @staticmethod
    def compress(pdf_path, output_path):
        with fitz.open(pdf_path) as doc:
            doc.save(output_path, garbage=4, deflate=True, clean=True, deflate_images=True, deflate_fonts=True)

    @staticmethod
    def create_contact_sheet(pdf_path, output_path, dpi=72, columns=3):
        from PIL import Image, ImageDraw
        with fitz.open(pdf_path) as doc:
            images = []
            for page in doc:
                pix = page.get_pixmap(dpi=dpi, alpha=False)
                images.append(Image.frombytes('RGB', [pix.width, pix.height], pix.samples))
        if not images:
            return
        width = max(image.width for image in images)
        height = max(image.height for image in images)
        rows = (len(images) + columns - 1) // columns
        sheet = Image.new('RGB', (width * columns, height * rows), 'white')
        draw = ImageDraw.Draw(sheet)
        for index, image in enumerate(images):
            x = (index % columns) * width
            y = (index // columns) * height
            sheet.paste(image, (x, y))
            draw.rectangle((x, y, x + image.width - 1, y + image.height - 1), outline='#999999')
            draw.text((x + 8, y + 8), str(index + 1), fill='#cc0000')
        sheet.save(output_path)

    @staticmethod
    def check_document(doc, max_render_pages=300):
        """Belgeyi sınar: sayfalar yüklenip çiziliyor mu, otomatik onarım yapıldı mı, MuPDF uyarıları. Metin satırları döner."""
        try:
            fitz.TOOLS.mupdf_warnings()  # önceki uyarıları sıfırla
        except Exception:  # noqa: BLE001
            pass
        lines = [f"Sayfa sayısı: {len(doc)}", f"Şifreli: {'evet' if doc.is_encrypted else 'hayır'}"]
        repaired = bool(getattr(doc, "is_repaired", False))
        bad = []
        for index in range(len(doc)):
            try:
                page = doc[index]
                page.get_text("text")
                if index < max_render_pages:
                    page.get_pixmap(matrix=fitz.Matrix(0.1, 0.1), alpha=False)
            except Exception as exc:  # noqa: BLE001
                bad.append(f"Sayfa {index + 1}: {exc}")
        lines.append("Dosya yapısı: " + ("BOZUKTU, açılırken otomatik onarıldı" if repaired else "sağlam görünüyor"))
        if len(doc) > max_render_pages:
            lines.append(f"Not: yalnızca ilk {max_render_pages} sayfa çizilerek denendi; kalanında metin okuması denendi.")
        if bad:
            lines.append(f"Sorunlu sayfa sayısı: {len(bad)}")
            lines.extend(bad[:10])
        else:
            lines.append("Sorunlu sayfa: yok")
        try:
            warnings_text = (fitz.TOOLS.mupdf_warnings() or "").strip()
        except Exception:  # noqa: BLE001
            warnings_text = ""
        if warnings_text:
            warn_lines = warnings_text.splitlines()
            lines.append(f"MuPDF uyarıları ({len(warn_lines)}):")
            lines.extend("  " + line for line in warn_lines[:8])
        return lines, (repaired or bool(bad))

    @staticmethod
    def repair_pdf(pdf_path, output_path):
        """Dosyayı açıp (MuPDF yapıyı onarır) yapıyı baştan yazarak temiz bir kopya üretir."""
        with fitz.open(pdf_path) as doc:
            doc.save(output_path, garbage=4, deflate=True, clean=True)

    @staticmethod
    def crop_margins(doc, margin_pt, page_indices):
        """Seçili sayfaların dört kenarından margin_pt (punto) kırpar; sayfa dönüklüğüne uyar."""
        for index in page_indices:
            page = doc[index]
            rect = page.rect
            if margin_pt * 2 >= min(rect.width, rect.height) - 10:
                raise ValueError(f"Sayfa {index + 1} için kenar boşluğu çok büyük.")
            inner = fitz.Rect(rect.x0 + margin_pt, rect.y0 + margin_pt, rect.x1 - margin_pt, rect.y1 - margin_pt)
            page.set_cropbox(PDFProcessor.cropbox_from_visible(page, inner))  # daha önce kırpılmış sayfada kayma olmaz

    @staticmethod
    def cropbox_from_visible(page, visible_rect):
        """Ekranda görünen (kırpık/döndürülmüş sayfa uzayındaki) dikdörtgeni set_cropbox'ın beklediği döndürülmemiş
        MediaBox uzayına çevirir. Sayfa daha önce kırpıldıysa kırpma kutusunun kaymasını (ofsetini) da ekler."""
        rect = fitz.Rect(visible_rect).normalize() & page.rect
        if rect.is_empty or rect.width < 10 or rect.height < 10:
            raise ValueError("Kırpılacak alan çok küçük.")
        api = rect * page.derotation_matrix
        origin = page.cropbox.tl
        return (api + (origin.x, origin.y, origin.x, origin.y)) & page.mediabox

    @staticmethod
    def crop_to_rect(doc, page_indices, visible_rect):
        """Seçili sayfaları görünen dikdörtgene kırpar (her sayfa kendi görünen uzayında, sayfa dışı kısım atılır)."""
        for index in page_indices:
            page = doc[index]
            page.set_cropbox(PDFProcessor.cropbox_from_visible(page, visible_rect))

    @staticmethod
    def reset_crop(doc, page_indices):
        """Kırpmayı kaldırır (kırpma kutusu = MediaBox); kırpılmış olan sayfa sayısını döndürür."""
        count = 0
        for index in page_indices:
            page = doc[index]
            if page.cropbox != page.mediabox:
                page.set_cropbox(page.mediabox)
                count += 1
        return count

    @staticmethod
    def insert_blank_pages(doc, after_index, count=1):
        """after_index sayfasından sonra (-1: en başa) count boş sayfa ekler; boyut, komşu sayfanın görünen boyutudur."""
        if not 1 <= count <= 500:
            raise ValueError("Boş sayfa sayısı 1 ile 500 arasında olmalı.")
        reference = doc[min(max(after_index, 0), len(doc) - 1)]
        width, height = reference.rect.width, reference.rect.height
        for k in range(count):
            doc.new_page(pno=after_index + 1 + k, width=width, height=height)

    @staticmethod
    def delete_page_numbers(doc, numbers):
        """1 tabanlı sayfa numaralarını siler; en az bir sayfa kalmalıdır. Silinen sayfa sayısını döndürür."""
        indices = sorted({n - 1 for n in numbers if 1 <= n <= len(doc)})
        if not indices:
            raise ValueError("Belgede bu aralıkta sayfa yok.")
        if len(indices) >= len(doc):
            raise ValueError("Belgede en az bir sayfa kalmalıdır.")
        doc.delete_pages(indices)
        return len(indices)

    # --- Office biçimleri (Word / Excel / PowerPoint) -------------------------------------------
    _XML_BAD = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")

    @staticmethod
    def _clean_text(value):
        """Office XML'inde geçersiz denetim karakterlerini atar (python-docx/openpyxl bunlarda hata verir)."""
        return PDFProcessor._XML_BAD.sub("", "" if value is None else str(value))

    @staticmethod
    def _page_images(doc, page, min_size=40):
        """Sayfadaki (bbox, bayt) resimleri; çok küçükler ve okunamayanlar atlanır."""
        found = []
        try:
            infos = page.get_image_info(xrefs=True)
        except Exception:  # noqa: BLE001
            return found
        for info in infos:
            xref = info.get("xref") or 0
            if xref <= 0 or info.get("width", 0) < min_size or info.get("height", 0) < min_size:
                continue
            try:
                data = doc.extract_image(xref)["image"]
            except Exception:  # noqa: BLE001
                continue
            found.append((fitz.Rect(info["bbox"]), data))
        return found

    @staticmethod
    def export_docx(pdf_path, output_path, editable=True, progress=None):
        """Word (.docx): editable=True sayfa düzenini korur: sayfa boyutu ve kenar boşlukları, yazı tipi/boyut/renk,
        kalın-italik-altı çizili, hizalama ve girinti, satır aralığı, gerçek tablolar (birleşik hücre ve zemin rengi
        dahil), resimler ve vektör şekiller (resim olarak) yerinde, sütunlu sayfalarda okuma sırası. Metinsiz
        (taranmış) sayfalar ve editable=False durumunda sayfa resim olarak konur."""
        import io
        from collections import Counter
        from docx import Document
        from docx.enum.section import WD_SECTION
        from docx.enum.table import WD_ROW_HEIGHT_RULE
        from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Pt, RGBColor

        out = Document()
        alignments = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                      "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}

        def style_run(run, st):
            run.font.name = st["font"]
            rpr = run._element.get_or_add_rPr()
            fonts = rpr.find(qn("w:rFonts"))
            if fonts is not None:
                fonts.set(qn("w:eastAsia"), st["font"])
                fonts.set(qn("w:cs"), st["font"])
            run.font.size = Pt(max(1, min(round(st["size"] * 2) / 2, 400)))
            run.bold = st["bold"] or None
            run.italic = st["italic"] or None
            if st["underline"]:
                run.underline = True
            if st["strike"]:
                run.font.strike = True
            if st["sup"]:
                run.font.superscript = True
            if st["color"] & 0xFFFFFF:
                run.font.color.rgb = RGBColor((st["color"] >> 16) & 255, (st["color"] >> 8) & 255, st["color"] & 255)

        def fill_par(par, info):
            run = None
            for text, st in _of_join(info):
                if text is None:
                    (run or par.add_run()).add_break()
                    continue
                if not text:
                    continue
                run = par.add_run(text)
                style_run(run, st)

        def set_spacing(par, info, before):
            pf = par.paragraph_format
            pf.space_before = Pt(max(0.0, min(before, 600.0)))
            pf.space_after = Pt(0)
            pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            pf.line_spacing = Pt(max(info["pitch"], info["size"] * 1.05))

        def add_text(container, info, left, right, cursor):
            par = container.add_paragraph()
            first = info["lines"][0]
            align = _of_align(info["bbox"], left, right, len(info["lines"]), info["justified"])
            par.paragraph_format.alignment = alignments[align]
            if align in ("left", "justify"):
                x0 = info["bbox"][0]
                if len(info["lines"]) > 1:
                    rest = min(ln["bbox"][0] for ln in info["lines"][1:])
                    if first["bbox"][0] - rest > 3 and _OF_BULLET.match("".join(t for t, _s in first["spans"])):
                        par.paragraph_format.left_indent = Pt(max(rest - left, 0))  # madde imi: asılı girinti
                        par.paragraph_format.first_line_indent = Pt(first["bbox"][0] - rest)
                    else:
                        par.paragraph_format.left_indent = Pt(max(rest - left, 0))
                        if first["bbox"][0] - rest > 3:
                            par.paragraph_format.first_line_indent = Pt(first["bbox"][0] - rest)
                        elif rest - first["bbox"][0] > 3:
                            par.paragraph_format.first_line_indent = Pt(first["bbox"][0] - rest)
                elif x0 - left > 3:
                    par.paragraph_format.left_indent = Pt(x0 - left)
            before = info["bbox"][1] - cursor if cursor is not None else info["bbox"][1]
            set_spacing(par, info, before)
            fill_par(par, info)
            return info["bbox"][1] + len(info["lines"]) * max(info["pitch"], info["size"] * 1.05)

        def gap_before(item, cursor):
            gap = item["bbox"].y0 - cursor if cursor is not None else 0
            if gap > 2:
                par = out.add_paragraph()
                pf = par.paragraph_format
                pf.space_before = pf.space_after = Pt(0)
                pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
                pf.line_spacing = Pt(min(gap, 600.0))

        def add_cols(item, page, left, cursor):
            gap_before(item, cursor)
            cols = item["cols"]
            tbl = out.add_table(rows=1, cols=len(cols))
            tbl.autofit = False
            ind = OxmlElement("w:tblInd")
            ind.set(qn("w:w"), str(int(max(cols[0][0]["bbox"].x0 - left, 0) * 20)))
            ind.set(qn("w:type"), "dxa")
            tbl._tbl.tblPr.append(ind)
            starts = [g[0]["bbox"].x0 for g in cols]
            ends = starts[1:] + [max(i["bbox"].x1 for i in cols[-1])]
            for j, group in enumerate(cols):
                cell = tbl.cell(0, j)
                cell.width = Pt(ends[j] - starts[j])
                first_par = cell.paragraphs[0]
                top_cursor = item["bbox"].y0
                cur = None
                for item2 in group:
                    cur = add_text(cell, item2["info"], starts[j], max(i["bbox"].x1 for i in group), cur if cur is not None else top_cursor)
                first_par._element.getparent().remove(first_par._element)
            return item["bbox"].y1

        def add_image(item, left, right, pw, cursor):
            bbox = item["bbox"]
            par = out.add_paragraph()
            pf = par.paragraph_format
            pf.space_after = Pt(0)
            pf.space_before = Pt(max(0.0, min(bbox.y0 - cursor, 600.0)) if cursor is not None else bbox.y0)
            align = _of_align(bbox, left, right)
            if align == "center":
                pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif bbox.x0 - left > 3:
                pf.left_indent = Pt(bbox.x0 - left)
            width = min(bbox.width, max(pw - left - right, 40))
            try:
                par.add_run().add_picture(io.BytesIO(item["data"]), width=Pt(width), height=Pt(bbox.height * width / bbox.width))
            except Exception:  # noqa: BLE001 - tanınmayan resim biçimi
                return cursor
            return bbox.y1

        def add_table(item, page, hlines, fills, left, cursor):
            table = item["table"]
            xs, ys, spans = _of_grid(table)
            if len(xs) < 2 or len(ys) < 2 or not spans:
                return cursor
            gap_before(item, cursor)
            tbl = out.add_table(rows=len(ys) - 1, cols=len(xs) - 1)
            if item.get("ruled", True):
                tbl.style = "Table Grid"
            col_align = _of_col_aligns(page, spans, hlines)
            tbl.autofit = False
            props = tbl._tbl.tblPr
            ind = OxmlElement("w:tblInd")
            ind.set(qn("w:w"), str(int(max(xs[0] - left, 0) * 20)))
            ind.set(qn("w:type"), "dxa")
            props.append(ind)
            for i, row in enumerate(tbl.rows):
                row.height = Pt(max(ys[i + 1] - ys[i], 4))
                row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
                for j, cell in enumerate(row.cells):
                    cell.width = Pt(xs[j + 1] - xs[j])
            for r0, c0, r1, c1, rect in spans:
                cell = tbl.cell(r0, c0)
                if r1 > r0 or c1 > c0:
                    cell = cell.merge(tbl.cell(r1, c1))
                fill = _of_cell_fill(rect, fills)
                if fill:
                    shade = OxmlElement("w:shd")
                    shade.set(qn("w:val"), "clear")
                    shade.set(qn("w:color"), "auto")
                    shade.set(qn("w:fill"), "%02X%02X%02X" % fill)
                    cell._tc.get_or_add_tcPr().append(shade)
                blocks = _of_cell_blocks(page, fitz.Rect(rect.x0 + 0.5, rect.y0 + 0.5, rect.x1 - 0.5, rect.y1 - 0.5), hlines)
                first = True
                for info in blocks:
                    par = cell.paragraphs[0] if first else cell.add_paragraph()
                    first = False
                    par.paragraph_format.alignment = alignments[col_align.get(c0) or _of_align(info["bbox"], rect.x0, rect.x1)]
                    set_spacing(par, info, 0)
                    fill_par(par, info)
            return table.bbox[3]

        with fitz.open(pdf_path) as doc:
            total = len(doc)
            normal = out.styles["Normal"]
            normal.paragraph_format.space_after = Pt(0)
            normal.paragraph_format.space_before = Pt(0)
            for number, page in enumerate(doc, 1):
                rect = page.rect
                section = out.sections[0] if number == 1 else out.add_section(WD_SECTION.NEW_PAGE)
                if number > 1:
                    marker = out.paragraphs[-1]  # bölüm sonu paragrafı sayfayı taşırmasın
                    marker.paragraph_format.space_before = Pt(0)
                    marker.paragraph_format.space_after = Pt(0)
                    marker.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
                    marker.paragraph_format.line_spacing = Pt(1)
                section.page_width, section.page_height = Pt(rect.width), Pt(rect.height)
                section.header_distance = section.footer_distance = Pt(0)
                has_text = any(b.get("type") == 0 and any(sp["text"].strip() for ln in b["lines"] for sp in ln["spans"])
                               for b in page.get_text("dict")["blocks"])
                if not editable or not has_text:
                    section.left_margin = section.right_margin = section.top_margin = Pt(0)
                    section.bottom_margin = Pt(0)
                    par = out.add_paragraph()
                    par.paragraph_format.space_after = Pt(0)
                    par.add_run().add_picture(io.BytesIO(page.get_pixmap(dpi=150).tobytes("png")), width=Pt(rect.width - 3))
                else:
                    items, (ml, mr, mt, mb), hlines, fills = _of_page_items(doc, page)
                    items = _of_columns(items, ml, rect.width - mr)
                    if number == 1:
                        sizes = Counter((it["info"]["lines"][0]["spans"][0][1]["font"], round(it["info"]["size"]))
                                        for it in items if it["kind"] == "text")
                        if sizes:
                            (font, size), _n = sizes.most_common(1)[0]
                            normal.font.name = font
                            normal.font.size = Pt(size)
                    section.left_margin, section.right_margin = Pt(ml), Pt(mr)
                    section.top_margin, section.bottom_margin = Pt(mt), Pt(min(mb, 28.0))  # taşma payı
                    cursor = mt
                    for item in items:
                        if item["kind"] == "text":
                            cursor = add_text(out, item["info"], ml, rect.width - mr, cursor)
                        elif item["kind"] == "cols":
                            cursor = add_cols(item, page, ml, cursor)
                        elif item["kind"] == "table":
                            cursor = add_table(item, page, hlines, fills, ml, cursor)
                        else:
                            cursor = add_image(item, ml, mr, rect.width, cursor)
                if progress:
                    progress(number, total)
        out.save(output_path)

    @staticmethod
    def export_xlsx(pdf_path, output_path, progress=None):
        """Excel (.xlsx): sayfalardaki tabloları (çizgili; çizgisiz sütunlu düzen de denenir) gerçek hücrelere yazar:
        sayılar sayı olarak (Türkçe biçim), yazı biçimi, hizalama, zemin rengi, birleşik hücre, sütun genişliği ve satır
        yüksekliği korunur; sayfalara bölünmüş tablo tek sayfada birleştirilir. Hiç tablo yoksa metin satır satır
        'Metin' sayfasına konur. Tablo sayısını döndürür (0 = yalnızca metin)."""
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
        clean = PDFProcessor._clean_text
        book = Workbook()
        book.remove(book.active)
        thin = Side(style="thin", color="999999")
        tables_found = 0
        last = None  # önceki tablo: devam kontrolü için
        with fitz.open(pdf_path) as doc:
            total = len(doc)
            for number, page in enumerate(doc, 1):
                hlines, fills = _of_marks(page)
                found = []
                try:
                    found = [(t, True) for t in page.find_tables().tables if t.col_count >= 2 and len(t.cells) >= 2]
                except Exception:  # noqa: BLE001
                    found = []
                if not found:  # çizgisiz sütunlu düzen: yalnızca tablo gibi görünüyorsa kabul edilir
                    try:
                        for t in page.find_tables(strategy="text").tables:
                            rows = t.extract()
                            cells = [c for r in rows for c in r]
                            filled = [c for c in cells if c and str(c).strip()]
                            if (t.col_count >= 2 and len(rows) >= 3 and len(filled) >= 0.6 * len(cells)
                                    and sum(len(str(c)) for c in filled) / max(len(filled), 1) < 30):
                                found.append((t, False))
                    except Exception:  # noqa: BLE001
                        pass
                for table, ruled in found:
                    xs, ys, spans = _of_grid(table)
                    if len(xs) < 2 or len(ys) < 2 or not spans:
                        continue
                    ncols = len(xs) - 1
                    near_bottom = last and last["page"] == number - 1 and last["bottom"] > 0.8 * last["height"]
                    if near_bottom and last["ncols"] == ncols and table.bbox[1] < 0.25 * page.rect.height:
                        sheet, base = last["sheet"], last["next_row"]
                        skip_first = True
                    else:
                        tables_found += 1
                        sheet = book.create_sheet(f"S{number}-T{tables_found}"[:31])
                        base, skip_first = 1, False
                    row_shift = 0
                    for r0, c0, r1, c1, rect in spans:
                        if skip_first and r0 == 0:
                            row_shift = max(row_shift, r1 + 1)
                    cells_text = {}
                    col_align = _of_col_aligns(page, spans, hlines)
                    for r0, c0, r1, c1, rect in spans:
                        if skip_first and r0 < row_shift:
                            continue
                        blocks = _of_cell_blocks(page, fitz.Rect(rect.x0 + 0.5, rect.y0 + 0.5, rect.x1 - 0.5, rect.y1 - 0.5), hlines)
                        text = "\n".join("".join(t if t is not None else "\n" for t, _s in _of_join(b)).strip() for b in blocks)
                        text = clean(text).strip()
                        row = base + r0 - (row_shift if skip_first else 0)
                        cell = sheet.cell(row=row, column=c0 + 1)
                        number_value = _of_number(text) if "\n" not in text else None
                        if number_value:
                            cell.value, cell.number_format = number_value
                        else:
                            cell.value = text or None
                        st = blocks[0]["lines"][0]["spans"][0][1] if blocks else None
                        if st:
                            cell.font = Font(name=st["font"], size=max(6, min(round(st["size"]), 72)), bold=st["bold"],
                                             italic=st["italic"], underline="single" if st["underline"] else None,
                                             strike=st["strike"] or None,
                                             color=_of_hex(st["color"]) if st["color"] & 0xFFFFFF else None)
                        align = _of_align(blocks[0]["bbox"], rect.x0, rect.x1) if blocks else "left"
                        align = col_align.get(c0) or align
                        if number_value:
                            align = "right" if align == "left" else align
                        cell.alignment = Alignment(horizontal={"justify": "left"}.get(align, align), vertical="center",
                                                   wrap_text=True)
                        fill = _of_cell_fill(rect, fills)
                        if fill:
                            cell.fill = PatternFill("solid", fgColor="%02X%02X%02X" % fill)
                        if ruled:
                            for rr in range(r0, r1 + 1):
                                for cc in range(c0, c1 + 1):
                                    sheet.cell(row=base + rr - (row_shift if skip_first else 0), column=cc + 1).border = Border(
                                        left=thin, right=thin, top=thin, bottom=thin)
                        if r1 > r0 or c1 > c0:
                            sheet.merge_cells(start_row=row, start_column=c0 + 1,
                                              end_row=base + r1 - (row_shift if skip_first else 0), end_column=c1 + 1)
                        cells_text[(r0, c0)] = text
                    for j in range(ncols):
                        letter = get_column_letter(j + 1)
                        width = max((xs[j + 1] - xs[j]) / 5.2, 4)
                        current = sheet.column_dimensions[letter].width or 0
                        sheet.column_dimensions[letter].width = min(max(current, width), 80)
                    for i in range(len(ys) - 1):
                        if skip_first and i < row_shift:
                            continue
                        idx = base + i - (row_shift if skip_first else 0)
                        sheet.row_dimensions[idx].height = min(max(ys[i + 1] - ys[i], 10), 409)
                    rows_written = len(ys) - 1 - (row_shift if skip_first else 0)
                    last = {"sheet": sheet, "ncols": ncols, "page": number, "bottom": table.bbox[3],
                            "height": page.rect.height, "next_row": base + rows_written}
                if progress:
                    progress(number, total)
            if not tables_found:
                sheet = book.create_sheet("Metin")
                sheet.append(["Sayfa", "Metin"])
                for number, page in enumerate(doc, 1):
                    for line in page.get_text("text", sort=True).splitlines():
                        if line.strip():
                            sheet.append([number, clean(line)])
                sheet.column_dimensions["B"].width = 100
        book.save(output_path)
        return tables_found

    @staticmethod
    def export_pptx(pdf_path, output_path, editable=True, progress=None):
        """PowerPoint (.pptx): her sayfa bir slayt. editable=True yazıları (yazı tipi, boyut, renk, hizalama dahil)
        kendi konumlarında düzenlenebilir metin kutusu yapar; yazı çıkarılmış sayfa görünümü (çizgiler, tablo
        kenarlıkları, şekiller, resimler) arka plan resmi olur. editable=False sayfayı tek resim olarak koyar."""
        import io
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
        from pptx.util import Emu, Pt
        prs = Presentation()
        alignments = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.LEFT}
        with fitz.open(pdf_path) as doc:
            first = doc[0].rect
            prs.slide_width = Emu(12192000)
            prs.slide_height = Emu(int(12192000 * first.height / first.width))
            blank = prs.slide_layouts[6]
            total = len(doc)
            for number, page in enumerate(doc, 1):
                slide = prs.slides.add_slide(blank)
                rect = page.rect
                emu = prs.slide_width / rect.width  # nokta başına EMU
                font_scale = 960.0 / rect.width     # slayt 960 pt genişliğinde
                hlines, _fills = _of_marks(page)
                blocks = []
                flags = fitz.TEXT_PRESERVE_WHITESPACE | fitz.TEXT_MEDIABOX_CLIP
                try:
                    tables = [t for t in page.find_tables().tables if t.row_count >= 1 and t.col_count >= 2 and len(t.cells) >= 2]
                except Exception:  # noqa: BLE001
                    tables = []
                tbboxes = [fitz.Rect(t.bbox) for t in tables]
                for block in page.get_text("dict", flags=flags)["blocks"]:
                    if block.get("type") == 0:
                        center = fitz.Point((block["bbox"][0] + block["bbox"][2]) / 2, (block["bbox"][1] + block["bbox"][3]) / 2)
                        if not any(center in tb for tb in tbboxes):
                            blocks.extend(_of_block_infos(block, hlines))
                for table in tables:  # tablo: her hücre ayrı metin kutusu
                    _xs, _ys, spans = _of_grid(table)
                    col_align = _of_col_aligns(page, spans, hlines)
                    for _r0, _c0, _r1, _c1, crect in spans:
                        for info in _of_cell_blocks(page, fitz.Rect(crect.x0 + 0.5, crect.y0 + 0.5, crect.x1 - 0.5, crect.y1 - 0.5), hlines):
                            info["cell"] = crect
                            info["align"] = col_align.get(_c0)
                            blocks.append(info)
                text = PDFProcessor._clean_text(page.get_text("text"))
                background = None
                if editable and blocks:
                    try:
                        with fitz.open() as tmp:
                            tmp.insert_pdf(doc, from_page=number - 1, to_page=number - 1)
                            tpage = tmp[0]
                            for info in blocks:
                                tpage.add_redact_annot(fitz.Rect(info["bbox"]), fill=False)
                            tpage.apply_redactions(images=fitz.PDF_REDACT_IMAGE_NONE, graphics=fitz.PDF_REDACT_LINE_ART_NONE)
                            background = tpage.get_pixmap(dpi=170).tobytes("png")
                    except Exception:  # noqa: BLE001 - yazı çıkarılamazsa slayt resim olarak konur
                        background = None
                if background is None:
                    slide.shapes.add_picture(io.BytesIO(page.get_pixmap(dpi=150).tobytes("png")), 0, 0,
                                             prs.slide_width, prs.slide_height)
                else:
                    slide.shapes.add_picture(io.BytesIO(background), 0, 0, prs.slide_width, int(rect.height * emu))
                    for info in blocks:
                        x0, y0, x1, y1 = info["bbox"]
                        cell = info.get("cell")
                        align = _of_align(info["bbox"], cell.x0 if cell else 0, cell.x1 if cell else rect.width,
                                          len(info["lines"]), info["justified"])
                        if info.get("align"):
                            align = info["align"]
                        if cell:
                            x0, x1 = cell.x0 + 1, cell.x1 - 1
                        line_h = max(info["lines"][0]["bbox"][3] - info["lines"][0]["bbox"][1], 1.0)
                        pitch = max(info["pitch"], info["size"] * 1.05)
                        lead = max(pitch - line_h, 0.0)
                        slack = 0 if cell else 0.08 * (x1 - x0) + 4
                        bx = {"left": x0, "justify": x0, "center": x0 - slack / 2, "right": x0 - slack}[align]
                        box = slide.shapes.add_textbox(Emu(int(bx * emu)), Emu(int((y0 - lead * 0.8) * emu)),
                                                       Emu(int((x1 - x0 + slack) * emu)), Emu(int((y1 - y0 + lead) * emu)))
                        frame = box.text_frame
                        frame.word_wrap = True
                        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
                        frame.vertical_anchor = MSO_ANCHOR.TOP
                        for li, line in enumerate(info["lines"]):
                            paragraph = frame.paragraphs[0] if li == 0 else frame.add_paragraph()
                            paragraph.alignment = alignments[align]
                            paragraph.space_before = Pt(0)
                            paragraph.space_after = Pt(0)
                            paragraph.line_spacing = Pt(pitch * font_scale)
                            for chunk, st in line["spans"]:
                                run = paragraph.add_run()
                                run.text = chunk
                                run.font.name = st["font"]
                                run.font.size = Pt(max(4, min(round(st["size"] * font_scale * 2) / 2, 200)))
                                run.font.bold = st["bold"]
                                run.font.italic = st["italic"]
                                if st["underline"]:
                                    run.font.underline = True
                                run.font.color.rgb = RGBColor((st["color"] >> 16) & 255, (st["color"] >> 8) & 255, st["color"] & 255)
                slide.notes_slide.notes_text_frame.text = text
                if progress:
                    progress(number, total)
        prs.save(output_path)

    @staticmethod
    def make_booklet(pdf_path, output_path):
        """Kitapçık dizilimi: sayfalar 4'ün katına tamamlanır, iki yana yan yana yerleştirilir (çift yüz basıp ortadan katlanır)."""
        with fitz.open(pdf_path) as src, fitz.open() as out:
            count = len(src)
            if count == 0:
                raise ValueError("Belgede sayfa yok.")
            padded = (count + 3) // 4 * 4
            width = max(page.rect.width for page in src)
            height = max(page.rect.height for page in src)
            sheets = []
            for k in range(padded // 4):
                sheets.append((padded - 1 - 2 * k, 2 * k))        # ön yüz
                sheets.append((2 * k + 1, padded - 2 - 2 * k))    # arka yüz
            for left, right in sheets:
                page = out.new_page(width=width * 2, height=height)
                for number, rect in ((left, fitz.Rect(0, 0, width, height)), (right, fitz.Rect(width, 0, width * 2, height))):
                    if number < count:
                        page.show_pdf_page(rect, src, number)
            out.save(output_path, garbage=3, deflate=True)

    # --- Temizlik, gizlilik ve sayfa düzeni araçları (Stirling PDF incelemesinden) ---
    @staticmethod
    def sanitize(doc):
        """JavaScript, gömülü/ekli dosyalar, küçük resimler, belge bilgileri ve XML meta verisini siler; bağlantı ve form alanlarına dokunmaz."""
        doc.scrub(attached_files=True, clean_pages=False, embedded_files=True, hidden_text=False, javascript=True,
                  metadata=True, redactions=False, redact_images=0, remove_links=False, reset_fields=False,
                  reset_responses=False, thumbnails=True, xml_metadata=True)

    @staticmethod
    def flatten(doc):
        """Açıklamaları ve form alanlarını sayfa içeriğine gömer (artık düzenlenemez)."""
        doc.bake(annots=True, widgets=True)

    @staticmethod
    def find_blank_pages(doc):
        """Metni, görseli, çizimi, açıklaması ve form alanı olmayan sayfaların 0 tabanlı indeksleri."""
        blank = []
        for index, page in enumerate(doc):
            if page.get_text().strip() or page.get_images() or page.get_drawings():
                continue
            if any(True for _ in page.annots()) or any(True for _ in page.widgets()):
                continue
            blank.append(index)
        return blank

    @staticmethod
    def remove_blank_pages(doc):
        blank = PDFProcessor.find_blank_pages(doc)
        if len(blank) == len(doc):
            raise ValueError("Tüm sayfalar boş görünüyor; hiçbiri silinmedi.")
        for index in reversed(blank):
            doc.delete_page(index)
        return len(blank)

    @staticmethod
    def extract_images(doc, folder):
        """Belgedeki görselleri (her biri bir kez) klasöre özgün biçiminde yazar; sayısını döndürür."""
        os.makedirs(folder, exist_ok=True)
        seen, count = set(), 0
        for number, page in enumerate(doc, 1):
            for image in PDFProcessor._real_images(page):
                xref = image[0]
                if xref in seen:
                    continue
                seen.add(xref)
                info = doc.extract_image(xref)
                if not info:
                    continue
                with open(os.path.join(folder, f"sayfa{number}_{xref}.{info.get('ext', 'png')}"), "wb") as handle:
                    handle.write(info["image"])
                count += 1
        return count

    @staticmethod
    def _real_images(page):
        """Sayfadaki görseller; silinen görselin yerine konan 1x1 şeffaf yer tutucu ve tek noktalık görseller sayılmaz."""
        return [image for image in page.get_images(full=True) if image[2] > 1 or image[3] > 1]

    @staticmethod
    def count_images(doc):
        return len({image[0] for page in doc for image in PDFProcessor._real_images(page)})  # tekil görseller

    @staticmethod
    def remove_images(doc):
        count = 0
        for page in doc:
            for image in PDFProcessor._real_images(page):
                page.delete_image(image[0])
                count += 1
        return count

    @staticmethod
    def save_with_permissions(doc, path, allow_print, allow_copy, allow_edit):
        """Açılırken parola istemeyen ama yazdırma/kopyalama/düzenleme izinleri kısıtlı kopya yazar (AES-256, rastgele sahip parolası)."""
        perm = fitz.PDF_PERM_ACCESSIBILITY
        if allow_print:
            perm |= fitz.PDF_PERM_PRINT | fitz.PDF_PERM_PRINT_HQ
        if allow_copy:
            perm |= fitz.PDF_PERM_COPY
        if allow_edit:
            perm |= fitz.PDF_PERM_MODIFY | fitz.PDF_PERM_ANNOTATE | fitz.PDF_PERM_FORM | fitz.PDF_PERM_ASSEMBLE
        doc.save(path, encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw=secrets.token_urlsafe(16), user_pw="",
                 permissions=perm, garbage=3, deflate=True)

    @staticmethod
    def stamp_pages(doc, text, size, position, page_indices, color=(0.75, 0.1, 0.1)):
        """Çerçeveli metin damgası. position: 1-9 (sayısal tuş takımı düzeni: 7 sol üst, 5 orta, 3 sağ alt)."""
        if position not in range(1, 10):
            raise ValueError("Konum 1-9 arasında olmalı (7 sol üst, 9 sağ üst, 5 orta, 1 sol alt, 3 sağ alt).")
        col = (position - 1) % 3  # 0 sol, 1 orta, 2 sağ
        row = 2 - (position - 1) // 3  # 0 üst, 1 orta, 2 alt
        width = text_width_tr(text, size)
        pad = size * 0.4
        box_w, box_h, margin = width + 2 * pad, size * 1.6, 36.0
        for index in page_indices:
            page = doc[index]
            view = page.rect
            x = margin if col == 0 else (view.width - box_w) / 2 if col == 1 else view.width - margin - box_w
            y = margin if row == 0 else (view.height - box_h) / 2 if row == 1 else view.height - margin - box_h
            x, y = max(0.0, x), max(0.0, y)
            box = fitz.Rect(x, y, x + box_w, y + box_h)
            page.draw_quad(box.quad * page.derotation_matrix, color=color, width=1.5)
            start = fitz.Point(x + pad, y + box_h * 0.72) * page.derotation_matrix
            insert_text_tr(page, start, text, fontsize=size, color=color, angle=page.rotation)

    @staticmethod
    def n_up(pdf_path, output_path, per_sheet):
        """Bir yaprağa 2, 4, 6 veya 9 sayfa yerleştirir."""
        layouts = {2: (2, 1), 4: (2, 2), 6: (3, 2), 9: (3, 3)}
        if per_sheet not in layouts:
            raise ValueError("Yaprak başına sayfa sayısı 2, 4, 6 veya 9 olmalı.")
        cols, rows = layouts[per_sheet]
        with fitz.open(pdf_path) as src, fitz.open() as out:
            if len(src) == 0:
                raise ValueError("Belgede sayfa yok.")
            width = max(page.rect.width for page in src)
            height = max(page.rect.height for page in src)
            sheet_w, sheet_h = (max(width, height), min(width, height)) if cols > rows else (width, height)
            cell_w, cell_h, gap = sheet_w / cols, sheet_h / rows, 4.0
            for start in range(0, len(src), per_sheet):
                sheet = out.new_page(width=sheet_w, height=sheet_h)
                for slot, number in enumerate(range(start, min(start + per_sheet, len(src)))):
                    r, c = divmod(slot, cols)
                    rect = fitz.Rect(c * cell_w + gap, r * cell_h + gap, (c + 1) * cell_w - gap, (r + 1) * cell_h - gap)
                    sheet.show_pdf_page(rect, src, number)
            out.save(output_path, garbage=3, deflate=True)

    @staticmethod
    def single_long_page(pdf_path, output_path):
        """Tüm sayfaları alt alta dizerek tek uzun sayfa yapar (PDF sayfa boyutu sınırı 14400 punto)."""
        with fitz.open(pdf_path) as src, fitz.open() as out:
            if len(src) == 0:
                raise ValueError("Belgede sayfa yok.")
            width = max(page.rect.width for page in src)
            total = sum(page.rect.height for page in src)
            if total > 14400:
                raise ValueError(f"Sayfaların toplam yüksekliği {total:.0f} punto; PDF sınırı 14400. Önce sayfa aralığı seçin.")
            sheet = out.new_page(width=width, height=total)
            y = 0.0
            for number, page in enumerate(src):
                w, h = page.rect.width, page.rect.height
                sheet.show_pdf_page(fitz.Rect((width - w) / 2, y, (width + w) / 2, y + h), src, number)
                y += h
            out.save(output_path, garbage=3, deflate=True)

    @staticmethod
    def count_matches(doc, query):
        return sum(len(page.search_for(query)) for page in doc)

    @staticmethod
    def redact_text(doc, query):
        """Metni kalıcı olarak siler ve yerine siyah kutu koyar (altındaki metin içerikten çıkarılır); eşleşme sayısını döndürür."""
        if not query:
            raise ValueError("Karartılacak metin boş.")
        count = 0
        for page in doc:
            hits = page.search_for(query)
            for rect in hits:
                page.add_redact_annot(rect, fill=(0, 0, 0))
                count += 1
            if hits:
                page.apply_redactions()
        return count

    @staticmethod
    def background_color(page, clip):
        """Seçili alanın hemen çevresindeki halkanın ortanca rengi (0-1 aralığında RGB). Sayfa zemini beyaz değilse silinen
        yerin çevresiyle uyumlu görünmesi için kullanılır; halka alınamazsa beyaz döner."""
        try:
            view = (fitz.Rect(clip) * page.rotation_matrix).normalize()
            pad = 6
            outer = fitz.Rect(view.x0 - pad, view.y0 - pad, view.x1 + pad, view.y1 + pad) & page.rect
            if outer.is_empty:
                return (1, 1, 1)
            pix = page.get_pixmap(clip=outer, dpi=72, alpha=False, colorspace=fitz.csRGB, annots=False)
            sx, sy = outer.width / pix.width, outer.height / pix.height
            n, samples = pix.n, pix.samples
            channels = ([], [], [])
            for row in range(pix.height):
                y = outer.y0 + (row + 0.5) * sy
                inside_rows = view.y0 <= y <= view.y1
                for col in range(pix.width):
                    x = outer.x0 + (col + 0.5) * sx
                    if inside_rows and view.x0 <= x <= view.x1:
                        continue
                    base = (row * pix.width + col) * n
                    for k in range(3):
                        channels[k].append(samples[base + k])
            if len(channels[0]) < 4:
                return (1, 1, 1)
            return tuple(sorted(c)[len(c) // 2] / 255 for c in channels)
        except Exception:  # noqa: BLE001
            return (1, 1, 1)

    @staticmethod
    def erase_area(page, clip):
        """Silgi: seçili alandaki yazıyı ve (taranmış sayfa gibi) resim piksellerini siler; yerini alanın çevresindeki
        zemin rengiyle doldurur. Çizimlere (arka plan kutuları, çizgiler) dokunmaz. Silinen sözcük ve resim sayısını döndürür."""
        words = page.get_text("words", clip=clip)
        images = [info for info in page.get_image_info() if fitz.Rect(info["bbox"]).intersects(clip)]
        if not words and not images:
            return 0
        color = PDFProcessor.background_color(page, clip)
        for w in words:
            page.add_redact_annot(fitz.Rect(w[:4]), fill=color)
        if images:
            page.add_redact_annot(fitz.Rect(clip), fill=color)
        kwargs = {"images": fitz.PDF_REDACT_IMAGE_PIXELS if images else fitz.PDF_REDACT_IMAGE_NONE}
        if hasattr(fitz, "PDF_REDACT_LINE_ART_NONE"):
            kwargs["graphics"] = fitz.PDF_REDACT_LINE_ART_NONE
        try:
            page.apply_redactions(**kwargs)
        except Exception:  # noqa: BLE001  (bazı resim türlerinde piksel silme olmayabilir)
            for annot in list(page.annots(types=[fitz.PDF_ANNOT_REDACT]) or []):
                page.delete_annot(annot)
        if images and not PDFProcessor._area_is_blank(page, clip, color):
            # Resim sıfırlanamadıysa (JBIG2/maske gibi) alanı zemin renginde bir dikdörtgenle kapat.
            cover = page.add_rect_annot(fitz.Rect(clip))
            cover.set_colors(stroke=color, fill=color)
            cover.set_border(width=0)
            cover.update()
        return len(words) + len(images)

    @staticmethod
    def _area_is_blank(page, clip, color=(1, 1, 1)):
        """Alan (gösterilen haliyle) zemin renginde boş mu? Silgi sonrası doğrulama için."""
        view = (fitz.Rect(clip) * page.rotation_matrix).normalize()
        if view.width > 8 and view.height > 8:
            view = fitz.Rect(view.x0 + 2, view.y0 + 2, view.x1 - 2, view.y1 - 2)  # kenar piksellerinin yumuşaması sayılmasın
        pix = page.get_pixmap(clip=view, dpi=30, alpha=False, colorspace=fitz.csRGB, annots=True)
        target = [round(c * 255) for c in color]
        samples = pix.samples
        return all(abs(samples[i] - target[i % 3]) <= 24 for i in range(len(samples)))

    @staticmethod
    def toc_to_text(toc):
        return "\n".join(f"{level} | {str(title).replace('|', '/')} | {max(int(page), 1)}" for level, title, page in toc)

    @staticmethod
    def parse_toc_text(text, page_count):
        """'seviye | başlık | sayfa' satırlarını doğrulayarak yer imi listesine çevirir."""
        toc, previous = [], 0
        for number, line in enumerate(text.splitlines(), 1):
            if not line.strip():
                continue
            parts = [part.strip() for part in line.split("|")]
            if len(parts) != 3:
                raise ValueError(f"Satır {number}: 'seviye | başlık | sayfa' biçiminde olmalı.")
            try:
                level, page = int(parts[0]), int(parts[2])
            except ValueError:
                raise ValueError(f"Satır {number}: seviye ve sayfa sayı olmalı.") from None
            if not parts[1]:
                raise ValueError(f"Satır {number}: başlık boş.")
            if level < 1 or (not toc and level != 1) or level > previous + 1:
                raise ValueError(f"Satır {number}: seviye 1'den başlamalı ve en fazla bir basamak artmalı.")
            if not 1 <= page <= page_count:
                raise ValueError(f"Satır {number}: sayfa 1-{page_count} arasında olmalı.")
            toc.append([level, parts[1], page])
            previous = level
        return toc

    @staticmethod
    def compare_documents(doc_a, doc_b, name_a, name_b, max_lines=40):
        """İki belgenin sayfa metinlerini karşılaştırır. (rapor satırları, farklı sayfa sayısı) döndürür."""
        import difflib
        lines = [f"Karşılaştırma: {name_a}  ↔  {name_b}", f"Sayfa sayıları: {len(doc_a)} / {len(doc_b)}", ""]
        different = 0
        for index in range(max(len(doc_a), len(doc_b))):
            a = doc_a[index].get_text().splitlines() if index < len(doc_a) else []
            b = doc_b[index].get_text().splitlines() if index < len(doc_b) else []
            if a == b:
                continue
            different += 1
            diff = [d for d in difflib.unified_diff(a, b, lineterm="", n=0) if not d.startswith(("---", "+++", "@@"))]
            removed = sum(1 for d in diff if d.startswith("-"))
            added = sum(1 for d in diff if d.startswith("+"))
            lines.append(f"Sayfa {index + 1}: {removed} satır yalnızca birinci belgede, {added} satır yalnızca ikincide")
            lines.extend(f"    {d}" for d in diff[:max_lines])
        lines.insert(2, f"Metni farklı sayfa: {different}" if different else "Metin açısından fark bulunamadı.")
        return lines, different

    @staticmethod
    def overlay_pdf(base_path, overlay_path, output_path):
        """İkinci PDF'in sayfalarını birincinin sayfalarının üstüne bindirir (sayfa sayısı azsa başa döner)."""
        with fitz.open(base_path) as base, fitz.open(overlay_path) as top:
            if len(top) == 0:
                raise ValueError("Bindirilecek PDF'te sayfa yok.")
            if top.needs_pass:
                raise ValueError("Şifreli PDF bindirilemez.")
            for index, page in enumerate(base):
                page.show_pdf_page(page.rect, top, index % len(top), overlay=True)
            base.save(output_path, garbage=3, deflate=True)

    @staticmethod
    def list_attachments(doc):
        """[(ad, bayt)] listesi (belgeye gömülü dosyalar)."""
        return [(name, doc.embfile_info(name).get("size", 0)) for name in doc.embfile_names()]

    @staticmethod
    def add_attachment(doc, path):
        base = os.path.basename(path)
        name, number = base, 1
        existing = set(doc.embfile_names())
        while name in existing:
            number += 1
            stem, ext = os.path.splitext(base)
            name = f"{stem}_{number}{ext}"
        with open(path, "rb") as handle:
            doc.embfile_add(name, handle.read())
        return name

    @staticmethod
    def extract_attachment(doc, name, output_path):
        with open(output_path, "wb") as handle:
            handle.write(doc.embfile_get(name))

    @staticmethod
    def remove_attachment(doc, name):
        doc.embfile_del(name)

    @staticmethod
    def detect_page_fix(page, min_chars=20):
        """Sayfa metni yan/ters duruyorsa okunur hâle getirmek için eklenecek saat yönü dönüş (0/90/180/270)."""
        counts = {0: 0, 90: 0, 180: 0, 270: 0}
        theta = math.radians(page.rotation)
        for block in page.get_text("dict").get("blocks", []):
            for line in block.get("lines", []):
                chars = sum(len(span.get("text", "")) for span in line.get("spans", []))
                dx, dy = line.get("dir", (1.0, 0.0))
                ex = dx * math.cos(theta) - dy * math.sin(theta)  # sayfa dönüşüyle görünür yön
                ey = dx * math.sin(theta) + dy * math.cos(theta)
                if abs(ex) >= abs(ey):
                    fix = 0 if ex > 0 else 180
                else:
                    fix = 90 if ey < 0 else 270
                counts[fix] += chars
        total = sum(counts.values())
        if total < min_chars:
            return 0
        best = max(counts, key=counts.get)
        return best if counts[best] / total >= 0.6 else 0

    @staticmethod
    def auto_rotate(doc, dry_run=False):
        """Metin yönüne göre sayfaları düzeltir; düzeltilen (veya düzeltilecek) sayfa indekslerini döndürür."""
        fixes = {}
        for index, page in enumerate(doc):
            fix = PDFProcessor.detect_page_fix(page)
            if fix:
                fixes[index] = fix
        if not dry_run:
            for index, fix in fixes.items():
                doc[index].set_rotation((doc[index].rotation + fix) % 360)
        return fixes

    @staticmethod
    def merge_pdfs(pdf_list, output_path):
        merger = fitz.open()
        try:
            for pdf in pdf_list:
                with fitz.open(pdf) as m_pdf:
                    if m_pdf.needs_pass:
                        raise ValueError(f"Şifreli PDF birleştirilemez: {os.path.basename(pdf)}")
                    merger.insert_pdf(m_pdf)
            merger.save(output_path, garbage=3, deflate=True)  # kullanılmayan nesneleri at, akışları sıkıştır
        finally:
            merger.close()


    # Resimden PDF: sayfa boyutu seçenekleri (genişlik, yükseklik pt); None = resmin kendi boyutu.
    IMAGE_PAGE_SIZES = {"A4": (595.28, 841.89), "Letter": (612.0, 792.0), "Resim boyutunda": None}
    IMAGE_EXTENSIONS = (".png", ".jpg", ".jpeg", ".bmp", ".gif", ".tif", ".tiff", ".webp")

    @staticmethod
    def _image_frames(path):
        """Resmi (EXIF yönü uygulanmış, saydamlığı beyaza basılmış) RGB/gri PNG bayt dizileri olarak verir; çok
        sayfalı TIFF/GIF'te her kare ayrı gelir. JPEG'de yön değişmiyorsa özgün baytlar korunur (yeniden sıkıştırma yok)."""
        from PIL import Image, ImageOps, ImageSequence
        import io
        with Image.open(path) as im:
            fmt = (im.format or "").upper()
            for frame in ImageSequence.Iterator(im):
                frame = ImageOps.exif_transpose(frame)
                if fmt == "JPEG" and frame.mode in ("RGB", "L") and im.getexif().get(0x0112, 1) == 1:
                    with open(path, "rb") as handle:
                        yield handle.read()
                    return
                if frame.mode in ("RGBA", "LA", "P"):
                    rgba = frame.convert("RGBA")
                    base = Image.new("RGB", rgba.size, "white")
                    base.paste(rgba, mask=rgba.getchannel("A"))
                    frame = base
                elif frame.mode not in ("RGB", "L"):
                    frame = frame.convert("RGB")
                buf = io.BytesIO()
                frame.save(buf, "PNG")
                yield buf.getvalue()

    @staticmethod
    def images_to_pdf(image_paths, output_path, page_size="A4", margin=18.0, progress=None):
        """Resimleri verilen sırayla, her biri ayrı sayfada olacak biçimde PDF'e çevirir; sayfa sayısını döndürür.
        page_size: IMAGE_PAGE_SIZES anahtarı. Sabit boyutta sayfa yönü resme göre seçilir, resim oranı korunarak ortalanır."""
        if not image_paths:
            raise ValueError("Çevrilecek resim seçilmedi.")
        if page_size not in PDFProcessor.IMAGE_PAGE_SIZES:
            raise ValueError(f"Bilinmeyen sayfa boyutu: {page_size}")
        fixed = PDFProcessor.IMAGE_PAGE_SIZES[page_size]
        out = fitz.open()
        try:
            for number, path in enumerate(image_paths, 1):
                try:
                    frames = list(PDFProcessor._image_frames(path))
                except Exception as exc:  # noqa: BLE001
                    raise ValueError(f"Resim okunamadı: {os.path.basename(path)} ({exc})") from exc
                for data in frames:
                    with fitz.open(stream=data) as probe:
                        width, height = probe[0].rect.width, probe[0].rect.height
                    if fixed:
                        pw, ph = (fixed[1], fixed[0]) if width > height else fixed
                        page = out.new_page(width=pw, height=ph)
                        box = fitz.Rect(margin, margin, pw - margin, ph - margin)
                        page.insert_image(box, stream=data, keep_proportion=True)
                    else:
                        page = out.new_page(width=width, height=height)
                        page.insert_image(page.rect, stream=data)
                if progress:
                    progress(number, len(image_paths))
            count = out.page_count
            out.save(output_path, garbage=3, deflate=True)
            return count
        finally:
            out.close()


# --- E-kitap ve belge biçimleri -------------------------------------------------------------------
# EPUB, MOBI, AZW3, IBA ve RTF önce okunabilir (akışlı) bir belgeye çevrilir, sonra EPUB ile aynı yoldan
# sayfalı PDF'e dönüşür. EPUB'u MuPDF doğrudan açar; MOBI/AZW3 için `mobi` (KindleUnpack), RTF ve IBA için
# aşağıdaki okuyucular kullanılır.
BOOK_FORMATS = {".epub": "EPUB", ".mobi": "MOBI", ".azw3": "AZW3", ".azw": "AZW3", ".iba": "IBA", ".rtf": "RTF",
                ".fb2": "FB2", ".xps": "XPS", ".oxps": "XPS", ".cbz": "CBZ", ".svg": "SVG"}
# MuPDF'in doğrudan açtığı biçimler; XPS, CBZ ve SVG sabit sayfalıdır (yeniden akıtılmaz), FB2 akışlıdır.
MUPDF_NATIVE_EXTENSIONS = (".epub", ".fb2", ".xps", ".oxps", ".cbz", ".svg")
BOOK_EXTENSIONS = tuple(BOOK_FORMATS)
DOC_EXTENSIONS = (".pdf",) + BOOK_EXTENSIONS
BOOK_PATTERNS = {"EPUB": "*.epub", "MOBI": "*.mobi", "AZW3": "*.azw3 *.azw", "IBA": "*.iba", "RTF": "*.rtf",
                 "FB2": "*.fb2", "XPS": "*.xps *.oxps", "CBZ": "*.cbz", "SVG": "*.svg"}
_ALL_BOOK_PATTERN = " ".join(BOOK_PATTERNS.values())
OPEN_FILETYPES = [("PDF ve e-kitaplar", "*.pdf " + _ALL_BOOK_PATTERN), ("PDF Dosyaları", "*.pdf")] + [
    (f"{name} dosyaları", pattern) for name, pattern in BOOK_PATTERNS.items()]
BOOK_FILETYPES = [("Tüm e-kitaplar", _ALL_BOOK_PATTERN)] + [
    (f"{name} dosyaları", pattern) for name, pattern in BOOK_PATTERNS.items()]


def is_openable_document(path):
    return bool(path) and path.lower().endswith(DOC_EXTENSIONS)


def is_book_file(path):
    """PDF olmayan, sayfalı PDF'e çevrilerek açılan e-kitap/belge dosyası mı?"""
    return bool(path) and path.lower().endswith(BOOK_EXTENSIONS)


def book_format_name(path):
    return BOOK_FORMATS.get(os.path.splitext(path or "")[1].lower(), "E-kitap")


MAX_BOOK_BYTES = 400 * 1024 * 1024  # bozuk/kötü niyetli dosyada bellek taşmasını önler

_RTF_CHARSET_CP = {0: 1252, 1: 1252, 77: 10000, 128: 932, 129: 949, 130: 1361, 134: 936, 136: 950, 161: 1253,
                   162: 1254, 163: 1258, 177: 1255, 178: 1256, 186: 1257, 204: 1251, 222: 874, 238: 1250,
                   254: 437, 255: 850}
_RTF_TOKEN = re.compile(rb"\\([a-zA-Z]{1,32})(-?\d{1,10})?[ ]?|\\'([0-9a-fA-F]{2})|\\([^a-zA-Z'])|([{}])|([^\\{}]+)", re.S)
_RTF_SKIP_DESTS = frozenset((
    "colortbl", "stylesheet", "header", "headerl", "headerr", "headerf", "footer", "footerl", "footerr", "footerf",
    "footnote", "fldinst", "object", "private", "themedata", "colorschememapping", "latentstyles", "datastore",
    "xmlnstbl", "listtable", "listoverridetable", "rsidtbl", "generator", "revtbl", "filetbl", "pgptbl",
    "operator", "userprops", "bkmkstart", "bkmkend", "pnseclvl", "listtext", "pntext"))
_RTF_CHARS = {"emdash": "\u2014", "endash": "\u2013", "bullet": "\u2022", "lquote": "\u2018", "rquote": "\u2019",
              "ldblquote": "\u201c", "rdblquote": "\u201d", "emspace": "\u2003", "enspace": "\u2002",
              "qmspace": "\u2005", "tab": "\t"}
_RTF_SYMBOLS = {b"~": "\u00a0", b"-": "\u00ad", b"_": "\u2011", b"{": "{", b"}": "}", b"\\": "\\"}
_PICT_KINDS = {"pngblip": ".png", "jpegblip": ".jpg"}


def _codec_for(codepage):
    if codepage == 10000:
        return "mac_roman"
    try:
        name = f"cp{int(codepage)}"
        "".encode(name)
        return name
    except (LookupError, TypeError, ValueError):
        return "cp1252"


class _RtfConverter:
    """RTF'yi HTML'e çevirir: paragraf, kalın/italik/altı çizili, hizalama, tablo, resim (PNG/JPEG), Unicode ve
    Windows kod sayfaları (Türkçe dahil). Üstbilgi, stil tabloları, nesneler ve dipnotlar atlanır."""

    def __init__(self, data):
        self.data = data
        self.cp = 1252
        self.cp_explicit = False
        self.fonts = {}
        self.font_def = None
        self.st = {"sink": None, "b": False, "i": False, "u": False, "s": False, "sup": False, "sub": False,
                   "font": None, "align": "", "uc": 1, "fonttbl": False, "info": False, "pict": None,
                   "pict_owner": False}
        self.stack = []
        self.body, self.images, self.meta = [], [], {}
        self.runs, self.pending = [], bytearray()
        self.uskip = 0
        self.in_row, self.cell, self.row, self.rows = False, [], [], []
        self.blank_run = 0

    # -- yardımcılar ------------------------------------------------------------------------------
    def _codec(self):
        font = self.st["font"]
        return _codec_for(self.fonts.get(font if font is not None else self.deff_font, self.cp))

    deff_font = None

    def _flush_bytes(self):
        if not self.pending:
            return
        raw, self.pending = bytes(self.pending), bytearray()
        if self.st["sink"] is None or isinstance(self.st["sink"], tuple):
            self._add(raw.decode(self._codec(), errors="replace"))

    def _add(self, text):
        if not text:
            return
        sink, st = self.st["sink"], self.st
        if sink is None:
            fmt = (st["b"], st["i"], st["u"], st["s"], st["sup"], st["sub"])
            if self.runs and self.runs[-1][1] == fmt:
                self.runs[-1] = (self.runs[-1][0] + text, fmt)
            else:
                self.runs.append((text, fmt))
        elif isinstance(sink, tuple):
            self.meta[sink[1]] = self.meta.get(sink[1], "") + text

    def _runs_html(self):
        parts = []
        for text, fmt in self.runs:
            text = text.encode("utf-16-le", "surrogatepass").decode("utf-16-le", "replace")
            text = html.escape(text, quote=False).replace("\n", "<br/>").replace("\t", "&#160;" * 4)
            for flag, tag in zip(fmt, ("b", "i", "u", "s", "sup", "sub")):
                if flag:
                    text = f"<{tag}>{text}</{tag}>"
            parts.append(text)
        return "".join(parts)

    def _has_text(self):
        return any(text.strip() for text, _ in self.runs)

    def _flush_table(self):
        if not self.rows:
            return
        rows = "".join("<tr>" + "".join(f"<td>{cell}</td>" for cell in row) + "</tr>" for row in self.rows if row)
        self.body.append('<table border="1" cellspacing="0" cellpadding="4" style="border-collapse:collapse">'
                         + rows + "</table>")
        self.rows = []

    def _end_paragraph(self):
        text = self._runs_html() if self._has_text() else ""
        self.runs = []
        if self.in_row:
            if text:
                self.cell.append(text)
            return
        self._flush_table()
        if not text:
            self.blank_run += 1
            if self.blank_run <= 2:
                self.body.append("<p>&#160;</p>")
            return
        self.blank_run = 0
        align = {"c": "center", "r": "right", "j": "justify"}.get(self.st["align"])
        self.body.append(f'<p style="text-align:{align}">{text}</p>' if align else f"<p>{text}</p>")

    def _end_cell(self):
        if self._has_text():
            self._end_paragraph_in_row()
        self.in_row = True
        self.row.append("<br/>".join(self.cell) or "&#160;")
        self.cell = []

    def _end_paragraph_in_row(self):
        was = self.in_row
        self.in_row = True
        self._end_paragraph()
        self.in_row = was

    def _end_row(self):
        if self._has_text():
            self._end_paragraph_in_row()
        if self.cell:
            self.row.append("<br/>".join(self.cell))
            self.cell = []
        self.rows.append(self.row)
        self.row = []
        self.in_row = False

    def _finish_pict(self, pict):
        kind = pict.get("kind")
        if kind not in _PICT_KINDS.values():
            return
        raw = bytes(pict["raw"])
        if not raw:
            hexa = bytes(pict["hex"])
            hexa = hexa[: len(hexa) // 2 * 2]
            try:
                raw = bytes.fromhex(hexa.decode("ascii"))
            except ValueError:
                return
        if not raw:
            return
        name = f"resim{len(self.images) + 1}{kind}"
        self.images.append((name, raw))
        size = f' width="{min(max(round(pict["w"] / 15), 8), 2000)}"' if pict.get("w") else ""
        if self._has_text() and not self.in_row:
            self._end_paragraph()
        tag = f'<img src="{name}"{size}/>'
        if self.in_row:
            self.cell.append(tag)
        else:
            self._flush_table()
            self.body.append(f'<p style="text-align:center">{tag}</p>')

    # -- belirteç işleme --------------------------------------------------------------------------
    def _control(self, word, param):
        st = self.st
        if st.get("star"):  # {\*\hedef …}: tanınmayan hedef atlanır (resim kapsayıcısı shppict hariç)
            st["star"] = False
            if word != "shppict":
                st["sink"] = "skip"
            return
        if st["fonttbl"]:
            if word == "f":
                self.font_def = param
            elif word == "fcharset" and self.font_def is not None and _RTF_CHARSET_CP.get(param):
                self.fonts[self.font_def] = _RTF_CHARSET_CP[param]
            return
        if word == "pict" and st["sink"] is None:
            st.update(sink="pict", pict={"kind": None, "hex": bytearray(), "raw": bytearray(), "w": None},
                      pict_owner=True)
            return
        if st["sink"] == "pict":
            if word in _PICT_KINDS:
                st["pict"]["kind"] = _PICT_KINDS[word]
            elif word == "picwgoal" and param:
                st["pict"]["w"] = param
            return
        if word == "fonttbl":
            st.update(fonttbl=True, sink="skip")
            return
        if word == "info":
            st.update(info=True, sink="skip")
            return
        if st["info"] and word in ("title", "author", "subject"):
            st["sink"] = ("meta", word)
            return
        if word in _RTF_SKIP_DESTS:
            st["sink"] = "skip"
            return
        if word == "ansicpg" and param:
            self.cp, self.cp_explicit = param, True
        elif word in ("ansi", "mac", "pc", "pca") and not self.cp_explicit:
            self.cp = {"ansi": 1252, "mac": 10000, "pc": 437, "pca": 850}[word]
        elif word == "deff":
            self.deff_font = param
        elif word == "uc":
            st["uc"] = 1 if param is None else param
        elif word == "f":
            st["font"] = param
        elif word in ("b", "i", "strike", "super", "sub"):
            st[{"strike": "s", "super": "sup"}.get(word, word)] = param != 0
        elif word == "ul":
            st["u"] = param != 0
        elif word in ("ulnone", "ul0"):
            st["u"] = False
        elif word == "nosupersub":
            st["sup"] = st["sub"] = False
        elif word == "plain":
            st.update(b=False, i=False, u=False, s=False, sup=False, sub=False)
        elif word == "pard":
            st["align"] = ""
        elif word in ("ql", "qr", "qc", "qj"):
            st["align"] = word[1]
        elif word in ("intbl", "trowd"):
            self.in_row = True
        elif st["sink"] is not None and not (word == "u" and isinstance(st["sink"], tuple)):
            return
        elif word == "u" and param is not None:
            code = param + 65536 if param < 0 else param
            self._add(chr(code))
            self.uskip = st["uc"]
        elif word in ("par", "sect"):
            self._end_paragraph()
        elif word == "line":
            self._add("\n")
        elif word == "cell":
            self._end_cell()
        elif word == "row":
            self._end_row()
        elif word in _RTF_CHARS:
            self._add(_RTF_CHARS[word])

    def _symbol(self, sym):
        if sym == b"*":
            if self.st["sink"] is None or self.st["sink"] == "pict":
                self.st["star"] = True
        elif sym in (b"\n", b"\r"):
            if self.st["sink"] is None:
                self._end_paragraph()
        elif sym in _RTF_SYMBOLS:
            self._add(_RTF_SYMBOLS[sym])

    def _text(self, raw):
        raw = raw.replace(b"\r", b"").replace(b"\n", b"")
        if not raw:
            return
        if self.uskip:
            cut = min(self.uskip, len(raw))
            raw, self.uskip = raw[cut:], self.uskip - cut
            if not raw:
                return
        if self.st["sink"] == "pict":
            self.st["pict"]["hex"] += bytes(c for c in raw if c in b"0123456789abcdefABCDEF")
        elif self.st["sink"] is None or isinstance(self.st["sink"], tuple):
            self._add(raw.decode(self._codec() if max(raw) > 127 else "ascii", errors="replace"))

    def run(self):
        data, pos, size = self.data, 0, len(self.data)
        if not data.lstrip()[:5] == b"{\\rtf":
            raise ValueError("Bu dosya geçerli bir RTF belgesi gibi görünmüyor.")
        while pos < size:
            match = _RTF_TOKEN.match(data, pos)
            if not match:
                pos += 1
                continue
            pos = match.end()
            word, param, hexv, sym, brace, text = match.groups()
            if hexv is not None:
                if self.uskip:
                    self.uskip -= 1
                else:
                    self.pending.append(int(hexv, 16))
                continue
            self._flush_bytes()
            if brace == b"{":
                if len(self.stack) < 400:
                    self.stack.append(self.st)
                    self.st = dict(self.st, pict_owner=False, star=False)
                self.uskip = 0
            elif brace == b"}":
                if self.stack:
                    closed, self.st = self.st, self.stack.pop()
                    if closed["pict_owner"]:
                        self._finish_pict(closed["pict"])
                self.uskip = 0
            elif word is not None:
                name = word.decode("ascii")
                value = int(param) if param is not None else None
                if name != "u":
                    self.uskip = 0
                if name == "bin" and value:
                    if self.st["sink"] == "pict":
                        self.st["pict"]["raw"] += data[pos:pos + value]
                    pos += value
                    continue
                self._control(name, value)
            elif sym is not None:
                self._symbol(sym)
            elif text is not None:
                self._text(text)
        self._flush_bytes()
        if self._has_text():
            self._end_paragraph()
        if self.cell or self.row:
            self._end_row()
        self._flush_table()
        return "\n".join(self.body), self.images, self.meta


def rtf_to_html(data):
    """RTF baytlarından (HTML gövdesi, [(resim adı, bayt)], {'title':…, 'author':…}) üretir."""
    return _RtfConverter(data).run()


_IBA_IMAGE_EXT = (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff")


def _xml_local(tag):
    return tag.rsplit("}", 1)[-1] if isinstance(tag, str) else ""


def _iba_text(element):
    parts = [element.text or ""]
    for child in element:
        name = _xml_local(child.tag)
        if name in ("br", "lnbr", "crbr", "pgbr", "intr"):
            parts.append("\n")
        elif name == "tab":
            parts.append("\t")
        elif name not in ("ghost-text", "attachment", "ghost-p"):
            parts.append(_iba_text(child))
        parts.append(child.tail or "")
    return "".join(parts)


def iba_to_html(path):
    """iBooks Author (.iba) belgesinden metin ve resimleri okur. IBA, içinde index.xml(.gz) bulunan bir paket
    (zip) dosyasıdır; metin <p> paragraflarından, resimler Data/ klasöründeki dosya başvurularından alınır.
    Widget, 3B ve etkileşimli öğeler okunamaz; düzen (sütun, kutu konumu) kitap akışına çevrilir."""
    import gzip
    import io
    import urllib.parse
    import xml.etree.ElementTree as ET
    import zipfile
    if not zipfile.is_zipfile(path):
        raise ValueError("Bu dosya geçerli bir IBA (iBooks Author) belgesi gibi görünmüyor.")
    images, paragraphs, seen = [], [], set()
    with zipfile.ZipFile(path) as pack:
        names = pack.namelist()
        lookup = {}
        for name in names:
            lookup[name.lower()] = name
            if "/" in name:
                lookup.setdefault(name.split("/", 1)[1].lower(), name)
        candidates = [n for n in names if n.lower().rsplit("/", 1)[-1] in ("index.xml.gz", "index.xml")]
        if not candidates:
            raise ValueError("IBA içinde index.xml bulunamadı; dosya bu sürümle uyumlu olmayabilir.")
        index = min(candidates, key=lambda n: (n.count("/"), not n.lower().endswith(".gz")))
        with pack.open(index) as source:
            raw = source.read(MAX_BOOK_BYTES + 1)
        if len(raw) <= MAX_BOOK_BYTES and (index.lower().endswith(".gz") or raw[:2] == b"\x1f\x8b"):
            with gzip.GzipFile(fileobj=io.BytesIO(raw)) as unzipped:
                raw = unzipped.read(MAX_BOOK_BYTES + 1)
        if len(raw) > MAX_BOOK_BYTES:
            raise ValueError("IBA içeriği çok büyük.")
        try:
            root = ET.fromstring(raw)
        except ET.ParseError as exc:
            raise ValueError(f"IBA içeriği okunamadı: {exc}") from exc

        def image_refs(element):
            for node in element.iter():
                for key, value in node.attrib.items():
                    if _xml_local(key) == "path" and value.lower().endswith(_IBA_IMAGE_EXT):
                        member = lookup.get(value.lower()) or lookup.get(urllib.parse.unquote(value).lower())
                        if member and member not in seen:
                            seen.add(member)
                            yield member

        def add_image(member):
            info = pack.getinfo(member)
            if info.file_size > 50 * 1024 * 1024:
                return
            ext = os.path.splitext(member)[1].lower()
            name = f"resim{len(images) + 1}{ext}"
            images.append((name, pack.read(member)))
            paragraphs.append(f'<p style="text-align:center"><img src="{name}"/></p>')

        stack = [root]
        while stack:
            element = stack.pop()
            name = _xml_local(element.tag)
            if name == "p":
                text = _iba_text(element).strip()
                if text:
                    style = " ".join(v for k, v in element.attrib.items() if _xml_local(k) == "style").lower()
                    tag = "p"
                    if "title" in style:
                        tag = "h1"
                    elif "heading" in style or "chapter" in style or "head" in style:
                        digits = re.findall(r"\d", style)
                        tag = f"h{min(int(digits[0]) + 1, 6)}" if digits else "h2"
                    body = html.escape(text, quote=False).replace("\n", "<br/>").replace("\t", "&#160;" * 4)
                    paragraphs.append(f"<{tag}>{body}</{tag}>")
                for member in image_refs(element):
                    add_image(member)
                continue
            for key, value in element.attrib.items():
                if _xml_local(key) == "path" and value.lower().endswith(_IBA_IMAGE_EXT):
                    for member in image_refs(element):
                        add_image(member)
                    break
            stack.extend(reversed(list(element)))
    if not paragraphs:
        raise ValueError("IBA dosyasında okunabilir metin veya resim bulunamadı (desteklenmeyen sürüm olabilir).")
    return "\n".join(paragraphs), images, {}


_IMAGE_MIME = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".gif": "image/gif",
               ".bmp": "image/bmp", ".tif": "image/tiff", ".tiff": "image/tiff"}


def _html_document(body, images, title):
    """Gövde ve resimlerden tek parça HTML üretir; resimler data: adresi olarak gömülür (MuPDF dosya yolundan
    açılan HTML'de göreli resimleri yüklemez, EPUB'da yükler)."""
    import base64
    for name, data in images:
        mime = _IMAGE_MIME.get(os.path.splitext(name)[1].lower(), "image/png")
        body = body.replace(f'src="{name}"', f'src="data:{mime};base64,{base64.b64encode(data).decode("ascii")}"')
    head = html.escape((title or "").strip(), quote=False)
    return ('<!DOCTYPE html><html><head><meta charset="utf-8"><title>' + head + "</title>"
            "<style>body{font-family:serif;line-height:1.35} p{margin:0 0 0.6em 0} img{max-width:100%}"
            "td{vertical-align:top}</style></head><body>" + body + "</body></html>").encode("utf-8")


def _open_kindle(path, temp_dirs):
    """MOBI/AZW3/AZW'yi açar: önce KindleUnpack (`mobi` paketi) ile EPUB'a açılır, olmazsa .mobi için MuPDF denenir.
    Açılan dosyalar geçici klasörde tutulur ve temp_dirs ile çağıran tarafından silinir."""
    error = None
    try:
        from mobi.kindleunpack import unpackBook
    except ImportError:
        unpackBook = None
    if unpackBook is not None:
        try:
            work = tempfile.mkdtemp(prefix="babapdf_kitap_")
            temp_dirs.append(work)
            source = os.path.join(work, "kitap.mobi")  # ASCII ad: Türkçe/uzun yol sorunlarını önler
            shutil.copyfile(path, source)
            unpackBook(source, work, epubver="A")
            for relative in (("mobi8", "kitap.epub"), ("mobi7", "book.html"), ("kitap.001.pdf",)):
                candidate = os.path.join(work, *relative)
                if os.path.exists(candidate):
                    return fitz.open(candidate)
            raise ValueError("Kitap içeriği çıkarılamadı.")
        except Exception as exc:  # noqa: BLE001
            error = exc
    if path.lower().endswith(".mobi"):
        try:
            return fitz.open(path)
        except Exception as exc:  # noqa: BLE001
            error = error or exc
    if error and "drm" in str(error).lower():
        raise ValueError("Şifreli (DRM korumalı) Kindle kitapları açılamıyor.")
    raise ValueError(f"Kindle kitabı açılamadı (bozuk, korumalı ya da desteklenmeyen biçim olabilir):\n{error}")


def _open_book(path):
    """E-kitabı okunabilir belge olarak açar: (belge, ek üst veri, temizlik işlevi). Geçici dosyalar temizlikle silinir."""
    ext = os.path.splitext(path)[1].lower()
    temp_dirs = []

    def cleanup():
        for folder in temp_dirs:
            shutil.rmtree(folder, ignore_errors=True)

    try:
        if ext in MUPDF_NATIVE_EXTENSIONS:
            return fitz.open(path), {}, cleanup
        if ext in (".mobi", ".azw3", ".azw"):
            return _open_kindle(path, temp_dirs), {}, cleanup
        if ext == ".rtf":
            if os.path.getsize(path) > MAX_BOOK_BYTES:
                raise ValueError("RTF dosyası çok büyük.")
            with open(path, "rb") as source:
                body, images, meta = rtf_to_html(source.read())
        elif ext == ".iba":
            body, images, meta = iba_to_html(path)
        else:
            raise ValueError(f"Desteklenmeyen biçim: {ext}")
        return fitz.open(stream=_html_document(body, images, meta.get("title")), filetype="html"), meta, cleanup
    except Exception:
        cleanup()
        raise

EPUB_FONT_SIZES = {"Küçük": 10, "Orta": 12, "Büyük": 15, "Çok büyük": 19}
EPUB_PAGE_SIZES = {"Telefon": (360, 640), "Kitap": (480, 680), "A5": (420, 595), "A4": (595, 842)}
EPUB_PRODUCER = "Baba PDF EPUB"
EPUB_SOURCES = {}  # üretilen PDF'in (küçük harfli tam yol) geldiği EPUB'un yolu


def epub_options():
    """Kayıtlı EPUB okuma ayarları: (yazı puntosu, (genişlik, yükseklik), kitabın yanına kaydet)."""
    settings = load_settings()
    font = EPUB_FONT_SIZES.get(settings.get("epub_font"), EPUB_FONT_SIZES["Orta"])
    page = EPUB_PAGE_SIZES.get(settings.get("epub_page"), EPUB_PAGE_SIZES["Kitap"])
    return font, page, bool(settings.get("epub_beside", True))


def _epub_output_candidates(epub_path, beside):
    if beside:
        base = os.path.splitext(epub_path)[0]
        if not epub_path.lower().endswith(".epub"):
            base += "_" + book_format_name(epub_path).lower()  # aynı adlı EPUB/MOBI/RTF… PDF'leri çakışmasın
        yield base + ".pdf"
        counter = 2
        while True:
            yield f"{base}_epub{counter}.pdf"
            counter += 1
    else:
        key = hashlib.sha1(os.path.abspath(epub_path).lower().encode("utf-8")).hexdigest()[:12]
        folder = os.path.join(app_data_dir(), "epub")
        os.makedirs(folder, exist_ok=True)
        yield os.path.join(folder, f"{os.path.splitext(os.path.basename(epub_path))[0]}_{key}.pdf")


def _epub_stamp(font, page):
    return f"{EPUB_PRODUCER} f{font} p{page[0]}x{page[1]}"


def _read_producer(path):
    try:
        with fitz.open(path) as probe:
            return (probe.metadata or {}).get("producer", "") or ""
    except Exception:  # noqa: BLE001
        return ""


def epub_cached_pdf(epub_path):
    """Geçerli ayarlarla daha önce üretilmiş, güncel PDF varsa yolunu verir; yoksa None."""
    font, page, beside = epub_options()
    stamp = _epub_stamp(font, page)
    for out in _epub_output_candidates(epub_path, beside):
        if not os.path.exists(out):
            return None
        producer = _read_producer(out)
        if producer == stamp and os.path.getmtime(out) >= os.path.getmtime(epub_path):
            EPUB_SOURCES[os.path.normcase(os.path.abspath(out))] = epub_path
            return out
        if producer.startswith(EPUB_PRODUCER):
            return None  # kendi eski çıktımız (başka ayarla): yeniden üretilip üzerine yazılacak
    return None


def epub_to_pdf(epub_path):
    """E-kitabı (EPUB, MOBI, AZW3, IBA, RTF) sayfalı PDF'e çevirir. Ayarlar (yazı/sayfa boyutu) değişmediyse önceki
    çıktıyı kullanır; kendi ürettiği eski çıktının üzerine yazar, başkasının dosyasına dokunmaz."""
    cached = epub_cached_pdf(epub_path)
    if cached:
        return cached
    font, page, beside = epub_options()
    stamp = _epub_stamp(font, page)
    out = None
    for candidate in _epub_output_candidates(epub_path, beside):
        if not os.path.exists(candidate) or _read_producer(candidate).startswith(EPUB_PRODUCER):
            out = candidate
            break
    kind = book_format_name(epub_path)
    try:
        book, extra_meta, cleanup = _open_book(epub_path)
    except ValueError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"{kind} dosyası açılamadı (bozuk ya da desteklenmeyen biçim olabilir):\n{exc}") from exc
    try:
        with book:
            if book.needs_pass:
                raise ValueError(f"Şifreli (DRM korumalı) {kind} dosyaları açılamıyor.")
            book_meta = dict(book.metadata or {})
            book_meta.update({key: value for key, value in extra_meta.items() if value})
            if book.is_pdf:  # bazı Kindle dosyaları (baskı kopyası) PDF olarak çıkar
                toc = book.get_toc()
                pdf_bytes = book.tobytes()
            elif not book.is_reflowable and os.path.splitext(epub_path)[1].lower() in (".xps", ".oxps", ".cbz", ".svg"):
                if book.page_count == 0:  # sabit sayfalı biçimler olduğu gibi PDF'e dönüştürülür
                    raise ValueError(f"{kind} içinde okunabilir sayfa bulunamadı.")
                toc = book.get_toc()
                pdf_bytes = book.convert_to_pdf()
            else:
                if not book.is_reflowable:
                    raise ValueError(f"Bu dosya geçerli bir {kind} kitabı gibi görünmüyor.")
                book.layout(width=page[0], height=page[1], fontsize=font)
                if book.page_count == 0:
                    raise ValueError(f"{kind} içinde okunabilir sayfa bulunamadı.")
                toc = book.get_toc()
                pdf_bytes = book.convert_to_pdf()
    finally:
        cleanup()
    with fitz.open("pdf", pdf_bytes) as pdf:
        if toc:
            try:
                pdf.set_toc(toc)
            except Exception:  # noqa: BLE001
                pass
        meta = dict(pdf.metadata or {})
        meta["producer"] = stamp
        for key in ("title", "author", "subject", "keywords"):
            if book_meta.get(key):
                meta[key] = book_meta[key]
        pdf.set_metadata(meta)
        tmp = out + ".tmp"
        pdf.save(tmp, garbage=3, deflate=True)
    os.replace(tmp, out)
    EPUB_SOURCES[os.path.normcase(os.path.abspath(out))] = epub_path
    return out


def resolve_document_path(path):
    """Açılacak dosya e-kitap (EPUB, MOBI, AZW3, IBA, RTF) ise PDF karşılığını üretip yolunu döndürür; değilse yolu aynen verir."""
    if path and is_book_file(path):
        return epub_to_pdf(path)
    return path


def epub_source_of(pdf_path):
    return EPUB_SOURCES.get(os.path.normcase(os.path.abspath(pdf_path))) if pdf_path else None


PAGE_MODES = {
    "normal": ("Normal (beyaz)", None),
    "sepya": ("Sepya", (0x3B2A14, 0xF1E4C3)),
    "koyu": ("Koyu sayfa", (0xE6E1D6, 0x1C1C22)),
}


def tint_pixmap(pix, mode):
    """Yalnızca ekrandaki görüntüye uygulanır; belgenin kendisi ve dışa aktarımlar değişmez."""
    colors = PAGE_MODES.get(mode, PAGE_MODES["normal"])[1]
    if colors:
        pix.tint_with(*colors)
    return pix


MARKUP_KINDS = {"highlight": "add_highlight_annot", "underline": "add_underline_annot",
                "strikeout": "add_strikeout_annot", "squiggly": "add_squiggly_annot"}


def add_markup(page, rects, kind="highlight"):
    """Dikdörtgenlere vurgulama / altı çizili / üstü çizili / dalgalı alt çizgi açıklaması ekler; eklenen sayısını verir."""
    method = getattr(page, MARKUP_KINDS[kind])
    count = 0
    for rect in rects:
        method(quads=[fitz.Rect(rect)]).update()
        count += 1
    return count


def add_ink(page, strokes, color=(0.85, 0.1, 0.1), width=2.0):
    """Serbest kalem çizgisi: strokes, sayfanın PyMuPDF (döndürülmemiş) uzayında nokta listeleri listesidir."""
    strokes = [[(float(x), float(y)) for x, y in stroke] for stroke in strokes if len(stroke) >= 2]
    if not strokes:
        return None
    annot = page.add_ink_annot(strokes)
    annot.set_colors(stroke=color)
    annot.set_border(width=width)
    annot.update()
    return annot


COVER_W, COVER_H = 96, 128


def library_info(path):
    """Kitaplık kartı için başlık, yazar ve kapak (ilk sayfa küçük resmi). Sonuç uygulama klasöründe önbelleğe alınır."""
    info = {"title": os.path.splitext(os.path.basename(path))[0], "author": "", "cover": None}
    try:
        stat = os.stat(path)
        key = hashlib.sha1(f"{os.path.abspath(path).lower()}|{stat.st_mtime_ns}|{stat.st_size}".encode("utf-8")).hexdigest()[:16]
        folder = os.path.join(app_data_dir(), "covers")
        os.makedirs(folder, exist_ok=True)
        png, meta_file = os.path.join(folder, key + ".png"), os.path.join(folder, key + ".json")
        if os.path.exists(meta_file):
            with open(meta_file, "r", encoding="utf-8") as source:
                cached = json.load(source)
            info.update(title=cached.get("title") or info["title"], author=cached.get("author", ""))
            info["cover"] = png if os.path.exists(png) else None
            return info
        source = path
        if is_book_file(path) and not path.lower().endswith(".epub"):
            source = epub_cached_pdf(path)  # MOBI/AZW3/IBA/RTF: kapak ve başlık dönüştürülmüş PDF'ten alınır
            if not source:
                return info
        with fitz.open(source) as doc:
            if doc.needs_pass or doc.page_count == 0:
                return info
            meta = doc.metadata or {}
            info["title"] = (meta.get("title") or "").strip() or info["title"]
            info["author"] = (meta.get("author") or "").strip()
            page = doc[0]
            scale = min(COVER_W / page.rect.width, COVER_H / page.rect.height)
            page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False).save(png)
            info["cover"] = png
        with open(meta_file, "w", encoding="utf-8") as target:
            json.dump({"title": info["title"], "author": info["author"]}, target, ensure_ascii=False)
    except Exception:  # noqa: BLE001
        pass
    return info


def reading_progress(path):
    """Kaydedilmiş konumdan (yüzde, sayfa, toplam) döndürür; konum yoksa None."""
    saved = recall_position(path)
    if not saved or saved[1] <= 0:
        return None
    index, count = saved
    index = min(max(index, 0), count - 1)
    return round((index + 1) / count * 100), index + 1, count


def load_positions():
    data = load_settings().get("positions")
    return data if isinstance(data, dict) else {}


def remember_position(key, page_index, page_count):
    if not key:
        return
    positions = load_positions()
    norm = os.path.normcase(os.path.abspath(key))
    positions.pop(norm, None)
    positions[norm] = [int(page_index), int(page_count)]
    while len(positions) > 200:
        positions.pop(next(iter(positions)))
    save_setting("positions", positions)


def recall_position(key):
    entry = load_positions().get(os.path.normcase(os.path.abspath(key))) if key else None
    try:
        return int(entry[0]), int(entry[1])
    except (TypeError, ValueError, IndexError):
        return None


def generate_edited_filepath(original_path):
    """
    Orijinal dosya yoluna göre aynı klasör içerisinde D1, D2, D3...
    şeklinde benzersiz düzenlenmiş dosya adı üretir.
    """
    folder, filename = os.path.split(original_path)
    base, ext = os.path.splitext(filename)

    counter = 1
    while True:
        new_filename = f"{base}_D{counter}{ext}"
        new_path = os.path.join(folder, new_filename)
        if not os.path.exists(new_path):
            return new_path
        counter += 1


SAFE_LINK_SCHEMES = ("http://", "https://", "mailto:")


def find_link_at(page, point, links=None):
    """Sayfanın GÖRÜNEN (döndürülmüş) koordinatındaki noktanın altındaki bağlantıyı verir; yoksa None.
    Bağlantı dikdörtgenleri döndürülmemiş sayfa uzayında olduğundan nokta önce o uzaya çevrilir."""
    api_point = fitz.Point(point) * page.derotation_matrix
    for link in (page.get_links() if links is None else links):
        rect = link.get("from")
        if rect is not None and fitz.Rect(rect).contains(api_point):
            return link
    return None


def describe_link(link, page_count=0):
    """Bağlantının hedefini kısa Türkçe metinle anlatır (ipucu balonu için); desteklenmeyen türde None."""
    kind = link.get("kind")
    if kind == fitz.LINK_GOTO and link.get("page", -1) >= 0:
        return f"Sayfa {link['page'] + 1}"
    if kind == fitz.LINK_URI:
        uri = str(link.get("uri", ""))
        return uri if len(uri) <= 90 else uri[:90] + "…"
    if kind == fitz.LINK_NAMED:
        names = {"NextPage": "Sonraki sayfa", "PrevPage": "Önceki sayfa", "FirstPage": "İlk sayfa", "LastPage": "Son sayfa"}
        return names.get(link.get("name"))
    return None


class PDFViewer(TkinterDnD.Tk if DND_AVAILABLE else tk.Tk):
    """Ana PDF Görüntüleyici ve Düzenleme Penceresi."""

    TOOLS = (
        ("gez", "✋", "Gez: sayfayı sürükleyerek kaydır"),
        ("kopyala", "🔤", "Metin kopyala: sürükleyerek seç, panoya kopyalanır"),
        ("vurgula", "🖍", "Vurgula: sürükleyerek metni vurgula"),
        ("altcizgi", "🅄", "Altını çiz: sürükleyerek metnin altını çiz"),
        ("ustucizgi", "🅂", "Üstünü çiz: sürükleyerek metnin üstünü çiz"),
        ("dalga", "〰", "Dalgalı alt çizgi: sürükleyerek metni işaretle"),
        ("kalem", "✏", "Kalem: sürükleyerek serbest çizim yap"),
        ("not", "📝", "Not: tıkladığın yere not ekle"),
        ("metin", "🅰", "Yazı: tıkladığın yere yazı ekle"),
        ("cizgi", "📏", "Çizgi: sürükleyerek çizgi çiz"),
        ("gorsel", "🖼", "Görsel/imza: tıkladığın yere resim yerleştir"),
        ("silgi", "⌫", "Silgi: sürükleyerek seçtiğin yazıyı sil"),
        ("gizle", "⬛", "Gizle: seçili alanı kalıcı olarak sil (redaksiyon)"),
        ("kirp", "✂", "Kırp: sürükleyerek sayfada tutulacak alanı seç"),
    )

    def __init__(self, pdf_path, welcome=False):
        super().__init__()
        self._is_edited = False
        self.pdf_path = pdf_path
        self.epub_source = epub_source_of(pdf_path)
        self.is_welcome = False
        self.doc = fitz.open(pdf_path)
        self._password = None

        if self.doc.needs_pass:
            pw = simpledialog.askstring("Şifre", "Bu PDF şifreli. Şifreyi girin:", show="*", parent=self)
            if pw is None or not self.doc.authenticate(pw):
                self.doc.close()
                self.destroy()
                raise ValueError("PDF şifresi yanlış veya girilmedi.")
            self._password = pw

        self.zoom = 1.0
        self.page_images = {}      # sayfa indeksi -> PhotoImage (yalnızca görünür sayfalar)
        self.thumb_objects = []
        self.thumb_cards = []
        self._link_cache = {}      # sayfa indeksi -> get_links() sonucu (render_all_pages'te sıfırlanır)
        self._link_tip = None
        self._nav_back, self._nav_fwd = [], []   # gezinti geçmişi: kaydırma oranları
        self.is_toc_open = False
        self._toc_cache = None
        self._toc_targets = {}
        self.page_positions = []
        self.page_x = []
        self.total_height = 1
        self.search_hits = {}      # sayfa indeksi -> [Rect]
        self._hit_index = 0
        self._search_none = False
        self.widget_hits = {}      # sayfa indeksi -> [Rect]
        self._zoom_timer = None
        self._visible_job = None
        self._bg_state = None
        self._undo_stack = []
        self._redo_stack = []
        self.drag_data = {"widget": None, "index": None}
        self.selected_pages = set()
        self.page_orientation = None  # None: sayfalar dosyadaki gibi, "dikey" veya "yatay"
        self.tool_mode = "gez"
        self._te = None  # yazı düzenleme oturumu (sayfada önizleme)
        self._drag_start = None
        self._sel_id = None
        self._save_mode = None        # None: sorulmadı, "overwrite" veya "copy"
        self._toast_job = None
        self._redact_scrub_asked = False
        self._welcome_frame = None
        self.dark_mode_var = tk.BooleanVar(value=bool(load_settings().get("dark_mode")))
        saved_mode = load_settings().get("page_mode")
        self.page_mode_var = tk.StringVar(value=saved_mode if saved_mode in PAGE_MODES else "normal")
        self.theme = THEMES["dark" if self.dark_mode_var.get() else "light"]

        self.title(self._window_title())
        self.geometry("1180x780")
        self.minsize(900, 560)
        self.protocol("WM_DELETE_WINDOW", self._on_close)
        self.report_callback_exception = self._report_callback_exception

        # --- UYGULAMA İKONUNU PROGRAMA GÖMME ---
        try:
            ico_file = resource_path("app_icon.ico")
            if os.path.exists(ico_file):
                self.iconbitmap(default=ico_file)  # başlık çubuğu ve görev çubuğu (Windows)
            icon_file = resource_path("app_icon.png")
            if os.path.exists(icon_file):
                self._icon_img = tk.PhotoImage(file=icon_file)
                self.iconphoto(True, self._icon_img)
        except Exception as e:
            print(f"İkon yüklenirken hata oluştu: {e}")

        self.fullscreen_var = tk.BooleanVar(value=False)
        self.resizable_var = tk.BooleanVar(value=True)
        self.portrait_var = tk.BooleanVar(value=False)
        self.landscape_var = tk.BooleanVar(value=False)

        self._build_menus()
        self._build_toolbars()
        self._build_statusbar()
        self._build_body()
        self._bind_shortcuts()

        self.is_edited = self._is_edited  # düğme durumunu eşitle
        self.apply_theme()
        self.after(30000, self._autosave_recovery)
        self.after(100, self.init_view)
        if welcome:
            self.after(150, self.show_welcome)

    # ------------------------------------------------------------------
    # Pencere iskeleti: menüler, araç çubukları, durum çubuğu, gövde
    # ------------------------------------------------------------------
    def _build_menus(self):
        menubar = tk.Menu(self)

        def submenu(parent, label, items):
            menu = tk.Menu(parent, tearoff=0)
            for item in items:
                if item is None:
                    menu.add_separator()
                else:
                    text, command, *accel = item
                    menu.add_command(label=text, command=command, accelerator=accel[0] if accel else None)
            parent.add_cascade(label=label, menu=menu)
            return menu

        file_menu = tk.Menu(menubar, tearoff=0)
        file_menu.add_command(label="PDF Aç…", command=self.action_open_new, accelerator="Ctrl+O")
        file_menu.add_command(label="E-Kitap Aç…", command=self.action_open_book)
        self.recent_menu = tk.Menu(file_menu, tearoff=0)
        file_menu.add_cascade(label="Son Açılanlar", menu=self.recent_menu)
        self._refresh_recent_menu()
        file_menu.add_command(label="Ana Sayfa", command=self.show_welcome)
        file_menu.add_separator()
        file_menu.add_command(label="Kaydet", command=self.save_pdf, accelerator="Ctrl+S")
        file_menu.add_command(label="Farklı Kaydet…", command=self.save_pdf_as, accelerator="Ctrl+Shift+S")
        file_menu.add_separator()
        file_menu.add_command(label="Yazdır", command=self.action_print, accelerator="Ctrl+P")
        submenu(file_menu, "Office Biçiminde Kaydet", [
            ("Word Belgesi Olarak Kaydet (.docx)…", lambda: self.action_save_office("docx")),
            ("Excel Çalışma Kitabı Olarak Kaydet (.xlsx)…", lambda: self.action_save_office("xlsx")),
            ("PowerPoint Sunusu Olarak Kaydet (.pptx)…", lambda: self.action_save_office("pptx")),
        ])
        submenu(file_menu, "Dışa Aktar", [
            ("Metne Dönüştür (.txt)", self.action_extract_text),
            ("OCR ile Metin Çıkar (.txt)", self.action_ocr),
            ("Toplu Metin Çıkarma", self.action_batch_extract),
            None,
            ("Sayfaları Resme Çevir (.png)", self.action_convert_images),
            ("Resimlerden PDF Oluştur…", self.action_images_to_pdf),
            ("Sayfa Aralığını PDF Olarak Kaydet", self.action_extract_pages),
            ("Seçilen Sayfaları PDF Olarak Kaydet", self.action_export_selected_pages),
            ("Sayfa İletişim Paftası Oluştur", self.action_contact_sheet),
            None,
            ("Kitapçık Düzeni Oluştur…", self.action_booklet),
            ("Çok Sayfalı Düzen (yaprağa 2/4/6/9 sayfa)…", self.action_nup),
            ("Tüm Sayfaları Tek Uzun Sayfaya Birleştir…", self.action_single_page),
        ])
        file_menu.add_separator()
        file_menu.add_command(label="Çıkış", command=self._on_close)
        menubar.add_cascade(label="Dosya", menu=file_menu)

        edit_menu = tk.Menu(menubar, tearoff=0)
        edit_menu.add_command(label="Geri Al", command=self.undo, accelerator="Ctrl+Z")
        edit_menu.add_command(label="Yinele", command=self.redo, accelerator="Ctrl+Y")
        edit_menu.add_separator()
        # Ara + Sonraki + Önceki tek takımdır (araç çubuğundaki 🔍 ▲ ▼ ile aynı işlevler); gezinme yalnızca sonuç varken açık.
        edit_menu.add_command(label="Komut Ara…", command=self.action_command_palette, accelerator="Ctrl+K")
        edit_menu.add_command(label="PDF İçinde Ara…", command=self.action_search, accelerator="Ctrl+F")
        edit_menu.add_command(label="Sonraki Eşleşme", command=lambda: self.search_step(1), accelerator="F3", state="disabled")
        edit_menu.add_command(label="Önceki Eşleşme", command=lambda: self.search_step(-1), accelerator="Shift+F3", state="disabled")
        self._search_menu = (edit_menu, ("Sonraki Eşleşme", "Önceki Eşleşme"))
        edit_menu.add_separator()
        submenu(edit_menu, "Sayfalar", [
            ("Sayfayı Döndür", self.rotate_current_page),
            ("Yan/Ters Sayfaları Otomatik Döndür", self.action_auto_rotate),
            ("Sayfayı Sil", self.delete_current_page),
            ("Seçilen Sayfaları Sil", self.action_delete_selected_pages),
            ("Boş Sayfaları Sil", self.action_remove_blank),
            None,
            ("PDF'ye Sayfa Ekle…", self.action_append_pages),
            ("PDF'yi Sayfalara Böl…", self.action_split),
            ("Birden Fazla PDF'i Birleştir…", self.action_merge),
            None,
            ("Sayfa Boyutu, Kenar Boşluğu ve Arka Plan…", self.action_page_layout),
            ("Sayfa Kenarlarını Kırp…", self.action_crop),
            ("Alan Seçerek Kırp (sürükle)", self.action_crop_tool),
            ("Kırpmayı Kaldır…", self.action_crop_reset),
            ("Boş Sayfa Ekle…", self.action_insert_blank),
            ("Sayfa Aralığını Sil…", self.action_delete_page_range),
            ("Üst Bilgi, Alt Bilgi ve Sayfa Numarası…", self.action_headers_footers),
        ])
        submenu(edit_menu, "Filigran, Damga ve Bindirme", [
            ("Filigran Ekle…", self.action_watermark),
            ("Damga Ekle…", self.action_stamp),
            ("Başka Bir PDF'i Üstüne Bindir (Overlay)…", self.action_overlay),
        ])
        submenu(edit_menu, "Açıklamalar", [
            ("Metni Ara ve Vurgula", lambda: self.action_annotation("highlight")),
            ("Metni Ara ve Altını Çiz", lambda: self.action_annotation("underline")),
            ("Metni Ara ve Üstünü Çiz", lambda: self.action_annotation("strikeout")),
            ("Metni Ara ve Dalgalı Altını Çiz", lambda: self.action_annotation("squiggly")),
            None,
            ("Araç: Altını Çiz (sürükle)", lambda: self._select_tool("altcizgi")),
            ("Araç: Üstünü Çiz (sürükle)", lambda: self._select_tool("ustucizgi")),
            ("Araç: Dalgalı Alt Çizgi (sürükle)", lambda: self._select_tool("dalga")),
            ("Araç: Serbest Kalem (sürükle)", lambda: self._select_tool("kalem")),
            ("Araç: Silgi — yazıyı sil (sürükle)", lambda: self._select_tool("silgi")),
            None,
            ("Tüm Açıklamaları Sil", self.action_remove_annotations),
            None,
            ("Açıklamaları ve Formları Sayfaya Düzleştir", self.action_flatten),
        ])
        submenu(edit_menu, "Yer İmleri", [
            ("Yer İmi Ekle", self.action_add_bookmark),
            ("Yer İmlerini Göster, Düzenle veya Sil…", self.action_edit_toc),
        ])
        submenu(edit_menu, "Formlar", [
            ("Form Alanlarını Doldur", self.action_fill_form),
            ("Form Editörü", self.action_form_editor),
        ])
        menubar.add_cascade(label="Düzenle", menu=edit_menu)

        security_menu = tk.Menu(menubar, tearoff=0)
        security_menu.add_command(label="PDF'yi Parolayla Şifrele…", command=self.action_encrypt)
        security_menu.add_command(label="PDF Şifresini Kaldır…", command=self.action_decrypt)
        security_menu.add_command(label="İzinleri Değiştir ve Kopya Kaydet…", command=self.action_permissions)
        security_menu.add_separator()
        security_menu.add_command(label="Metni Kalıcı Karart (Redact)…", command=self.action_redact)
        security_menu.add_command(label="Gizli Verileri ve Belge Bilgilerini Temizle", command=self.action_sanitize)
        security_menu.add_separator()
        security_menu.add_command(label="Görsel / İmza Resmi Yerleştir…", command=self.action_add_signature)
        menubar.add_cascade(label="Güvenlik", menu=security_menu)

        tools_menu = tk.Menu(menubar, tearoff=0)
        tools_menu.add_command(label="PDF Boyutunu Küçült…", command=self.action_compress)
        tools_menu.add_command(label="Belge Bilgileri", command=self.action_document_info)
        tools_menu.add_command(label="Metadata Düzenle…", command=self.action_metadata)
        tools_menu.add_separator()
        tools_menu.add_command(label="PDF'yi Doğrula (Sağlık Denetimi)", command=self.action_check_document)
        tools_menu.add_command(label="Bozuk Dosyayı Onar ve Kopya Kaydet…", command=self.action_repair)
        submenu(tools_menu, "Görseller", [
            ("Görselleri Dışa Aktar…", self.action_extract_images),
            ("Tüm Görselleri Sil", self.action_remove_images),
        ])
        tools_menu.add_command(label="İki PDF'i Karşılaştır…", command=self.action_compare)
        tools_menu.add_command(label="Gömülü Ekleri Yönet…", command=self.action_attachments)
        menubar.add_cascade(label="Araçlar", menu=tools_menu)

        view_menu = tk.Menu(menubar, tearoff=0)
        view_menu.add_checkbutton(label="Tam ekran", accelerator="F11",
                                  variable=self.fullscreen_var, command=self.toggle_fullscreen)
        view_menu.add_checkbutton(label="Ayarlanabilir Pencere Boyutu",
                                  variable=self.resizable_var, command=self.toggle_resizable)
        view_menu.add_checkbutton(label="Koyu tema", variable=self.dark_mode_var, command=self.toggle_dark_mode)
        view_menu.add_separator()
        mode_menu = tk.Menu(view_menu, tearoff=0)
        for key, (label, _colors) in PAGE_MODES.items():
            mode_menu.add_radiobutton(label=label, value=key, variable=self.page_mode_var, command=self.set_page_mode)
        view_menu.add_cascade(label="Sayfa Rengi", menu=mode_menu)
        view_menu.add_command(label="E-Kitap Okuma Ayarları…", command=self.action_epub_settings)
        self.epub_beside_var = tk.BooleanVar(value=epub_options()[2])
        view_menu.add_checkbutton(label="E-kitaptan üretilen PDF'i kitabın yanına kaydet", variable=self.epub_beside_var,
                                  command=lambda: save_setting("epub_beside", bool(self.epub_beside_var.get())))
        view_menu.add_separator()
        view_menu.add_command(label="İçindekiler Paneli", accelerator="F4", command=self.toggle_toc_panel)
        view_menu.add_command(label="Katmanlar…", command=self.action_layers)
        view_menu.add_command(label="Geri Git (önceki konum)", accelerator="Alt+←", command=self.nav_back)
        view_menu.add_command(label="İleri Git", accelerator="Alt+→", command=self.nav_forward)
        view_menu.add_separator()
        view_menu.add_checkbutton(label="Dikey Sayfa", variable=self.portrait_var,
                                  command=lambda: self.set_orientation("dikey"))
        view_menu.add_checkbutton(label="Yatay Sayfa", variable=self.landscape_var,
                                  command=lambda: self.set_orientation("yatay"))
        menubar.add_cascade(label="Görünüm", menu=view_menu)

        help_menu = tk.Menu(menubar, tearoff=0)
        help_menu.add_command(label="Klavye Kısayolları", command=self.action_shortcuts)
        help_menu.add_command(label="Hata Günlüğünü Aç", command=self.action_open_error_log)
        help_menu.add_command(label="Varsayılan PDF Programı Yap…", command=self.action_set_default_app)
        help_menu.add_command(label="Güncellemeleri Denetle", command=self.action_check_update)
        help_menu.add_command(label="Lisans ve Kaynak Kodu", command=self.action_license)
        help_menu.add_separator()
        help_menu.add_command(label=f"{APP_NAME} Hakkında", command=self.action_about)
        menubar.add_cascade(label="Yardım", menu=help_menu)

        self.config(menu=menubar)

    # Açılış sayfasında (belge yokken) yalnızca belge istemeyen komutlar açık kalır.
    _WELCOME_MENU_OK = {"PDF Aç…", "E-Kitap Aç…", "Son Açılanlar", "Çıkış", "Toplu Metin Çıkarma", "Komut Ara…",
                        "Birden Fazla PDF'i Birleştir…", "Görünüm", "Yardım"}
    _SEARCH_STEP_LABELS = {"Sonraki Eşleşme", "Önceki Eşleşme"}  # durumlarını arama sonucu belirler

    def _refresh_menu_state(self):
        """Belge yokken (açılış sayfası) belgeye bağlı menü öğelerini kapatır; belge açılınca yeniden açar."""
        welcome = self.is_welcome

        def walk(menu, always_on):
            any_on = False
            last = menu.index("end")
            for index in range(0 if last is None else last + 1):
                kind = menu.type(index)
                if kind in ("separator", "tearoff"):
                    continue
                label = menu.entrycget(index, "label")
                keep = always_on or label in self._WELCOME_MENU_OK
                if kind == "cascade":
                    on = True if label == "Son Açılanlar" else walk(self.nametowidget(menu.entrycget(index, "menu")), keep)  # alt öğeleri son dosyalar belirler
                    menu.entryconfigure(index, state="normal" if on else "disabled")
                else:
                    on = keep or not welcome
                    if label not in self._SEARCH_STEP_LABELS:
                        menu.entryconfigure(index, state="normal" if on else "disabled")
                    elif welcome:
                        menu.entryconfigure(index, state="disabled")
                any_on = any_on or on
            return any_on

        walk(self.nametowidget(self.cget("menu")), False)

    def _build_toolbars(self):
        # 1. satır: dosya, sayfa, geri al, yakınlaştırma, sayfa işlemleri, arama
        self.bar = ttk.Frame(self, style="Bar.TFrame")
        self.bar.pack(side="top", fill="x")
        bar = ttk.Frame(self.bar, style="Bar.TFrame")
        bar.pack(fill="x", padx=8, pady=(6, 3))

        def add_button(text, command, tip, width=3):
            button = ttk.Button(bar, text=text, width=width, command=command, style="Icon.TButton")
            button.pack(side="left", padx=1)
            self._tooltip(button, tip)
            return button

        def add_separator():
            ttk.Separator(bar, orient="vertical").pack(side="left", fill="y", padx=6, pady=2)

        add_button("🏠", self.show_welcome, "Ana sayfa")
        add_button("📂", self.action_open_new, "PDF aç (Ctrl+O)")
        add_separator()
        add_button("◀", lambda: self._step_page(-1, force=True), "Önceki sayfa (PageUp)", width=2)
        self.page_var = tk.StringVar(value="1")
        self.page_entry = ttk.Entry(bar, textvariable=self.page_var, width=4, justify="center")
        self.page_entry.pack(side="left", padx=2)
        self.page_entry.bind("<Return>", self.goto_page)
        self.total_label = ttk.Label(bar, text=f"/ {len(self.doc)}", style="Bar.TLabel")
        self.total_label.pack(side="left", padx=(2, 2))
        add_button("▶", lambda: self._step_page(1, force=True), "Sonraki sayfa (PageDown)", width=2)
        add_separator()
        add_button("↶", self.undo, "Geri al (Ctrl+Z)")
        add_button("↷", self.redo, "Yinele (Ctrl+Y)")
        add_separator()
        add_button("−", lambda: self.set_zoom(self.zoom / 1.2), "Uzaklaştır (Ctrl+-)")
        self.zoom_label = ttk.Label(bar, text="100%", width=5, anchor="center", style="Bar.TLabel")
        self.zoom_label.pack(side="left")
        add_button("+", lambda: self.set_zoom(self.zoom * 1.2), "Yakınlaştır (Ctrl++)")
        add_button("↔", self.fit_width, "Genişliğe sığdır (Ctrl+0)")
        add_separator()
        add_button("⟳", self.rotate_current_page, "Sayfayı döndür")
        add_button("🗑", self.delete_current_page, "Sayfayı sil")
        add_button("☰", self.toggle_thumbnail_panel, "Sayfa sıralama panelini aç/kapat")
        add_button("📑", self.toggle_toc_panel, "İçindekiler panelini aç/kapat (F4)")
        add_separator()
        add_button("🔍", self.action_search, "Ara (Ctrl+F)")
        # Ara + önceki/sonraki tek takımdır: gezinme düğmeleri yalnızca aktif bir arama sonucu varken kullanılabilir.
        self.search_prev_button = add_button("▲", lambda: self.search_step(-1), "Önceki eşleşme (Shift+F3)", width=2)
        self.search_next_button = add_button("▼", lambda: self.search_step(1), "Sonraki eşleşme (F3)", width=2)
        self.search_prev_button.state(["disabled"])
        self.search_next_button.state(["disabled"])
        self.search_label = ttk.Label(bar, text="", width=14, style="Bar.TLabel")
        self.search_label.pack(side="left", padx=4)

        print_button = ttk.Button(bar, text="🖨", width=3, command=self.action_print)
        print_button.pack(side="right", padx=1)
        self._tooltip(print_button, "Yazdır (Ctrl+P)")
        theme_box = ttk.Checkbutton(bar, text="🌙", variable=self.dark_mode_var, command=self.toggle_dark_mode,
                                    style="Bar.TCheckbutton")
        theme_box.pack(side="right", padx=4)
        self._tooltip(theme_box, "Koyu tema")
        # Düzenleme yapılınca görünen tek "Kaydet" düğmesi (Yazdır'ın soluna paketlenir)
        self.save_btn = ttk.Button(bar, text="💾 Kaydet", style="Save.TButton", command=self.save_pdf)
        self._tooltip(self.save_btn, "Kaydet (Ctrl+S)")

        # 2. satır: sayfa üstü araçlar
        tools = ttk.Frame(self.bar, style="Tools.TFrame")
        tools.pack(fill="x", padx=8, pady=(0, 6))
        ttk.Label(tools, text="Araçlar", style="Tools.TLabel").pack(side="left", padx=(8, 6), pady=3)
        self.tool_var = tk.StringVar(value="gez")
        for key, icon, tip in self.TOOLS:
            rb = ttk.Radiobutton(tools, text=icon, value=key, variable=self.tool_var,
                                 style="Tool.Toolbutton", command=self._on_tool_change, width=3)
            rb.pack(side="left", padx=1, pady=3)
            self._tooltip(rb, tip)
        self.tool_hint = ttk.Label(tools, text=self.TOOLS[0][2], style="Tools.TLabel")
        self.tool_hint.pack(side="left", padx=12)

    def _build_statusbar(self):
        self.statusbar = ttk.Frame(self, style="Bar.TFrame")
        self.statusbar.pack(side="bottom", fill="x")
        self.operation_progress = ttk.Progressbar(self.statusbar, orient="horizontal", length=130, mode="determinate")
        self.operation_progress.pack(side="left", padx=(10, 6), pady=4)
        self.operation_status = ttk.Label(self.statusbar, text="Hazır", style="Bar.TLabel")
        self.operation_status.pack(side="left")
        self.toast_frame = tk.Frame(self.statusbar)
        self.toast_label = tk.Label(self.toast_frame, text="", padx=10, pady=2, font=(FONT_UI, 9, "bold"))
        self.toast_label.pack(side="left")
        self.toast_link = tk.Label(self.toast_frame, text="Klasörde göster", padx=8, pady=2,
                                   font=(FONT_UI, 9, "underline"), cursor="hand2")

    def _build_body(self):
        self.paned = ttk.PanedWindow(self, orient="horizontal")
        self.paned.pack(fill="both", expand=True)

        # --- SOL PANEL: Görsel Sayfa Sıralama (Thumbnails) ---
        self.thumb_frame = ttk.Frame(self.paned, width=200)
        thumb_label = ttk.Label(self.thumb_frame, text="Sayfa Sıralama (Sürükle & Bırak)",
                                font=(FONT_UI, 9, "bold"))
        thumb_label.pack(side="top", fill="x", pady=5, padx=5)
        self.thumb_canvas = tk.Canvas(self.thumb_frame, highlightthickness=0)
        thumb_vbar = ttk.Scrollbar(self.thumb_frame, orient="vertical", command=self.thumb_canvas.yview)
        self.thumb_canvas.configure(yscrollcommand=thumb_vbar.set)
        thumb_vbar.pack(side="right", fill="y")
        self.thumb_canvas.pack(side="left", fill="both", expand=True)
        self.thumb_container = ttk.Frame(self.thumb_canvas)
        self.thumb_canvas.create_window((0, 0), window=self.thumb_container, anchor="nw")
        self.thumb_container.bind("<Configure>", lambda e: self.thumb_canvas.configure(scrollregion=self.thumb_canvas.bbox("all")))

        # --- SAĞ PANEL: Ana Görünüm Canvas ---
        self.body = ttk.Frame(self.paned)
        self.paned.add(self.body, weight=4)
        self.is_thumb_open = False  # sayfa sıralama paneli varsayılan olarak kapalı

        # --- SOL PANEL: İçindekiler (yer imleri ağacı); varsayılan kapalı ---
        self.toc_frame = ttk.Frame(self.paned, width=240)
        ttk.Label(self.toc_frame, text="İçindekiler", font=(FONT_UI, 9, "bold")).pack(side="top", fill="x", pady=5, padx=5)
        self.toc_note = ttk.Label(self.toc_frame, text="Bu belgede içindekiler (yer imi) yok.", wraplength=200)
        toc_bar = ttk.Scrollbar(self.toc_frame, orient="vertical")
        self.toc_tree = ttk.Treeview(self.toc_frame, show="tree", selectmode="browse", style="Toc.Treeview",
                                     yscrollcommand=toc_bar.set)
        toc_bar.config(command=self.toc_tree.yview)
        toc_bar.pack(side="right", fill="y")
        self.toc_tree.pack(side="left", fill="both", expand=True)
        self.toc_tree.bind("<ButtonRelease-1>", self._on_toc_click)
        self.toc_tree.bind("<Return>", self._on_toc_click)

        self.canvas = tk.Canvas(self.body, highlightthickness=0)
        self.vbar = ttk.Scrollbar(self.body, orient="vertical", command=self.canvas.yview)
        hbar = ttk.Scrollbar(self.body, orient="horizontal", command=self.canvas.xview)
        self.canvas.configure(yscrollcommand=self.on_scroll, xscrollcommand=hbar.set)
        self.vbar.pack(side="right", fill="y")
        hbar.pack(side="bottom", fill="x")
        self.canvas.pack(side="left", fill="both", expand=True)

    def _bind_shortcuts(self):
        self.bind("<Control-z>", lambda e: self.undo())
        self.bind("<Control-y>", lambda e: self.redo())
        self.bind("<Control-plus>", lambda e: self.set_zoom(self.zoom * 1.2))
        self.bind("<Control-equal>", lambda e: self.set_zoom(self.zoom * 1.2))
        self.bind("<Control-minus>", lambda e: self.set_zoom(self.zoom / 1.2))
        self.bind("<Control-o>", lambda e: self.action_open_new())
        self.bind("<Control-s>", lambda e: self.save_pdf() if self.is_edited else None)
        self.bind("<Control-S>", lambda e: self.save_pdf_as())
        self.bind("<Control-p>", lambda e: self.action_print())
        self.bind("<Control-f>", lambda e: self.action_search())
        self.bind("<Control-k>", lambda e: self.action_command_palette())
        self.bind("<F3>", lambda e: self.search_step(1))
        self.bind("<Shift-F3>", lambda e: self.search_step(-1))
        self.bind("<Control-0>", lambda e: self.fit_width())
        self.bind("<Prior>", lambda e: self._step_page(-1))
        self.bind("<Next>", lambda e: self._step_page(1))
        self.bind("<F11>", self.toggle_fullscreen)
        self.bind("<F4>", lambda e: self.toggle_toc_panel())
        self.bind("<Alt-Left>", lambda e: self.nav_back())
        self.bind("<Alt-Right>", lambda e: self.nav_forward())
        self.bind("<Escape>", self.exit_fullscreen)
        self.canvas.bind("<MouseWheel>", self.on_wheel)
        self.canvas.bind("<Control-MouseWheel>", self.on_ctrl_wheel)
        self.canvas.bind("<ButtonPress-1>", self.on_canvas_press)
        self.canvas.bind("<B1-Motion>", self.on_canvas_drag)
        self.canvas.bind("<ButtonRelease-1>", self.on_canvas_release)
        self.canvas.bind("<Motion>", self.on_canvas_motion)
        self.canvas.bind("<Leave>", lambda e: self._set_link_hover(None, e))
        if DND_AVAILABLE:
            for target in (self, self.canvas):
                target.drop_target_register(DND_FILES)
                target.dnd_bind("<<Drop>>", self.on_files_dropped)

    def _on_tool_change(self):
        self.tool_mode = self.tool_var.get()
        if self._te and self.tool_mode != "metin":
            self._te_cancel()
        cursors = {"gez": "fleur", "kopyala": "xterm", "not": "plus", "metin": "xterm", "gorsel": "plus"}
        self.canvas.configure(cursor=cursors.get(self.tool_mode, "crosshair"))
        for key, _icon, tip in self.TOOLS:
            if key == self.tool_mode:
                self.tool_hint.config(text=tip)

    def _toast(self, message, path=None, seconds=8):
        """Başarı bildirimini pencere açmadan durum çubuğunda gösterir; dosya yolu varsa 'Klasörde göster' sunar."""
        t = self.theme
        self.toast_label.config(text="✔ " + message, bg=t["mint"], fg=t["text"])
        self.toast_frame.config(bg=t["mint"])
        if path:
            self.toast_link.config(bg=t["mint"], fg=t["accent_dark"])
            self.toast_link.bind("<Button-1>", lambda _e, p=path: reveal_in_folder(p))
            self.toast_link.pack(side="left")
        else:
            self.toast_link.pack_forget()
        self.toast_frame.pack(side="right", padx=10, pady=3)
        if self._toast_job is not None:
            self.after_cancel(self._toast_job)
        self._toast_job = self.after(seconds * 1000, self.toast_frame.pack_forget)

    def _tooltip(self, widget, text):
        """Düğmenin üzerinde bekleyince küçük açıklama balonu gösterir."""
        state = {"win": None}

        def show(_event):
            if state["win"] is not None:
                return
            win = tk.Toplevel(widget)
            win.wm_overrideredirect(True)
            win.wm_geometry(f"+{widget.winfo_rootx() + 4}+{widget.winfo_rooty() + widget.winfo_height() + 4}")
            tk.Label(win, text=text, bg=self.theme["lemon"], fg=self.theme["text"], relief="solid", borderwidth=1,
                     padx=6, pady=2, font=(FONT_UI, 9)).pack()
            state["win"] = win

        def hide(_event):
            if state["win"] is not None:
                state["win"].destroy()
                state["win"] = None

        widget.bind("<Enter>", show, add="+")
        widget.bind("<Leave>", hide, add="+")
        widget.bind("<ButtonPress>", hide, add="+")

    def _report_callback_exception(self, exc_type, exc_value, exc_traceback):
        try:
            with open(error_log_path(), "a", encoding="utf-8") as log:
                log.write(f"\n[{datetime.now().isoformat(timespec='seconds')}] Tkinter callback\n")
                traceback.print_exception(exc_type, exc_value, exc_traceback, file=log)
            messagebox.showerror("Uygulama hatası", f"Beklenmeyen hata oluştu. Ayrıntılar:\n{error_log_path()}", parent=self)
        except Exception:
            pass

    # ------------------------------------------------------------------
    # Durum / yardımcılar
    # ------------------------------------------------------------------
    @property
    def is_edited(self):
        return self._is_edited

    @is_edited.setter
    def is_edited(self, value):
        self._is_edited = bool(value)
        btn = self.__dict__.get("save_btn")
        if btn is not None:
            if self._is_edited:
                btn.pack(side="right", padx=5)
            else:
                btn.pack_forget()

    def restore_recovery(self, recovery_file):
        recovered = None
        try:
            recovered = fitz.open(recovery_file)
            old_doc = self.doc
            self.doc = recovered
            old_doc.close()
            self.is_edited = True
            self.search_hits = {}
            self.render_thumbnails()
            self.render_all_pages()
            self.title(f"{os.path.basename(self.pdf_path)} - Kurtarılan çalışma - {APP_NAME}")
        except Exception as exc:
            if recovered is not None and recovered is not self.doc:
                recovered.close()  # yarım kalan kurtarma belgesini serbest bırak
            show_runtime_error(self, "Kurtarma dosyası açma", exc)

    def _refresh_recent_menu(self):
        self.recent_menu.delete(0, "end")
        recent = load_recent_files()
        if not recent:
            self.recent_menu.add_command(label="Yok", state="disabled")
            return
        for path in recent:
            self.recent_menu.add_command(label=os.path.basename(path), command=lambda p=path: self.open_pdf_path(p))
        self.recent_menu.add_separator()
        self.recent_menu.add_command(label="Listeyi Temizle", command=self.action_clear_recent)

    def action_clear_recent(self):
        if not load_recent_files():
            return
        if not messagebox.askyesno("Son Açılanlar", "Son açılanlar listesi temizlensin mi?\n(Dosyalarınız silinmez.)", parent=self):
            return
        if not clear_recent_files():
            messagebox.showerror("Hata", "Liste temizlenemedi.", parent=self)
            return
        self._refresh_recent_menu()
        if getattr(self, "_welcome_frame", None) is not None and self._welcome_frame.winfo_exists():
            self.show_welcome()

    def _autosave_recovery(self):
        try:
            if self.is_edited and self.doc and self.pdf_path:
                self.doc.save(recovery_path(self.pdf_path), garbage=4, deflate=True, clean=True)
        except Exception as exc:
            log_exception("Otomatik kurtarma", exc)
        finally:
            self.after(30000, self._autosave_recovery)

    def _clear_recovery(self):
        self._remove_quietly(recovery_path(self.pdf_path))

    def _position_key(self):
        return getattr(self, "epub_source", None) or self.pdf_path

    def _window_title(self):
        source = getattr(self, "epub_source", None)
        if source:
            meta = self.doc.metadata or {}
            label = " – ".join(part for part in (meta.get("title"), meta.get("author")) if part)
            name = os.path.basename(source)
            return f"{label} ({name}) - {APP_NAME}" if label else f"{name} - {APP_NAME}"
        return f"{os.path.basename(self.pdf_path)} - {APP_NAME}"

    def _save_position(self):
        """Okunan sayfayı kaydeder; kitap/belge aynı dosyadan yeniden açılınca kaldığı yerden devam eder."""
        try:
            if self.is_welcome or not self.page_positions:
                return
            remember_position(self._position_key(), self.get_current_page_index(), len(self.doc))
        except Exception:  # noqa: BLE001
            pass

    def _restore_position(self, start_page=None):
        """Belirtilen sayfaya (yoksa son okunan sayfaya) gider. Sayfa sayısı değiştiyse oranla eşler."""
        if start_page is None:
            saved = recall_position(self._position_key())
            if not saved:
                return
            index, count = saved
            if count != len(self.doc) and count > 0:
                index = round(index / count * len(self.doc))
        else:
            index = start_page
        index = min(max(index, 0), len(self.doc) - 1)
        if index <= 0:
            return

        def go():
            if not self.page_positions:
                return
            self.page_var.set(str(index + 1))
            self.goto_page()

        self.after(120, go)

    def _confirm_discard(self):
        """Kaydedilmemiş değişiklik varsa kullanıcıya sorar. Devam edilebilirse True döner."""
        if not self.is_edited:
            return True
        return messagebox.askyesno(
            "Kaydedilmemiş değişiklikler",
            "Bu PDF'de kaydedilmemiş değişiklikler var. Kaybolacak, devam edilsin mi?",
            icon="warning", parent=self,
        )

    def _on_close(self):
        if self._bg_state is not None and not messagebox.askyesno(
                "İşlem sürüyor", "Arka planda bir işlem çalışıyor. Yine de çıkılsın mı?", parent=self):
            return
        if self.is_edited:
            answer = messagebox.askyesnocancel(
                "Kaydedilmemiş değişiklikler",
                "Bu PDF'de kaydedilmemiş değişiklikler var.\n\nEvet: düzenlenmiş kopya olarak kaydet ve çık\n"
                "Hayır: kaydetmeden çık\nİptal: geri dön",
                icon="warning", parent=self)
            if answer is None:
                return
            if answer:
                self.save_pdf()
                if self.is_edited:  # kaydetme başarısız oldu; veri kaybolmasın
                    return
        self._save_position()
        self._clear_recovery()
        try:
            self.doc.close()
        except Exception:
            pass
        self.destroy()

    def _is_same_file(self, path):
        try:
            return os.path.exists(path) and os.path.samefile(path, self.pdf_path)
        except OSError:
            return False

    def _working_pdf_path(self):
        """
        Dış araçların (metin çıkarma, OCR, yazdırma) kullanacağı PDF yolu.
        Bellekte düzenleme varsa veya PDF şifreliyse, güncel hali geçici bir dosyaya yazılır.
        Döndürülen (yol, geçici_mi).
        """
        if not self.is_edited and not self.doc.is_encrypted and os.path.exists(self.pdf_path):
            return self.pdf_path, False
        tmp = os.path.join(tempfile.gettempdir(), f"baba_pdf_{uuid.uuid4().hex}.pdf")
        self.doc.save(tmp, encryption=fitz.PDF_ENCRYPT_NONE)
        return tmp, True

    def _export_from_working(self, func, success_text, on_ok=None):
        """func(kaynak_pdf_yolu)'nu güncel (düzenlenmiş) belgeyle çalıştırır; hata ve temizliği yönetir."""
        source, is_tmp = None, False
        try:
            source, is_tmp = self._working_pdf_path()
            func(source)
            if on_ok:
                on_ok()
            elif success_text:
                self._toast(success_text)
        except Exception as exc:
            log_exception("Dışa aktarma", exc)
            messagebox.showerror("Hata", str(exc), parent=self)
        finally:
            if is_tmp and source:
                self._remove_quietly(source)

    @staticmethod
    def _remove_quietly(path):
        try:
            os.remove(path)
        except OSError:
            pass

    def _run_background(self, task, on_done, label):
        """task(progress)'i ayrı iş parçacığında çalıştırır; ilerlemeyi ana iş parçacığında yoklar."""
        if self._bg_state is not None:
            messagebox.showwarning("Uyarı", "Başka bir işlem hâlâ sürüyor.", parent=self)
            return False
        state = {"done": 0, "total": 1, "label": label, "finished": False, "result": None, "error": None}
        self._bg_state = state

        def progress(done, total, text=None):
            state["done"], state["total"] = done, total
            if text:
                state["label"] = text

        def worker():
            try:
                state["result"] = task(progress)
            except Exception as exc:
                state["error"] = exc
            finally:
                state["finished"] = True

        self.operation_progress.configure(value=0, maximum=100)
        self.operation_status.configure(text=f"{label}...")
        self.config(cursor="watch")
        threading.Thread(target=worker, daemon=True).start()
        self._poll_background(on_done)
        return True

    def _poll_background(self, on_done):
        state = self._bg_state
        if state is None:
            return
        if state["finished"]:
            self._bg_state = None
            self.config(cursor="")
            self.operation_progress.configure(value=0)
            self.operation_status.configure(text="Hazır")
            on_done(state["result"], state["error"])
            return
        total = max(state["total"], 1)
        self.operation_progress.configure(value=(state["done"] / total) * 100)
        self.operation_status.configure(text=f"{state['label']} {state['done']}/{state['total']}")
        self.after(200, lambda: self._poll_background(on_done))

    # ------------------------------------------------------------------
    # Açılış (ana) sayfası
    # ------------------------------------------------------------------
    def hide_welcome(self):
        if self._welcome_frame is not None:
            self._welcome_frame.destroy()
            self._welcome_frame = None

    def _open_then(self, action=None):
        """Dosya seçtirip açar, ardından (varsa) action'ı çalıştırır; açılış sayfasındaki kısayol kartları için."""
        path = filedialog.askopenfilename(title="PDF veya e-kitap seçin", filetypes=OPEN_FILETYPES, parent=self)
        if not path:
            return
        try:
            self.open_pdf_path(path)
        except Exception as exc:  # noqa: BLE001
            messagebox.showerror("PDF Açma Hatası", f"PDF açılamadı:\n{exc}", parent=self)
            return
        if action and not self.is_welcome and self._is_same_file(path):
            self.after(150, action)

    def _welcome_run(self, action, needs_doc=True):
        """Açılış sayfasındaki düğme: belge gerekmiyorsa doğrudan; gerekiyorsa açık belgede ya da dosya seçtirip çalıştırır."""
        if not needs_doc:
            action()
        elif self.is_welcome:
            self._open_then(action)
        else:
            self.hide_welcome()
            self.after(120, action)

    def _popup_page_mode(self):
        menu = tk.Menu(self, tearoff=0)
        for key, (label, _colors) in PAGE_MODES.items():
            menu.add_radiobutton(label=label, value=key, variable=self.page_mode_var, command=self.set_page_mode)
        menu.tk_popup(self.winfo_pointerx(), self.winfo_pointery())

    def _toggle_check(self, var, command):
        var.set(not var.get())
        command()

    def _welcome_feature_groups(self):
        """Menülerdeki her özelliğin açılış sayfasında düğmesi olması için tek liste: (başlık, renk, [(simge, ad, eylem,
        belge_gerekir)]). Büyük kartlarda olan işler (birleştir, böl, OCR, imza, şifre, kitap açma…) burada tekrarlanmaz.
        tests/test_acilis_ozellikleri.py menüdeki her eylemin bu listede ya da kartlarda geçtiğini denetler."""
        return (
            ("📤  Dışa aktarma ve yazdırma", "mint", [
                ("📝", "Word (.docx) Olarak Kaydet", lambda: self.action_save_office("docx"), True),
                ("📊", "Excel (.xlsx) Olarak Kaydet", lambda: self.action_save_office("xlsx"), True),
                ("📽", "PowerPoint (.pptx) Olarak Kaydet", lambda: self.action_save_office("pptx"), True),
                ("📃", "Metne Dönüştür (.txt)", self.action_extract_text, True),
                ("🖼", "Sayfaları Resme Çevir", self.action_convert_images, True),
                ("🖼", "Görselleri Dışa Aktar", self.action_extract_images, True),
                ("📑", "Sayfa Aralığını PDF Kaydet", self.action_extract_pages, True),
                ("☑", "Seçilen Sayfaları PDF Kaydet", self.action_export_selected_pages, True),
                ("🗂", "İletişim Paftası Oluştur", self.action_contact_sheet, True),
                ("📘", "Kitapçık Düzeni", self.action_booklet, True),
                ("🔢", "Çok Sayfalı Düzen (2/4/6/9)", self.action_nup, True),
                ("📜", "Tek Uzun Sayfa", self.action_single_page, True),
                ("🖨", "Yazdır", self.action_print, True),
                ("📚", "Toplu Metin Çıkarma", self.action_batch_extract, False),
            ]),
            ("🗂  Sayfalar", "lavender", [
                ("✂", "Alan Seçerek Kırp", self.action_crop_tool, True),
                ("🔲", "Kenarları Kırp (mm)", self.action_crop, True),
                ("↩", "Kırpmayı Kaldır", self.action_crop_reset, True),
                ("➕", "Boş Sayfa Ekle", self.action_insert_blank, True),
                ("📎", "PDF'ye Sayfa Ekle", self.action_append_pages, True),
                ("🗑", "Sayfayı Sil", self.delete_current_page, True),
                ("☑", "Seçilen Sayfaları Sil", self.action_delete_selected_pages, True),
                ("🧹", "Sayfa Aralığını Sil", self.action_delete_page_range, True),
                ("⬜", "Boş Sayfaları Sil", self.action_remove_blank, True),
                ("⟳", "Sayfayı Döndür", self.rotate_current_page, True),
                ("🧭", "Yan/Ters Sayfaları Düzelt", self.action_auto_rotate, True),
                ("📐", "Sayfa Boyutu ve Arka Plan", self.action_page_layout, True),
                ("🔖", "Üst/Alt Bilgi ve Sayfa No", self.action_headers_footers, True),
            ]),
            ("🖍  Filigran, açıklama, form ve yer imi", "peach", [
                ("💧", "Filigran Ekle", self.action_watermark, True),
                ("🔏", "Damga Ekle", self.action_stamp, True),
                ("📄", "Başka PDF'i Üstüne Bindir", self.action_overlay, True),
                ("🖍", "Metni Ara ve Vurgula", lambda: self.action_annotation("highlight"), True),
                ("🔤", "Metni Ara ve Altını Çiz", lambda: self.action_annotation("underline"), True),
                ("🅂", "Metni Ara ve Üstünü Çiz", lambda: self.action_annotation("strikeout"), True),
                ("〰", "Metni Ara ve Dalgalı Altını Çiz", lambda: self.action_annotation("squiggly"), True),
                ("🅄", "Araç: Altını Çiz (sürükle)", lambda: self._select_tool("altcizgi"), True),
                ("🅂", "Araç: Üstünü Çiz (sürükle)", lambda: self._select_tool("ustucizgi"), True),
                ("〰", "Araç: Dalgalı Alt Çizgi (sürükle)", lambda: self._select_tool("dalga"), True),
                ("✏", "Araç: Serbest Kalem (sürükle)", lambda: self._select_tool("kalem"), True),
                ("⌫", "Araç: Silgi — yazıyı sil (sürükle)", lambda: self._select_tool("silgi"), True),
                ("🧽", "Tüm Açıklamaları Sil", self.action_remove_annotations, True),
                ("🧱", "Açıklama ve Formları Düzleştir", self.action_flatten, True),
                ("🔖", "Yer İmi Ekle", self.action_add_bookmark, True),
                ("📑", "Yer İmlerini Düzenle", self.action_edit_toc, True),
                ("📝", "Form Alanlarını Doldur", self.action_fill_form, True),
                ("🛠", "Form Editörü", self.action_form_editor, True),
                ("✒", "Görsel / İmza Resmi Yerleştir", self.action_add_signature, True),
            ]),
            ("🔐  Güvenlik", "rose", [
                ("🔓", "Şifreyi Kaldır", self.action_decrypt, True),
                ("🛂", "İzinleri Değiştir", self.action_permissions, True),
                ("⬛", "Metni Kalıcı Karart", self.action_redact, True),
                ("🧼", "Gizli Verileri Temizle", self.action_sanitize, True),
            ]),
            ("🛠  Denetim ve bakım", "sky", [
                ("ℹ", "Belge Bilgileri", self.action_document_info, True),
                ("🏷", "Metadata Düzenle", self.action_metadata, True),
                ("🩺", "PDF'yi Doğrula", self.action_check_document, True),
                ("🔧", "Bozuk Dosyayı Onar", self.action_repair, True),
                ("🗑", "Tüm Görselleri Sil", self.action_remove_images, True),
                ("⚖", "İki PDF'i Karşılaştır", self.action_compare, True),
                ("📎", "Gömülü Ekleri Yönet", self.action_attachments, True),
                ("🔍", "PDF İçinde Ara", self.action_search, True),
                ("⌨", "Komut Ara (Ctrl+K)", self.action_command_palette, False),
            ]),
            ("👁  Görünüm", "lemon", [
                ("📑", "İçindekiler Paneli", self.toggle_toc_panel, True),
                ("🧅", "Katmanlar", self.action_layers, True),
                ("🎞", "Sayfa Sıralama Paneli", self.toggle_thumbnail_panel, True),
                ("⬅", "Geri Git (Alt+←)", self.nav_back, True),
                ("➡", "İleri Git (Alt+→)", self.nav_forward, True),
                ("🎨", "Sayfa Rengi…", self._popup_page_mode, False),
                ("↕", "Dikey Sayfa", lambda: self._toggle_check(self.portrait_var, lambda: self.set_orientation("dikey")), True),
                ("↔", "Yatay Sayfa", lambda: self._toggle_check(self.landscape_var, lambda: self.set_orientation("yatay")), True),
                ("⛶", "Tam Ekran (F11)", lambda: self.toggle_fullscreen(True), False),
                ("🪟", "Pencere Boyutu Kilidi", lambda: self._toggle_check(self.resizable_var, self.toggle_resizable), False),
                ("📖", "E-Kitap Okuma Ayarları", self.action_epub_settings, False),
                ("📂", "E-kitap PDF'ini Yanına Kaydet",
                 lambda: self._toggle_check(self.epub_beside_var,
                                            lambda: save_setting("epub_beside", bool(self.epub_beside_var.get()))), False),
                ("💾", "Farklı Kaydet…", self.save_pdf_as, True),
                ("↶", "Geri Al", self.undo, True),
                ("↷", "Yinele", self.redo, True),
            ]),
            ("❓  Yardım ve ayarlar", "lavender", [
                ("⌨", "Klavye Kısayolları", self.action_shortcuts, False),
                ("📋", "Hata Günlüğünü Aç", self.action_open_error_log, False),
                ("⭐", "Varsayılan PDF Programı Yap", self.action_set_default_app, False),
                ("🔄", "Güncellemeleri Denetle", self.action_check_update, False),
                ("⚖", "Lisans ve Kaynak Kodu", self.action_license, False),
                ("ℹ", f"{APP_NAME} Hakkında", self.action_about, False),
            ]),
        )

    def show_welcome(self):
        """Pastel kartlı açılış sayfasını pencerenin üstüne getirir (belge kapanmaz). Sayfa kaydırılabilir; tüm
        özellikler için büyük kart ya da düğme içerir."""
        t = self.theme
        self.hide_welcome()
        if self.is_welcome:
            self.title(f"{APP_NAME} - Ana Sayfa")
        frame = tk.Frame(self, bg=t["bg"])
        self._welcome_frame = frame
        frame.place(x=0, y=0, relwidth=1, relheight=1)
        frame.lift()
        if DND_AVAILABLE:  # sürükle-bırak, kayıtlı pencere (widget) üzerinde çalışır
            frame.drop_target_register(DND_FILES)
            frame.dnd_bind("<<Drop>>", self.on_files_dropped)

        scroll = tk.Canvas(frame, bg=t["bg"], highlightthickness=0)
        scroll_bar = ttk.Scrollbar(frame, orient="vertical", command=scroll.yview)
        scroll.configure(yscrollcommand=scroll_bar.set)
        scroll_bar.pack(side="right", fill="y")
        scroll.pack(side="left", fill="both", expand=True)
        wrap = tk.Frame(scroll, bg=t["bg"])
        window = scroll.create_window((0, 0), window=wrap, anchor="n")

        def relayout(_event=None):
            scroll.coords(window, scroll.winfo_width() // 2, 8)
            scroll.configure(scrollregion=(0, 0, scroll.winfo_width(), wrap.winfo_reqheight() + 70))

        scroll.bind("<Configure>", relayout)
        wrap.bind("<Configure>", relayout)

        theme_toggle = tk.Label(frame, text="☀ Açık tema" if self.dark_mode_var.get() else "🌙 Koyu tema",
                                bg=t["panel"], fg=t["text"], font=(FONT_UI, 9), padx=10, pady=4, cursor="hand2")
        theme_toggle.place(relx=1.0, x=-30, y=14, anchor="ne")
        theme_toggle.bind("<Button-1>", lambda _e: (self.dark_mode_var.set(not self.dark_mode_var.get()),
                                                    self.toggle_dark_mode()))

        title_row = tk.Frame(wrap, bg=t["bg"])
        title_row.pack(pady=(14, 0))
        tk.Label(title_row, text="📄", bg=t["bg"], font=("Segoe UI Emoji", 34)).pack(side="left", padx=(0, 10))
        tk.Label(title_row, text=APP_NAME, bg=t["bg"], fg=t["text"], font=(FONT_UI, 32, "bold")).pack(side="left")
        tk.Label(wrap, text="PDF'lerinizi düzenleyin, birleştirin, şifreleyin; EPUB, MOBI, AZW3, IBA ve RTF belgelerinizi okuyun — hepsi bilgisayarınızda, çevrimdışı.",
                 bg=t["bg"], fg=t["muted"], font=(FONT_UI, 11)).pack(pady=(2, 14))

        if not self.is_welcome:  # açık bir belge varsa ona dönüş yolu
            back = tk.Label(wrap, text=f"↩  {os.path.basename(self.pdf_path)} belgesine dön", bg=t["accent"],
                            fg=t["text"], font=(FONT_UI, 10, "bold"), padx=16, pady=6, cursor="hand2")
            back.pack(pady=(0, 12))
            back.bind("<Button-1>", lambda _e: self.hide_welcome())

        def make_card(grid, row, col, icon, title, desc, color, command):
            bg = t[color]
            card = tk.Frame(grid, bg=bg, width=196, height=92, highlightthickness=2,
                            highlightbackground=bg, cursor="hand2")
            card.grid(row=row, column=col, padx=5, pady=3)
            card.pack_propagate(False)
            icon_label = tk.Label(card, text=icon, bg=bg, fg=t["text"], font=("Segoe UI Emoji", 22))
            icon_label.place(x=12, y=8)
            title_label = tk.Label(card, text=title, bg=bg, fg=t["text"], font=(FONT_UI, 11, "bold"), anchor="w")
            title_label.place(x=62, y=16)
            desc_label = tk.Label(card, text=desc, bg=bg, fg=t["muted"], font=(FONT_UI, 9), wraplength=172,
                                  justify="left", anchor="nw")
            desc_label.place(x=12, y=52)
            for widget in (card, icon_label, title_label, desc_label):
                widget.bind("<Enter>", lambda _e, c=card: c.config(highlightbackground=t["accent_dark"]))
                widget.bind("<Leave>", lambda _e, c=card, b=bg: c.config(highlightbackground=b))
                widget.bind("<Button-1>", lambda _e: command())

        def make_chip(grid, row, col, icon, title, color, command):
            bg = t[color]
            chip = tk.Frame(grid, bg=bg, width=216, height=36, highlightthickness=2,
                            highlightbackground=bg, cursor="hand2")
            chip.grid(row=row, column=col, padx=4, pady=2)
            chip.pack_propagate(False)
            icon_label = tk.Label(chip, text=icon, bg=bg, fg=t["text"], font=("Segoe UI Emoji", 12))
            icon_label.pack(side="left", padx=(8, 4))
            title_label = tk.Label(chip, text=title, bg=bg, fg=t["text"], font=(FONT_UI, 9, "bold"), anchor="w")
            title_label.pack(side="left", fill="x", expand=True)
            for widget in (chip, icon_label, title_label):
                widget.bind("<Enter>", lambda _e, c=chip: c.config(highlightbackground=t["accent_dark"]))
                widget.bind("<Leave>", lambda _e, c=chip, b=bg: c.config(highlightbackground=b))
                widget.bind("<Button-1>", lambda _e: command())

        pdf_cards = (
            ("📂", "PDF Aç", "Bir PDF seçip görüntülemeye ve düzenlemeye başlayın.", "lavender", self.action_open_new),
            ("🧩", "Birleştir", "Birden fazla PDF'i tek dosyada toplayın.", "mint", self.action_merge),
            ("🖼", "Resimden PDF", "JPG/PNG resimlerini tek PDF'e çevirin.", "peach", self.action_images_to_pdf),
            ("✂", "Sayfalara Böl", "Her sayfayı ayrı bir PDF dosyası yapın.", "sky", lambda: self._welcome_run(self.action_split)),
            ("🗜", "Boyut Küçült", "Dosyayı sıkıştırıp daha hafif hale getirin.", "lemon", lambda: self._welcome_run(self.action_compress)),
            ("🔎", "OCR ile Metin", "Taranmış belgelerden yazıyı çıkarın.", "rose", lambda: self._welcome_run(self.action_ocr)),
            ("🔒", "Şifrele", "PDF'yi parolayla koruyun.", "sky", lambda: self._welcome_run(self.action_encrypt)),
            ("📝", "Word / Excel / PPT", "PDF'yi Office biçiminde farklı kaydedin.", "mint", lambda: self._welcome_run(lambda: self.action_save_office("docx"))),
            ("✂", "Sayfa Kırp", "Alan seçerek kırpın; istediğiniz an kırpmayı kaldırın.", "lemon", lambda: self._welcome_run(self.action_crop_tool)),
        )
        book_cards = (
            ("📖", "EPUB Kitap Aç", "E-kitapları okuyun; kaldığınız yerden devam edin.", "mint", self.action_open_epub),
            ("📘", "MOBI Kitap Aç", "Kindle MOBI e-kitaplarını okuyun.", "sky", lambda: self.action_open_book("MOBI")),
            ("📙", "AZW3 Kitap Aç", "Kindle AZW3 kitaplarını okuyun (DRM'siz).", "lemon", lambda: self.action_open_book("AZW3")),
            ("📗", "IBA Kitap Aç", "iBooks Author belgelerinin metin ve resimleri.", "rose", lambda: self.action_open_book("IBA")),
            ("📝", "RTF Belgesi Aç", "Zengin metin (RTF) belgelerini görüntüleyin.", "peach", lambda: self.action_open_book("RTF")),
            ("📕", "FB2 Kitap Aç", "FictionBook (FB2) kitaplarını okuyun.", "lavender", lambda: self.action_open_book("FB2")),
            ("🗎", "XPS Belgesi Aç", "XPS / OXPS belgelerini görüntüleyin.", "mint", lambda: self.action_open_book("XPS")),
            ("🗯", "CBZ Çizgi Roman Aç", "CBZ çizgi roman arşivlerini okuyun.", "sky", lambda: self.action_open_book("CBZ")),
            ("🎨", "SVG Resmi Aç", "SVG çizimlerini PDF olarak görüntüleyin.", "lemon", lambda: self.action_open_book("SVG")),
        )

        def build_group(title, group_cards):
            tk.Label(wrap, text=title, bg=t["bg"], fg=t["muted"], font=(FONT_UI, 10, "bold"), anchor="w").pack(
                fill="x", padx=9, pady=(6, 0))
            grid = tk.Frame(wrap, bg=t["bg"])
            grid.pack()
            for index, card in enumerate(group_cards):
                make_card(grid, index // 5, index % 5, *card)

        build_group("📄  PDF araçları", pdf_cards)
        build_group("📚  E-kitaplar ve belgeler (EPUB · MOBI · AZW3 · IBA · RTF · FB2 · XPS · CBZ · SVG)", book_cards)

        # Özellik düğmeleri sayfayı doldurmaz: sol üstteki (Dosya menüsüyle hizalı) çekmecede, başlıklara göre açılır.
        drawer = tk.Frame(frame, bg=t["panel"], width=268, highlightthickness=1, highlightbackground=t["card_border"])
        drawer_canvas = tk.Canvas(drawer, bg=t["panel"], highlightthickness=0, width=250)
        drawer_bar = ttk.Scrollbar(drawer, orient="vertical", command=drawer_canvas.yview)
        drawer_canvas.configure(yscrollcommand=drawer_bar.set)
        drawer_bar.pack(side="right", fill="y")
        drawer_canvas.pack(side="left", fill="both", expand=True)
        drawer_body = tk.Frame(drawer_canvas, bg=t["panel"])
        drawer_canvas.create_window((0, 0), window=drawer_body, anchor="nw")
        drawer_body.bind("<Configure>", lambda _e: drawer_canvas.configure(scrollregion=drawer_canvas.bbox("all")))
        drawer_canvas.bind("<MouseWheel>", lambda e: drawer_canvas.yview_scroll(-1 if e.delta > 0 else 1, "units"))
        drawer_state = {"open": False}
        drawer_button = tk.Label(frame, bg=t["accent"], fg=t["text"], font=(FONT_UI, 10, "bold"),
                                 padx=12, pady=4, cursor="hand2")

        def set_drawer(open_):
            drawer_state["open"] = open_
            drawer_button.config(text="✕  Özellikler" if open_ else "☰  Özellikler")
            if open_:
                drawer.place(x=0, y=40, relheight=1.0, height=-40)
                drawer.lift()
            else:
                drawer.place_forget()
            drawer_button.lift()

        def run_from_drawer(action, needs_doc):
            set_drawer(False)
            self._welcome_run(action, needs_doc)

        def make_section(title, color, items):
            holder = tk.Frame(drawer_body, bg=t["panel"])
            holder.pack(fill="x", padx=6, pady=(4, 0))
            head = tk.Label(holder, text="▸  " + title, bg=t[color], fg=t["text"], font=(FONT_UI, 10, "bold"),
                            anchor="w", padx=8, pady=6, cursor="hand2", width=28)
            head.pack(fill="x")
            body = tk.Frame(holder, bg=t["panel"])
            grid_state = {"open": False, "built": False}

            def toggle(_event=None):
                if not grid_state["built"]:  # düğmeler ilk açılışta üretilir: açılış sayfası hızlı gelir
                    grid_state["built"] = True
                    for index, (icon, label, action, needs_doc) in enumerate(items):
                        make_chip(body, index, 0, icon, label, color,
                                  lambda a=action, n=needs_doc: run_from_drawer(a, n))
                    bind_drawer_wheel(body)
                grid_state["open"] = not grid_state["open"]
                head.config(text=("▾  " if grid_state["open"] else "▸  ") + title)
                if grid_state["open"]:
                    body.pack(fill="x")
                else:
                    body.pack_forget()

            head.bind("<Button-1>", toggle)

        for title, color, items in self._welcome_feature_groups():
            make_section(title, color, items)

        drawer_button.place(x=8, y=6)
        drawer_button.bind("<Button-1>", lambda _e: set_drawer(not drawer_state["open"]))
        set_drawer(False)

        recent = [path for path in load_recent_files() if os.path.exists(path)][:5]
        recent_head = tk.Frame(wrap, bg=t["bg"])
        recent_head.pack(pady=(14, 4))
        tk.Label(recent_head, text="Kitaplığım · kaldığınız yerden devam edin" if recent else "Son açılanlar",
                 bg=t["bg"], fg=t["text"], font=(FONT_UI, 11, "bold")).pack(side="left")
        if recent:
            clear = tk.Label(recent_head, text="🧹 Listeyi temizle", bg=t["panel"], fg=t["text"], font=(FONT_UI, 9),
                             padx=8, pady=2, cursor="hand2")
            clear.pack(side="left", padx=(12, 0))
            clear.bind("<Button-1>", lambda _e: self.action_clear_recent())
        recent_row = tk.Frame(wrap, bg=t["bg"])
        recent_row.pack()
        if not recent:
            tk.Label(recent_row, text="Henüz dosya açmadınız.", bg=t["bg"], fg=t["muted"], font=(FONT_UI, 10)).pack()
        self._library_images = []

        def on_wheel(event):
            scroll.yview_scroll(-1 if event.delta > 0 else 1, "units")

        def bind_wheel(widget):
            widget.bind("<MouseWheel>", on_wheel)
            for child in widget.winfo_children():
                bind_wheel(child)

        bind_wheel(scroll)
        bind_wheel(wrap)
        def bind_drawer_wheel(widget):
            widget.bind("<MouseWheel>", lambda e: drawer_canvas.yview_scroll(-1 if e.delta > 0 else 1, "units"))
            for child in widget.winfo_children():
                bind_drawer_wheel(child)

        bind_drawer_wheel(drawer_body)

        def add_library_cards(position=0):
            """Kitaplık kartları (kapak üretimi yavaş olabilir) sayfa görüntülendikten sonra tek tek eklenir."""
            if self._welcome_frame is not frame or position >= len(recent):
                return
            try:
                self._make_library_card(recent_row, recent[position])
                bind_wheel(recent_row)
            except Exception as exc:  # noqa: BLE001
                log_exception("Kitaplık kartı", exc)
            self.after(15, lambda: add_library_cards(position + 1))

        self.after(60, add_library_cards)

        hint = tk.Frame(frame, bg=t["bg"])
        hint.place(relx=0.5, rely=1.0, y=-14, anchor="s")
        tk.Label(hint, text="💡 PDF, EPUB, MOBI, AZW3, IBA veya RTF dosyasını bu pencereye sürükleyip bırakabilirsiniz  ·  Ctrl+O ile açın",
                 bg=t["bg"], fg=t["muted"], font=(FONT_UI, 9)).pack(side="left")
        about = tk.Label(hint, text=f"  ·  Sürüm {APP_VERSION}  ·  Hakkında", bg=t["bg"], fg=t["accent_dark"],
                         font=(FONT_UI, 9, "underline"), cursor="hand2")
        about.pack(side="left")
        about.bind("<Button-1>", lambda _e: self.action_about())
        hint.lift()
        self.after_idle(relayout)

    def _make_library_card(self, parent, path):
        """Kapak, başlık ve okuma ilerlemesiyle kitaplık kartı; tıklanınca kitabı kaldığı yerden açar."""
        t = self.theme
        info = library_info(path)
        progress = reading_progress(path)
        card = tk.Frame(parent, bg=t["entry"], width=150, height=196, highlightthickness=2,
                        highlightbackground=t["card_border"], cursor="hand2")
        card.pack(side="left", padx=6)
        card.pack_propagate(False)
        is_book = is_book_file(path)
        cover = None
        if info["cover"]:
            try:
                cover = tk.PhotoImage(file=info["cover"])
                self._library_images.append(cover)
            except tk.TclError:
                cover = None
        if cover:
            cover_label = tk.Label(card, image=cover, bg=t["entry"], bd=1, relief="solid")
        else:
            cover_label = tk.Label(card, text="📖" if is_book else "📄", bg=t["lavender"], fg=t["text"],
                                   font=("Segoe UI Emoji", 28), width=5, height=2)
        cover_label.pack(pady=(6, 3))
        name = info["title"]
        title_label = tk.Label(card, text=name if len(name) <= 36 else name[:33] + "…", bg=t["entry"], fg=t["text"],
                               font=(FONT_UI, 9, "bold"), wraplength=136, justify="center", height=2)
        title_label.pack(padx=6)
        bar = tk.Canvas(card, width=128, height=6, bg=t["thumb_bg"], highlightthickness=0)
        bar.pack(pady=(2, 2))
        if progress:
            bar.create_rectangle(0, 0, 128 * progress[0] / 100, 6, fill=t["accent_dark"], width=0)
            status = f"%{progress[0]} · s. {progress[1]}/{progress[2]}"
        else:
            status = "Yeni"
        status_label = tk.Label(card, text=status, bg=t["entry"], fg=t["muted"], font=(FONT_UI, 8))
        status_label.pack()
        widgets = (card, cover_label, title_label, bar, status_label)
        for widget in widgets:
            widget.bind("<Enter>", lambda _e, c=card: c.config(highlightbackground=t["accent_dark"]))
            widget.bind("<Leave>", lambda _e, c=card: c.config(highlightbackground=t["card_border"]))
            widget.bind("<Button-1>", lambda _e, p=path: self._open_then_path(p))
        if info["author"]:
            self._tooltip(title_label, f"{info['title']}\n{info['author']}")

    def _open_then_path(self, path):
        try:
            self.open_pdf_path(path)
        except Exception as exc:  # noqa: BLE001
            messagebox.showerror("PDF Açma Hatası", f"PDF açılamadı:\n{exc}", parent=self)

    # ------------------------------------------------------------------
    # Görünüm
    # ------------------------------------------------------------------
    def init_view(self):
        self.fit_width()

    def set_page_mode(self):
        save_setting("page_mode", self.page_mode_var.get())
        index = self.get_current_page_index() if self.page_positions else 0
        self.render_thumbnails()
        self.render_all_pages()
        if self.page_positions:
            self.page_var.set(str(index + 1))
            self.goto_page()

    def toggle_dark_mode(self):
        save_setting("dark_mode", bool(self.dark_mode_var.get()))
        self.apply_theme()
        self.render_thumbnails()
        self.render_all_pages()
        if self._welcome_frame is not None:
            self.show_welcome()

    def apply_theme(self):
        """Pastel açık / koyu temayı tüm ttk ve tk bileşenlerine uygular."""
        t = self.theme = THEMES["dark" if self.dark_mode_var.get() else "light"]
        style = ttk.Style(self)
        try:
            style.theme_use("clam")
        except tk.TclError:
            pass
        self.configure(bg=t["bg"])
        self.option_add("*Toplevel.background", t["bg"])
        self.option_add("*Menu.background", t["panel"])
        self.option_add("*Menu.foreground", t["text"])
        self.option_add("*Menu.activeBackground", t["accent"])
        self.option_add("*Menu.activeForeground", t["text"])
        font = (FONT_UI, 9)
        style.configure(".", background=t["bg"], foreground=t["text"], font=font,
                        fieldbackground=t["entry"], bordercolor=t["card_border"],
                        lightcolor=t["bg"], darkcolor=t["bg"], troughcolor=t["panel"])
        style.configure("TFrame", background=t["bg"])
        style.configure("Bar.TFrame", background=t["panel"])
        style.configure("Tools.TFrame", background=t["lavender"])
        style.configure("TLabel", background=t["bg"], foreground=t["text"])
        style.configure("Bar.TLabel", background=t["panel"], foreground=t["text"])
        style.configure("Tools.TLabel", background=t["lavender"], foreground=t["muted"])
        style.configure("TButton", background=t["lavender"], foreground=t["text"], borderwidth=0,
                        padding=(6, 4), focusthickness=0, focuscolor=t["lavender"])
        style.map("TButton", background=[("active", t["accent"]), ("pressed", t["accent_hover"])])
        style.configure("Icon.TButton", font=("Segoe UI Emoji", 11), padding=(4, 2))
        style.configure("Save.TButton", background=t["mint"], font=(FONT_UI, 9, "bold"))
        style.map("Save.TButton", background=[("active", t["accent"]), ("pressed", t["accent_hover"])])
        style.configure("Tool.Toolbutton", background=t["lavender"], foreground=t["text"], borderwidth=0,
                        padding=(6, 3), focuscolor=t["lavender"], font=("Segoe UI Emoji", 12))
        style.map("Tool.Toolbutton", background=[("selected", t["accent"]), ("active", t["peach"])])
        style.configure("TCheckbutton", background=t["bg"], foreground=t["text"])
        style.configure("Bar.TCheckbutton", background=t["panel"], foreground=t["text"], font=("Segoe UI Emoji", 11))
        style.map("TCheckbutton", background=[("active", t["bg"])])
        style.map("Bar.TCheckbutton", background=[("active", t["panel"])])
        style.configure("TEntry", fieldbackground=t["entry"], foreground=t["text"], bordercolor=t["card_border"])
        style.configure("TCombobox", fieldbackground=t["entry"], foreground=t["text"], background=t["lavender"])
        style.configure("TSeparator", background=t["card_border"])
        style.configure("TProgressbar", background=t["accent"], troughcolor=t["bg"], bordercolor=t["panel"])
        style.configure("TScrollbar", background=t["lavender"], troughcolor=t["panel"], bordercolor=t["panel"],
                        arrowcolor=t["text"])
        style.configure("TPanedwindow", background=t["panel"])
        style.configure("Treeview", background=t["entry"], fieldbackground=t["entry"], foreground=t["text"])
        style.configure("Treeview.Heading", background=t["lavender"], foreground=t["text"])
        style.map("Treeview", background=[("selected", t["accent"])], foreground=[("selected", t["text"])])
        self.canvas.configure(bg=t["canvas"])
        self.thumb_canvas.configure(bg=t["thumb_bg"])
        if self.toast_label.cget("text"):
            self.toast_label.config(bg=t["mint"], fg=t["text"])
            self.toast_frame.config(bg=t["mint"])

    def toggle_thumbnail_panel(self):
        """Sol taraftaki sayfa sıralama panelini açar veya kapatır."""
        if self.is_thumb_open:
            self.paned.forget(self.thumb_frame)
            self.is_thumb_open = False
        else:
            self.paned.insert(0, self.thumb_frame, weight=1)
            self.is_thumb_open = True
            self.render_thumbnails()

    def _rerender_keep_page(self):
        """Pencere veya yön değişince aynı sayfada kalarak sayfaları yeniden çizer."""
        idx = self.get_current_page_index()
        self.render_all_pages()
        self.page_var.set(str(idx + 1))
        self.goto_page()

    def toggle_fullscreen(self, _event=None):
        if _event is not None:  # F11 ile çağrıldıysa menüdeki işareti eşitle
            self.fullscreen_var.set(not self.fullscreen_var.get())
        full = self.fullscreen_var.get()
        if full:
            self.resizable(True, True)
        else:
            allowed = self.resizable_var.get()
            self.resizable(allowed, allowed)
        self.attributes("-fullscreen", full)
        self.after(250, self._rerender_keep_page)

    def exit_fullscreen(self, _event=None):
        if self.fullscreen_var.get():
            self.fullscreen_var.set(False)
            self.toggle_fullscreen()

    def toggle_resizable(self):
        """Pencere kenarlarından boyutlandırmaya izin verir veya pencere boyutunu kilitler."""
        allowed = self.resizable_var.get()
        if not self.fullscreen_var.get():
            self.resizable(allowed, allowed)

    def _page_is_rotated_for_view(self, page):
        wanted = self.page_orientation
        return bool(wanted) and ((wanted == "yatay") != (page.rect.width > page.rect.height))

    def _display_page_size(self, page):
        """Seçili yöne (dikey/yatay) göre sayfanın görünen genişlik ve yüksekliği."""
        w, h = page.rect.width, page.rect.height
        return (h, w) if self._page_is_rotated_for_view(page) else (w, h)

    def _render_matrix(self, page, scale):
        """Seçili yöne göre görüntüleme matrisi (dosyadaki sayfa değişmez)."""
        matrix = fitz.Matrix(scale, scale)
        if self._page_is_rotated_for_view(page):
            matrix.prerotate(90)
        return matrix

    def set_orientation(self, which):
        # Aynı seçeneğe tekrar tıklanınca yön zorlaması kapanır (sayfalar dosyadaki gibi görünür)
        if which == "dikey" and self.portrait_var.get():
            self.landscape_var.set(False)
        elif which == "yatay" and self.landscape_var.get():
            self.portrait_var.set(False)
        if self.portrait_var.get():
            self.page_orientation = "dikey"
        elif self.landscape_var.get():
            self.page_orientation = "yatay"
        else:
            self.page_orientation = None
        idx = self.get_current_page_index()
        self.render_thumbnails()
        self.fit_width()  # genişliğe göre ölçekler ve sayfaları yeniden çizer
        self.page_var.set(str(idx + 1))
        self.goto_page()

    # ------------------------------------------------------------------
    # Küçük resimler (sayfa sıralama)
    # ------------------------------------------------------------------
    def render_thumbnails(self):
        if not self.is_thumb_open:
            return

        for child in self.thumb_container.winfo_children():
            child.destroy()
        self.thumb_objects.clear()
        self.thumb_cards.clear()
        self._thumb_marker = None

        for i, page in enumerate(self.doc):
            pix = tint_pixmap(page.get_pixmap(matrix=self._render_matrix(page, 0.2), alpha=False), self.page_mode_var.get())
            img = tk.PhotoImage(data=pix.tobytes("ppm"))
            self.thumb_objects.append(img)

            t = self.theme
            card_bg = t["accent"] if i in self.selected_pages else t["lavender"]
            card = tk.Frame(self.thumb_container, bg=card_bg, bd=2, relief="groove", cursor="hand2")
            card.pack(fill="x", padx=10, pady=5)
            self.thumb_cards.append(card)

            lbl_img = tk.Label(card, image=img, bg="#ffffff")
            lbl_img.pack(pady=(5, 2), padx=5)

            lbl_num = tk.Label(card, text=f"Sayfa {i+1}", bg=card_bg, fg=t["text"], font=(FONT_UI, 9, "bold"))
            lbl_num.pack(pady=(0, 5))

            for w in (card, lbl_img, lbl_num):
                w.bind("<Button-1>", lambda e, idx=i: self.select_thumb(idx, bool(e.state & 0x0001)))
                w.bind("<B1-Motion>", lambda e, idx=i: self.on_thumb_drag(e, idx))
                w.bind("<ButtonRelease-1>", lambda e, idx=i: self.on_thumb_drop(e, idx))

    def select_thumb(self, idx, multi=False):
        self.drag_data["index"] = None
        if multi:
            self.toggle_page_selection(idx)
        else:
            self.page_var.set(str(idx + 1))
            self.goto_page()

    # ------------------------------------------------------------------
    # Sayfa üstü araçlar (kopyala, vurgula, not, yazı, çizgi, görsel, gizle)
    # ------------------------------------------------------------------
    def _view_to_page(self, cx, cy, index=None):
        """Canvas noktasını (sayfa indeksi, sayfanın görünen koordinatı) olarak çevirir; sayfa dışıysa (None, None).
        index verilirse nokta, imleç başka sayfada olsa bile o sayfanın düzleminde hesaplanır."""
        if index is None:
            for index, (top, bottom) in enumerate(self.page_positions):
                if top <= cy <= bottom:
                    break
            else:
                return None, None
        top = self.page_positions[index][0]
        page = self.doc[index]
        matrix = self._render_matrix(page, self.zoom)
        origin = page.rect * matrix
        point = fitz.Point(cx - self.page_x[index] + origin.x0, cy - top + origin.y0) * ~matrix
        return index, point

    @staticmethod
    def _api_rect(page, p0, p1):
        """Görünen koordinattaki iki köşeyi PyMuPDF'nin (döndürülmemiş) sayfa uzayında Rect'e çevirir."""
        rect = fitz.Rect(p0, p1).normalize() * page.derotation_matrix
        return rect.normalize() & page.mediabox

    def on_canvas_press(self, event):
        self.canvas.focus_set()
        cx, cy = self.canvas.canvasx(event.x), self.canvas.canvasy(event.y)
        if self._te and self.tool_mode == "metin" and not self.is_welcome:
            self._te_press(cx, cy)
            return
        self._drag_start = (cx, cy)
        mode = self.tool_mode
        if mode == "gez":
            self.canvas.scan_mark(event.x, event.y)
        elif mode in ("kopyala", "vurgula", "altcizgi", "ustucizgi", "dalga", "silgi", "gizle", "kirp"):
            self._sel_id = self.canvas.create_rectangle(cx, cy, cx, cy, outline=self.theme["accent_dark"],
                                                        width=2, dash=(4, 2))
        elif mode == "kalem":
            self._ink_points = [(cx, cy)]
            self._ink_items = []
        elif mode == "cizgi":
            self._sel_id = self.canvas.create_line(cx, cy, cx, cy, fill=self.theme["accent_dark"], width=2)

    def on_canvas_drag(self, event):
        if self._te and self._te.get("drag"):
            self._te_motion(self.canvas.canvasx(event.x), self.canvas.canvasy(event.y))
            return
        if self._drag_start is None:
            return
        if self.tool_mode == "gez":
            self.canvas.scan_dragto(event.x, event.y, gain=1)
            self._schedule_visible()
        elif self.tool_mode == "kalem":
            cx, cy = self.canvas.canvasx(event.x), self.canvas.canvasy(event.y)
            last = self._ink_points[-1]
            if abs(cx - last[0]) + abs(cy - last[1]) >= 2:
                self._ink_items.append(self.canvas.create_line(*last, cx, cy, fill="#D81B1B", width=2))
                self._ink_points.append((cx, cy))
        elif self._sel_id is not None:
            cx, cy = self.canvas.canvasx(event.x), self.canvas.canvasy(event.y)
            self.canvas.coords(self._sel_id, *self._drag_start, cx, cy)

    def on_canvas_release(self, event):
        if self._te and self._te.get("drag"):
            self._te["drag"] = None
            return
        start, self._drag_start = self._drag_start, None
        if self._sel_id is not None:
            self.canvas.delete(self._sel_id)
            self._sel_id = None
        mode = self.tool_mode
        for item in getattr(self, "_ink_items", ()):
            self.canvas.delete(item)
        self._ink_items = []
        points, self._ink_points = getattr(self, "_ink_points", []), []
        if start is None or self.is_welcome:
            return
        end = (self.canvas.canvasx(event.x), self.canvas.canvasy(event.y))
        if mode == "kalem":
            try:
                if len(points) >= 2:
                    self._tool_kalem(points)
            except Exception as exc:  # noqa: BLE001
                show_runtime_error(self, "Araç hatası", exc)
            return
        if mode == "gez":  # sürüklemeden tek tıklama: sayfadaki bağlantıyı izle
            if abs(end[0] - start[0]) < 4 and abs(end[1] - start[1]) < 4:
                try:
                    _index, link = self._link_under(*end)
                    if link:
                        self._follow_link(link)
                except Exception as exc:  # noqa: BLE001
                    show_runtime_error(self, "Bağlantı hatası", exc)
            return
        try:
            if mode in ("not", "metin", "gorsel"):
                index, point = self._view_to_page(*start)
                if index is not None:
                    getattr(self, f"_tool_{mode}")(index, point)
            else:
                if abs(end[0] - start[0]) < 4 and abs(end[1] - start[1]) < 4:
                    return
                index, p0 = self._view_to_page(*start)
                if index is None:
                    return
                _same, p1 = self._view_to_page(end[0], end[1], index=index)
                getattr(self, f"_tool_{mode}")(index, p0, p1)
        except Exception as exc:  # noqa: BLE001
            show_runtime_error(self, "Araç hatası", exc)

    def _tool_kopyala(self, index, p0, p1):
        page = self.doc[index]
        text = page.get_text("text", clip=self._api_rect(page, p0, p1)).replace(" ", " ").strip()
        if not text:
            self._toast("Bu alanda metin bulunamadı (taranmış sayfaysa OCR kullanın).")
            return
        self.clipboard_clear()
        self.clipboard_append(text)
        self._toast(f"{len(text)} karakter panoya kopyalandı")

    def _tool_markup(self, kind, index, p0, p1):
        page = self.doc[index]
        words = page.get_text("words", clip=self._api_rect(page, p0, p1))
        if not words:
            self._toast("Bu alanda işaretlenecek metin yok.")
            return
        self._record_edit()
        page = self.doc[index]
        add_markup(page, [w[:4] for w in words], kind)
        self.render_all_pages()

    def _tool_vurgula(self, index, p0, p1):
        self._tool_markup("highlight", index, p0, p1)

    def _tool_altcizgi(self, index, p0, p1):
        self._tool_markup("underline", index, p0, p1)

    def _tool_ustucizgi(self, index, p0, p1):
        self._tool_markup("strikeout", index, p0, p1)

    def _tool_dalga(self, index, p0, p1):
        self._tool_markup("squiggly", index, p0, p1)

    def _tool_silgi(self, index, p0, p1):
        page = self.doc[index]
        clip = self._api_rect(page, p0, p1)
        if clip.is_empty:
            return
        if not page.get_text("words", clip=clip) and not any(
                fitz.Rect(info["bbox"]).intersects(clip) for info in page.get_image_info()):
            self._toast("Bu alanda silinecek yazı ya da resim yok.")
            return
        self._record_edit()
        PDFProcessor.erase_area(self.doc[index], clip)
        self.render_all_pages()
        self._toast("Seçilen alan silindi (Ctrl+Z ile geri alınır)")

    def _tool_kalem(self, view_points):
        """Fare yolunu (canvas noktaları) ilk noktanın sayfasına serbest kalem çizgisi olarak ekler."""
        index, first = self._view_to_page(*view_points[0])
        if index is None:
            return
        page = self.doc[index]
        stroke = []
        for cx, cy in view_points:
            _same, point = self._view_to_page(cx, cy, index=index)
            stroke.append(point * page.derotation_matrix)
        self._record_edit()
        add_ink(self.doc[index], [[(p.x, p.y) for p in stroke]])
        self.render_all_pages()

    def _select_tool(self, key):
        """Araç çubuğundaki aracı seçer (menü ve açılış sayfasından çağrılır)."""
        self.tool_var.set(key)
        self._on_tool_change()

    def _tool_gizle(self, index, p0, p1):
        page = self.doc[index]
        rect = self._api_rect(page, p0, p1)
        if rect.is_empty:
            return
        if not messagebox.askyesno(
                "Kalıcı gizleme (redaksiyon)",
                "Seçili alandaki metin ve resim parçaları PDF'ten kalıcı olarak silinir "
                "(siyah kutu ile kapatılır).\n\nKaydetmeden önce Ctrl+Z ile geri alabilirsiniz. Devam edilsin mi?",
                icon="warning", parent=self):
            return
        self._record_edit()
        page = self.doc[index]
        page.add_redact_annot(rect, fill=(0, 0, 0))
        page.apply_redactions()
        if not self._redact_scrub_asked:
            self._redact_scrub_asked = True
            if any(self.doc.metadata.values()) and messagebox.askyesno(
                    "Belge bilgileri", "Başlık, yazar gibi belge bilgileri de silinsin mi?\n"
                    "(Gizlenen içerik bu alanlarda geçiyor olabilir.)", parent=self):
                self.doc.set_metadata({})
        self.render_all_pages()

    def _tool_kirp(self, index, p0, p1):
        rect = fitz.Rect(p0, p1).normalize()
        if rect.width < 20 or rect.height < 20:
            return
        answer = messagebox.askyesnocancel(
            "Sayfayı kırp", "Seçtiğiniz alan dışında kalan kısım gizlenir (kırpma kutusu; dosyadaki içerik silinmez, "
            "'Kırpmayı Kaldır' ile geri gelir).\n\nEvet: yalnızca bu sayfa\nHayır: tüm sayfalar (aynı alan)\n"
            "İptal: vazgeç", parent=self)
        if answer is None:
            return
        indices = [index] if answer else list(range(len(self.doc)))
        self._apply_edit(lambda: PDFProcessor.crop_to_rect(self.doc, indices, rect),
                         f"{len(indices)} sayfa kırpıldı (Ctrl+Z ile geri alınır)")

    def _tool_cizgi(self, index, p0, p1):
        page = self.doc[index]
        self._record_edit()
        page.add_line_annot(p0 * page.derotation_matrix, p1 * page.derotation_matrix).update()
        self.render_all_pages()

    def _tool_not(self, index, point):
        text = simpledialog.askstring("Not", "Not metni:", parent=self)
        if not text:
            return
        page = self.doc[index]
        self._record_edit()
        page.add_text_annot(point * page.derotation_matrix, text).update()
        self.render_all_pages()

    # --- Yazı ekleme: düzenleme doğrudan açık sayfada yapılır; küçük bir biçim paneli yanında durur ---
    def _tool_metin(self, index, point):
        """İlk tıklama yazı düzenlemeyi başlatır; sonraki tıklamalar yazı çerçevesini o noktaya getirir."""
        if getattr(self, "_te", None):
            self._te_move_to(index, point)
            return
        self._te_start(index, point)

    def _te_start(self, index, point):
        t = self.theme
        style0 = dict(getattr(self, "_text_style", None) or default_text_style())
        fonts = text_font_families()
        panel = tk.Toplevel(self)
        panel.title("Yazı biçimi")
        panel.transient(self)
        panel.resizable(False, False)
        te = {"index": index, "pos": [point.x, point.y], "panel": panel, "color": tuple(style0["color"]),
              "job": None, "drag": None, "key": None, "img": None}
        self._te = te
        v = te["vars"] = {
            "family": tk.StringVar(value=style0["family"] if style0["family"] in fonts else fonts[0]),
            "size": tk.StringVar(value="%g" % style0["size"]),
            "bold": tk.BooleanVar(value=style0["bold"]), "italic": tk.BooleanVar(value=style0["italic"]),
            "underline": tk.BooleanVar(value=style0["underline"]), "strike": tk.BooleanVar(value=style0["strike"]),
            "align": tk.StringVar(value=style0["align"]), "opacity": tk.IntVar(value=int(round(style0["opacity"] * 100))),
            "spacing": tk.StringVar(value="%g" % style0["spacing"]), "angle": tk.StringVar(value="%g" % style0["angle"])}
        body = ttk.Frame(panel, padding=10)
        body.grid(row=0, column=0)
        ttk.Label(body, text="Yazı (sayfada önizlenir)").grid(row=0, column=0, columnspan=4, sticky="w")
        box = tk.Text(body, width=32, height=4, wrap="word", bg=t["entry"], fg=t["text"], relief="flat",
                      highlightthickness=1, highlightbackground=t["card_border"], font=(FONT_UI, 10),
                      insertbackground=t["text"], undo=True)
        box.grid(row=1, column=0, columnspan=4, pady=(2, 6))
        te["box"] = box

        ttk.Label(body, text="Yazı tipi").grid(row=2, column=0, sticky="w", pady=2)
        ttk.Combobox(body, textvariable=v["family"], values=fonts, state="readonly", width=17).grid(
            row=2, column=1, columnspan=3, sticky="w")
        ttk.Label(body, text="Boyut (pt)").grid(row=3, column=0, sticky="w", pady=2)
        ttk.Spinbox(body, from_=4, to=300, increment=1, textvariable=v["size"], width=6).grid(row=3, column=1, sticky="w")
        toggles = ttk.Frame(body)
        toggles.grid(row=3, column=2, columnspan=2, sticky="e")
        for text, key, fnt in (("K", "bold", (FONT_UI, 10, "bold")), ("İ", "italic", (FONT_UI, 10, "italic")),
                               ("A", "underline", (FONT_UI, 10, "underline")), ("Ü", "strike", (FONT_UI, 10, "overstrike"))):
            tk.Checkbutton(toggles, text=text, variable=v[key], indicatoron=False, width=2, font=fnt, bg=t["panel"],
                           fg=t["text"], selectcolor=t["accent"], activebackground=t["accent"], relief="raised",
                           bd=1).pack(side="left", padx=1)
        ttk.Label(body, text="K kalın · İ italik · A altı çizili · Ü üstü çizili", foreground=t["muted"]).grid(
            row=4, column=0, columnspan=4, sticky="w")

        ttk.Label(body, text="Renk").grid(row=5, column=0, sticky="w", pady=(6, 2))
        color_row = ttk.Frame(body)
        color_row.grid(row=5, column=1, columnspan=3, sticky="w", pady=(6, 2))
        swatch = tk.Label(color_row, width=3, relief="solid", bd=1)
        swatch.pack(side="left", padx=(0, 4))

        def set_color(rgb):
            te["color"] = tuple(rgb)
            swatch.config(bg="#%02x%02x%02x" % tuple(int(round(c * 255)) for c in rgb))
            self._te_schedule()

        def pick_color():
            from tkinter import colorchooser
            init = "#%02x%02x%02x" % tuple(int(round(c * 255)) for c in te["color"])
            chosen = colorchooser.askcolor(color=init, parent=panel, title="Yazı rengi")
            if chosen and chosen[0]:
                set_color(tuple(c / 255.0 for c in chosen[0]))

        ttk.Button(color_row, text="Seç…", command=pick_color, width=5).pack(side="left", padx=(0, 4))
        for rgb in ((0.1, 0.1, 0.1), (0.8, 0.0, 0.0), (0.0, 0.2, 0.8), (0.0, 0.5, 0.2), (1.0, 1.0, 1.0)):
            chip = tk.Label(color_row, width=2, relief="solid", bd=1, cursor="hand2",
                            bg="#%02x%02x%02x" % tuple(int(round(c * 255)) for c in rgb))
            chip.pack(side="left", padx=1)
            chip.bind("<Button-1>", lambda _e, c=rgb: set_color(c))

        ttk.Label(body, text="Saydamlık").grid(row=6, column=0, sticky="w", pady=2)
        ttk.Scale(body, from_=10, to=100, variable=v["opacity"], orient="horizontal", length=120,
                  command=lambda _x: self._te_schedule()).grid(row=6, column=1, columnspan=2, sticky="w")
        te["opacity_label"] = ttk.Label(body, text="")
        te["opacity_label"].grid(row=6, column=3, sticky="w")
        ttk.Label(body, text="Hizalama").grid(row=7, column=0, sticky="w", pady=2)
        align_row = ttk.Frame(body)
        align_row.grid(row=7, column=1, columnspan=3, sticky="w")
        for label, value in (("Sol", "left"), ("Orta", "center"), ("Sağ", "right")):
            ttk.Radiobutton(align_row, text=label, value=value, variable=v["align"]).pack(side="left", padx=2)
        ttk.Label(body, text="Satır aralığı").grid(row=8, column=0, sticky="w", pady=2)
        ttk.Spinbox(body, from_=0.8, to=3.0, increment=0.1, textvariable=v["spacing"], width=6).grid(row=8, column=1, sticky="w")
        ttk.Label(body, text="Açı (°)").grid(row=9, column=0, sticky="w", pady=2)
        ttk.Spinbox(body, from_=-180, to=180, increment=5, textvariable=v["angle"], width=6).grid(row=9, column=1, sticky="w")
        ttk.Button(body, text="0°", width=3, command=lambda: v["angle"].set("0")).grid(row=9, column=2, sticky="w")

        te["pos_label"] = ttk.Label(body, text="", foreground=t["muted"])
        te["pos_label"].grid(row=10, column=0, columnspan=4, sticky="w", pady=(6, 0))
        ttk.Label(body, text="Yazıyı sayfada sürükleyin; sayfada başka yere tıklarsanız yazı oraya gelir. "
                             "Ok tuşları 1 nokta, Shift ile 10 nokta kaydırır.",
                  foreground=t["muted"], wraplength=290, justify="left").grid(row=11, column=0, columnspan=4,
                                                                               sticky="w", pady=(2, 8))
        buttons = ttk.Frame(body)
        buttons.grid(row=12, column=0, columnspan=4, sticky="e")
        ttk.Button(buttons, text="✔ Onayla ve ekle", style="Save.TButton", command=self._te_confirm).pack(side="left", padx=4)
        ttk.Button(buttons, text="İptal", command=self._te_cancel).pack(side="left")

        panel.protocol("WM_DELETE_WINDOW", self._te_cancel)
        panel.bind("<Escape>", lambda _e: self._te_cancel())
        panel.bind("<Control-Return>", lambda _e: self._te_confirm())
        for var in v.values():
            var.trace_add("write", lambda *_a: self._te_schedule())
        box.bind("<KeyRelease>", lambda _e: self._te_schedule())
        box.bind("<<Paste>>", lambda _e: panel.after(20, self._te_schedule))
        set_color(te["color"])
        if not getattr(self, "_te_keys_bound", False):  # ok tuşları: yalnız düzenleme sürerken yutulur
            self._te_keys_bound = True
            for key, dx, dy in (("Left", -1, 0), ("Right", 1, 0), ("Up", 0, -1), ("Down", 0, 1)):
                for prefix in ("", "Shift-"):
                    self.canvas.bind(f"<{prefix}{key}>", lambda e, a=dx, b=dy: self._te_key(a, b, e), add="+")
        # paneli pencerenin sağ üstüne koy (sayfayı örtmesin)
        self.update_idletasks()
        panel.update_idletasks()
        x = self.winfo_rootx() + max(self.winfo_width() - panel.winfo_reqwidth() - 24, 0)
        panel.geometry("+%d+%d" % (x, self.winfo_rooty() + 90))
        box.focus_set()
        self._te_draw()

    def _te_style(self):
        te = self._te
        v = te["vars"]

        def num(var, default, lo, hi):
            try:
                return max(lo, min(hi, float(var.get().replace(",", "."))))
            except ValueError:
                return default
        return {"family": v["family"].get(), "size": num(v["size"], 12.0, 4.0, 300.0), "bold": v["bold"].get(),
                "italic": v["italic"].get(), "underline": v["underline"].get(), "strike": v["strike"].get(),
                "color": te["color"], "opacity": max(0.1, min(1.0, v["opacity"].get() / 100.0)),
                "align": v["align"].get(), "spacing": num(v["spacing"], 1.2, 0.8, 3.0),
                "angle": num(v["angle"], 0.0, -180.0, 180.0)}

    def _te_bbox(self):
        te = self._te
        text = te["box"].get("1.0", "end-1c")
        return text_block_bbox(text if text.strip() else " ", self._te_style(), tuple(te["pos"]))

    def _te_clamp(self):
        """Çerçeveyi sayfa sınırı içinde tutar; sınır kutusunu döndürür."""
        te = self._te
        prect = fitz.Rect(self.doc[te["index"]].rect)
        bb = self._te_bbox()
        dx = (max(0.0, prect.x0 - bb.x0) or min(0.0, prect.x1 - bb.x1)) if bb.width <= prect.width else prect.x0 - bb.x0
        dy = (max(0.0, prect.y0 - bb.y0) or min(0.0, prect.y1 - bb.y1)) if bb.height <= prect.height else prect.y0 - bb.y0
        te["pos"][0] += dx
        te["pos"][1] += dy
        return fitz.Rect(bb.x0 + dx, bb.y0 + dy, bb.x1 + dx, bb.y1 + dy)

    def _te_schedule(self):
        te = getattr(self, "_te", None)
        if te and te["job"] is None:
            te["job"] = self.after(50, self._te_draw)

    def _te_draw(self):
        """Yazıyı ve kesikli çerçeveyi ana sayfa görünümünün üstüne çizer (sayfa yeniden çizilince de çağrılır)."""
        te = getattr(self, "_te", None)
        if not te:
            return
        te["job"] = None
        c = self.canvas
        c.delete("te")
        index = te["index"]
        if index >= len(self.doc) or index >= len(self.page_positions) or index >= len(self.page_x):
            return
        page = self.doc[index]
        text = te["box"].get("1.0", "end-1c")
        style = self._te_style()
        bb = self._te_clamp()
        matrix = self._render_matrix(page, self.zoom)
        origin = page.rect * matrix
        px, py = self.page_x[index], self.page_positions[index][0]

        def to_canvas(rect):
            r = fitz.Rect(rect) * matrix
            r.normalize()
            return px + r.x0 - origin.x0, py + r.y0 - origin.y0, px + r.x1 - origin.x0, py + r.y1 - origin.y0

        te["pos_label"].config(text="Sol üst köşe: x = %.1f, y = %.1f nokta (sayfa %d)" % (te["pos"][0], te["pos"][1], index + 1))
        te["opacity_label"].config(text="%d%%" % te["vars"]["opacity"].get())
        margin = 4.0
        if text.strip():
            key = (text, tuple(sorted((k, str(val)) for k, val in style.items())), round(self.zoom, 4), index,
                   self._page_is_rotated_for_view(page))
            if key != te["key"]:
                tmp = fitz.open()
                try:
                    tpage = tmp.new_page(width=bb.width + 2 * margin, height=bb.height + 2 * margin)
                    draw_text_block(tpage, (margin - (bb.x0 - te["pos"][0]), margin - (bb.y0 - te["pos"][1])), text, style)
                    pix = tpage.get_pixmap(matrix=matrix, alpha=True)
                    te["img"] = tk.PhotoImage(data=pix.tobytes("png"))
                    te["key"] = key
                finally:
                    tmp.close()
            x0, y0, _x1, _y1 = to_canvas(fitz.Rect(bb.x0 - margin, bb.y0 - margin, bb.x1 + margin, bb.y1 + margin))
            c.create_image(x0, y0, anchor="nw", image=te["img"], tags="te")
        fx0, fy0, fx1, fy1 = to_canvas(bb)
        c.create_rectangle(fx0 - 2, fy0 - 2, fx1 + 2, fy1 + 2, outline=self.theme["accent_dark"], width=2,
                           dash=(5, 3), tags="te")
        c.tag_raise("te")

    def _te_move_to(self, index, point):
        te = self._te
        te["index"] = index
        bb = self._te_bbox()
        te["pos"][0] = point.x - bb.width / 2 + (te["pos"][0] - bb.x0)
        te["pos"][1] = point.y - bb.height / 2 + (te["pos"][1] - bb.y0)
        self._te_draw()

    def _te_press(self, cx, cy):
        """Düzenleme sürerken tuval tıklaması: çerçevenin içindeyse sürükleme başlar, dışındaysa çerçeve oraya gelir."""
        te = self._te
        index, point = self._view_to_page(cx, cy)
        if index is None:
            return
        if index != te["index"] or not self._te_bbox().contains(point):
            self._te_move_to(index, point)
        te["drag"] = (self._view_to_page(cx, cy, index=te["index"])[1], list(te["pos"]))

    def _te_motion(self, cx, cy):
        te = self._te
        start, pos0 = te["drag"]
        now = self._view_to_page(cx, cy, index=te["index"])[1]
        te["pos"][0] = pos0[0] + now.x - start.x
        te["pos"][1] = pos0[1] + now.y - start.y
        self._te_draw()

    def _te_key(self, dx, dy, event):
        if not getattr(self, "_te", None):
            return None
        step = 10.0 if event.state & 1 else 1.0
        self._te["pos"][0] += dx * step
        self._te["pos"][1] += dy * step
        self._te_draw()
        return "break"

    def _te_close(self):
        te, self._te = getattr(self, "_te", None), None
        if te:
            if te["job"] is not None:
                try:
                    self.after_cancel(te["job"])
                except (ValueError, tk.TclError):
                    pass
            self.canvas.delete("te")
            try:
                te["panel"].destroy()
            except tk.TclError:
                pass

    def _te_cancel(self):
        self._te_close()

    def _te_confirm(self):
        te = self._te
        if not te:
            return
        text = te["box"].get("1.0", "end-1c").rstrip()
        if not text.strip():
            messagebox.showwarning("Yazı ekle", "Önce eklenecek yazıyı girin.", parent=te["panel"])
            return
        if te["index"] >= len(self.doc):
            self._te_close()
            return
        self._te_clamp()
        style, pos, index = self._te_style(), tuple(te["pos"]), te["index"]
        self._text_style = dict(style)
        self._te_close()
        page = self.doc[index]
        self._record_edit()
        draw_text_block(page, fitz.Point(pos), text, style)
        self.render_all_pages()
        self._toast("Yazı eklendi (Ctrl+Z ile geri alınır)")

    def _tool_gorsel(self, index, point):
        path = filedialog.askopenfilename(title="İmza veya görsel seçin",
                                          filetypes=[("Görseller", "*.png *.jpg *.jpeg")], parent=self)
        if not path:
            return
        page = self.doc[index]
        rect = (fitz.Rect(point.x - 80, point.y - 40, point.x + 80, point.y + 40) * page.derotation_matrix).normalize()
        rect = rect & page.mediabox
        if rect.is_empty:
            return
        self._record_edit()
        page.insert_image(rect, filename=path, keep_proportion=True, overlay=True, rotate=page.rotation)
        self.render_all_pages()

    def _text_dialog(self, title, prompt, size=12):
        """Çok satırlı metin ve yazı boyutu sorar; (metin, boyut) veya None döndürür."""
        t = self.theme
        win = tk.Toplevel(self)
        win.title(title)
        win.transient(self)
        win.resizable(False, False)
        ttk.Label(win, text=prompt).grid(row=0, column=0, columnspan=2, sticky="w", padx=12, pady=(12, 4))
        box = tk.Text(win, width=44, height=4, wrap="word", bg=t["entry"], fg=t["text"], relief="flat",
                      highlightthickness=1, highlightbackground=t["card_border"], font=(FONT_UI, 10),
                      insertbackground=t["text"])
        box.grid(row=1, column=0, columnspan=2, padx=12, pady=4)
        size_var = tk.StringVar(value=str(size))
        ttk.Label(win, text="Yazı boyutu").grid(row=2, column=0, sticky="w", padx=12, pady=6)
        ttk.Spinbox(win, from_=6, to=96, textvariable=size_var, width=6).grid(row=2, column=1, sticky="w", pady=6)
        result = {}

        def ok():
            text = box.get("1.0", "end").strip()
            try:
                value = max(6.0, min(96.0, float(size_var.get().replace(",", "."))))
            except ValueError:
                value = float(size)
            if text:
                result["v"] = (text, value)
            win.destroy()

        row = ttk.Frame(win)
        row.grid(row=3, column=0, columnspan=2, sticky="e", padx=12, pady=(4, 12))
        ttk.Button(row, text="Ekle", style="Save.TButton", command=ok).pack(side="left", padx=4)
        ttk.Button(row, text="İptal", command=win.destroy).pack(side="left")
        box.focus_set()
        win.grab_set()
        self.wait_window(win)
        return result.get("v")

    def _form_dialog(self, title, fields, ok_text="Uygula"):
        """fields: (anahtar, etiket, varsayılan, tür[str|int|float]) listesi. Değer sözlüğü veya None döndürür."""
        win = tk.Toplevel(self)
        win.title(title)
        win.transient(self)
        win.resizable(False, False)
        variables = {}
        for row, (key, label, default, _kind) in enumerate(fields):
            ttk.Label(win, text=label).grid(row=row, column=0, sticky="w", padx=12, pady=5)
            variables[key] = tk.StringVar(value=str(default))
            ttk.Entry(win, textvariable=variables[key], width=34).grid(row=row, column=1, padx=12, pady=5)
        result = {}

        def ok():
            try:
                values = {}
                for key, _label, _default, kind in fields:
                    raw = variables[key].get().strip()
                    values[key] = kind(raw.replace(",", ".")) if kind in (int, float) else raw
            except ValueError:
                messagebox.showerror("Geçersiz değer", "Sayısal alanlara sayı girin.", parent=win)
                return
            result["v"] = values
            win.destroy()

        row = ttk.Frame(win)
        row.grid(row=len(fields), column=0, columnspan=2, sticky="e", padx=12, pady=(6, 12))
        ttk.Button(row, text=ok_text, style="Save.TButton", command=ok).pack(side="left", padx=4)
        ttk.Button(row, text="İptal", command=win.destroy).pack(side="left")
        win.bind("<Return>", lambda _e: ok())
        win.grab_set()
        self.wait_window(win)
        return result.get("v")

    def toggle_page_selection(self, index):
        if index in self.selected_pages:
            self.selected_pages.remove(index)
        else:
            self.selected_pages.add(index)
        self.render_thumbnails()

    def action_delete_selected_pages(self):
        if not self.selected_pages:
            return
        if len(self.selected_pages) >= len(self.doc):
            messagebox.showwarning("Sayfa silme", "Belgede en az bir sayfa kalmalıdır.", parent=self)
            return
        if not messagebox.askyesno("Sayfaları sil", f"{len(self.selected_pages)} sayfa silinsin mi?", parent=self):
            return
        self._record_edit()
        for index in sorted(self.selected_pages, reverse=True):
            self.doc.delete_page(index)
        self.selected_pages.clear()
        self.render_thumbnails()
        self.render_all_pages()

    def action_export_selected_pages(self):
        if not self.selected_pages:
            return
        path = filedialog.asksaveasfilename(defaultextension=".pdf", filetypes=[("PDF", "*.pdf")])
        if path:
            self._export_from_working(
                lambda src: PDFProcessor.extract_pages(src, [i + 1 for i in sorted(self.selected_pages)], path),
                "Seçilen sayfalar dışa aktarıldı.")



    @staticmethod
    def reorder_indices(count, moved, slot):
        """`moved` indekslerini (sıralı) `slot` konumunun önüne taşıyan yeni sayfa sırasını verir: new[k] = eski indeks.
        slot 0..count: o sayfanın önüne ekle (count = sona). Taşınanların kendi aralarındaki sırası korunur."""
        moved = sorted(moved)
        rest = [i for i in range(count) if i not in moved]
        pos = sum(1 for i in rest if i < slot)
        return rest[:pos] + moved + rest[pos:]

    def _thumb_slot(self, y_root):
        """İmlecin altındaki ekleme aralığı (0..sayfa sayısı) ve canvas y'si; kartların orta çizgisine göre."""
        y = self.thumb_canvas.canvasy(y_root - self.thumb_canvas.winfo_rooty())
        for i, card in enumerate(self.thumb_cards):
            if y < card.winfo_y() + card.winfo_height() / 2:
                return i
        return len(self.thumb_cards)

    def _thumb_moved(self, src_idx):
        """Sürüklenen sayfa seçili sayfalardan biriyse seçilenlerin hepsi birlikte taşınır."""
        if src_idx in self.selected_pages and len(self.selected_pages) > 1:
            return sorted(self.selected_pages)
        return [src_idx]

    def _show_thumb_marker(self, slot):
        if not self.thumb_cards:
            return
        marker = getattr(self, "_thumb_marker", None)
        if marker is None or not marker.winfo_exists():
            marker = tk.Frame(self.thumb_container, bg=self.theme["accent_dark"], height=4)
            self._thumb_marker = marker
        if slot < len(self.thumb_cards):
            y = self.thumb_cards[slot].winfo_y() - 4
        else:
            last = self.thumb_cards[-1]
            y = last.winfo_y() + last.winfo_height() + 1
        marker.place(x=4, y=max(y, 0), width=max(self.thumb_container.winfo_width() - 8, 20), height=4)
        marker.lift()

    def _hide_thumb_marker(self):
        marker = getattr(self, "_thumb_marker", None)
        if marker is not None and marker.winfo_exists():
            marker.place_forget()

    def _thumb_autoscroll(self):
        """Sürüklerken imleç panelin üst/alt kenarına yaklaşınca listeyi kaydırır."""
        self._thumb_scroll_job = None
        if self.drag_data["index"] is None:
            return
        top = self.thumb_canvas.winfo_rooty()
        y = self._thumb_pointer_y
        step = -1 if y < top + 30 else 1 if y > top + self.thumb_canvas.winfo_height() - 30 else 0
        if step:
            self.thumb_canvas.yview_scroll(step, "units")
            self._show_thumb_marker(self._thumb_slot(y))
        self._thumb_scroll_job = self.after(40, self._thumb_autoscroll)

    def on_thumb_drag(self, event, idx):
        self.drag_data["index"] = idx
        self._thumb_pointer_y = event.y_root
        self._show_thumb_marker(self._thumb_slot(event.y_root))
        self.thumb_canvas.config(cursor="fleur")
        if getattr(self, "_thumb_scroll_job", None) is None:
            self._thumb_scroll_job = self.after(40, self._thumb_autoscroll)

    def on_thumb_drop(self, event, src_idx):
        self._hide_thumb_marker()
        self.thumb_canvas.config(cursor="")
        if self.drag_data["index"] is None or not self.thumb_cards:  # sürükleme değil, düz tıklama
            return
        self.drag_data["index"] = None
        count = len(self.doc)
        moved = self._thumb_moved(src_idx)
        order = self.reorder_indices(count, moved, self._thumb_slot(event.y_root))
        if order == list(range(count)):
            return
        self._record_edit()
        current = list(range(count))
        for k, page_no in enumerate(order):  # her sayfayı sırayla yerine getir (move_page hedefin önüne koyar)
            at = current.index(page_no)
            if at != k:
                self.doc.move_page(at, k)
                current.insert(k, current.pop(at))
        new_pos = {old: k for k, old in enumerate(order)}
        self.selected_pages = {new_pos[i] for i in self.selected_pages if i in new_pos}
        self.search_hits = {}
        self.render_thumbnails()
        self.render_all_pages()
        self.page_var.set(str(new_pos[src_idx] + 1))
        self.goto_page()

    # ------------------------------------------------------------------
    # Bağlantılar, gezinti geçmişi, içindekiler paneli, katmanlar
    # ------------------------------------------------------------------
    def _tool_cursor(self):
        cursors = {"gez": "fleur", "kopyala": "xterm", "not": "plus", "metin": "xterm", "gorsel": "plus"}
        return cursors.get(self.tool_mode, "crosshair")

    def _link_under(self, cx, cy):
        """Canvas noktasındaki (sayfa indeksi, bağlantı); yoksa (None, None). Sayfa bağlantıları önbelleğe alınır."""
        index, point = self._view_to_page(cx, cy)
        if index is None:
            return None, None
        links = self._link_cache.get(index)
        if links is None:
            links = self._link_cache[index] = self.doc[index].get_links()
        return index, find_link_at(self.doc[index], point, links)

    def on_canvas_motion(self, event):
        if self.is_welcome or self.tool_mode != "gez" or self._drag_start is not None or not self.page_positions:
            self._set_link_hover(None, event)
            return
        _index, link = self._link_under(self.canvas.canvasx(event.x), self.canvas.canvasy(event.y))
        self._set_link_hover(describe_link(link, len(self.doc)) if link else None, event)

    def _set_link_hover(self, text, event):
        """Bağlantı üzerinde el imleci ve hedefi gösteren küçük balon; bağlantı yoksa ikisini de kaldırır."""
        if text is None:
            if self._link_tip is not None and self._link_tip.winfo_ismapped():
                self._link_tip.place_forget()
                self.canvas.configure(cursor=self._tool_cursor())
            return
        t = self.theme
        if self._link_tip is None:
            self._link_tip = tk.Label(self.canvas, bg=t["panel"], fg=t["text"], relief="solid", bd=1,
                                      font=(FONT_UI, 9), padx=6, pady=2)
        if self._link_tip.cget("text") != text:
            self._link_tip.config(text=text)
        self._link_tip.place(x=event.x + 14, y=event.y + 20)
        self._link_tip.lift()
        self.canvas.configure(cursor="hand2")

    def _scroll_fraction(self):
        return self.canvas.yview()[0]

    def _nav_record(self):
        self._nav_back.append(self._scroll_fraction())
        del self._nav_back[:-50]
        self._nav_fwd.clear()

    def _nav_move(self, source, target):
        if self.is_welcome or not source:
            return
        target.append(self._scroll_fraction())
        self.canvas.yview_moveto(source.pop())

    def nav_back(self):
        self._nav_move(self._nav_back, self._nav_fwd)

    def nav_forward(self):
        self._nav_move(self._nav_fwd, self._nav_back)

    def _jump_to(self, index, to=None):
        """Sayfaya (varsa sayfa içindeki hedef noktaya) gider ve önceki konumu geçmişe yazar."""
        if not self.page_positions:
            return
        index = min(max(int(index), 0), len(self.page_positions) - 1)
        top, bottom = self.page_positions[index]
        y = top
        if to is not None:
            try:
                page = self.doc[index]
                matrix = self._render_matrix(page, self.zoom)
                point = fitz.Point(to) * page.rotation_matrix * matrix
                if math.isfinite(point.y):
                    y = top + min(max(point.y - (page.rect * matrix).y0, 0), bottom - top)
            except Exception:  # noqa: BLE001 - hedef noktası bozuksa sayfa başına gider
                y = top
        self._nav_record()
        self.canvas.yview_moveto(max(y - 10, 0) / max(self.total_height, 1))

    def _follow_link(self, link):
        kind = link.get("kind")
        if kind == fitz.LINK_GOTO:
            if link.get("page", -1) >= 0:
                self._jump_to(link["page"], link.get("to"))
        elif kind == fitz.LINK_NAMED:
            current = self.get_current_page_index()
            target = {"NextPage": current + 1, "PrevPage": current - 1, "FirstPage": 0,
                      "LastPage": len(self.doc) - 1}.get(link.get("name"))
            if target is not None:
                self._jump_to(target)
        elif kind == fitz.LINK_URI:
            self._open_external_link(str(link.get("uri", "")))
        else:
            messagebox.showinfo("Bağlantı", "Başka bir dosyayı veya programı açan bağlantılar güvenlik nedeniyle açılmaz.",
                                parent=self)

    def _open_external_link(self, uri):
        """Dış bağlantıyı yalnızca http/https/mailto ise ve kullanıcı onaylarsa tarayıcıda açar."""
        if not uri.lower().startswith(SAFE_LINK_SCHEMES):
            messagebox.showinfo("Bağlantı", f"Bu bağlantı türü açılmaz:\n{uri[:200]}", parent=self)
            return
        shown = uri if len(uri) <= 200 else uri[:200] + "…"
        if messagebox.askyesno("Dış bağlantı", f"Bu belge şu adresi açmak istiyor:\n\n{shown}\n\n"
                               "Tarayıcıda açılsın mı? (Yalnızca güvendiğiniz belgelerde açın.)",
                               icon="warning", parent=self):
            import webbrowser
            webbrowser.open(uri)

    def toggle_toc_panel(self):
        if self.is_welcome:
            return
        if self.is_toc_open:
            self.paned.forget(self.toc_frame)
            self.is_toc_open = False
            return
        self.paned.insert(0, self.toc_frame, weight=1)
        self.is_toc_open = True
        self._toc_cache = None
        self._refresh_toc()

    def _refresh_toc(self):
        """İçindekiler ağacını belgenin yer imlerinden kurar; yer imleri değişmediyse dokunmaz."""
        if not self.is_toc_open:
            return
        try:
            toc = self.doc.get_toc(simple=False)
        except Exception:  # noqa: BLE001
            toc = []
        key = [(entry[0], entry[1], entry[2]) for entry in toc]
        if key == self._toc_cache:
            return
        self._toc_cache = key
        t = self.theme
        style = ttk.Style(self)
        style.configure("Toc.Treeview", background=t["lavender"], fieldbackground=t["lavender"], foreground=t["text"],
                        rowheight=24)
        tree = self.toc_tree
        tree.delete(*tree.get_children())
        self._toc_targets = {}
        stack = []  # (düzey, iid)
        for entry in toc:
            level, title, page = entry[0], entry[1], entry[2]
            dest = entry[3] if len(entry) > 3 and isinstance(entry[3], dict) else {}
            while stack and stack[-1][0] >= level:
                stack.pop()
            iid = tree.insert(stack[-1][1] if stack else "", "end", text=(str(title).strip() or "(adsız)"),
                              open=len(toc) <= 60 or level < 2)
            stack.append((level, iid))
            self._toc_targets[iid] = (page - 1 if page >= 1 else -1, dest.get("to"))
        if toc:
            self.toc_note.pack_forget()
        else:
            self.toc_note.pack(side="top", fill="x", padx=8, pady=4)

    def _on_toc_click(self, event):
        tree = self.toc_tree
        iid = tree.identify_row(event.y) if getattr(event, "num", None) == 1 else tree.focus()
        target = self._toc_targets.get(iid)
        if target and target[0] >= 0:
            self._jump_to(target[0], target[1])

    def action_layers(self):
        """PDF katmanlarını (isteğe bağlı içerik grupları) açıp kapatan pencere."""
        if self.is_welcome:
            return
        try:
            layers = [item for item in self.doc.layer_ui_configs() if item.get("text")]
        except Exception:  # noqa: BLE001 - PDF olmayan belgelerde katman yok
            layers = []
        if not layers:
            messagebox.showinfo("Katmanlar", "Bu belgede katman yok.", parent=self)
            return
        win = tk.Toplevel(self)
        win.title("Katmanlar")
        win.transient(self)
        body = ttk.Frame(win, padding=12)
        body.pack(fill="both", expand=True)
        ttk.Label(body, text="Göstermek istediğiniz katmanları işaretleyin:").pack(anchor="w", pady=(0, 6))
        variables = []

        def toggle(number, var):
            self.doc.set_layer_ui_config(number, 0 if var.get() else 1)
            self._link_cache.clear()
            self.render_all_pages()
            self.render_thumbnails()

        for item in layers:
            var = tk.BooleanVar(value=bool(item.get("on")))
            variables.append(var)
            check = ttk.Checkbutton(body, text=item["text"], variable=var,
                                    command=lambda n=item["number"], v=var: toggle(n, v))
            if item.get("locked") or item.get("type") == "label":
                check.state(["disabled"])
            check.pack(anchor="w", padx=(12 * int(item.get("depth", 0)), 0))
        ttk.Button(body, text="Kapat", command=win.destroy).pack(anchor="e", pady=(10, 0))

    # ------------------------------------------------------------------
    # Sayfa çizimi (yalnızca görünen sayfalar çizilir)
    # ------------------------------------------------------------------
    def render_all_pages(self):
        if self._te and self._te["index"] >= len(self.doc):
            self._te_close()
        self._link_cache.clear()
        self.canvas.delete("all")
        self.page_images.clear()
        self.page_positions.clear()
        self.page_x.clear()

        y_offset = 20
        canvas_width = self.canvas.winfo_width()
        max_w = 0

        for i, page in enumerate(self.doc):
            w, h = self._display_page_size(page)
            w, h = int(math.ceil(w * self.zoom)), int(math.ceil(h * self.zoom))
            x_offset = max((canvas_width - w) // 2, 20)
            self.page_x.append(x_offset)
            self.page_positions.append((y_offset, y_offset + h))
            # Sayfa yer tutucusu (resim görünür olunca üstüne çizilir)
            self.canvas.create_rectangle(x_offset + 3, y_offset + 3, x_offset + w + 3, y_offset + h + 3,
                                         fill=self.theme["shadow"], outline="")
            self.canvas.create_rectangle(x_offset, y_offset, x_offset + w, y_offset + h,
                                         fill="#ffffff", outline="")
            y_offset += h + 30
            max_w = max(max_w, w)

        self.total_height = max(y_offset, 1)
        self.canvas.configure(scrollregion=(0, 0, max(max_w + 40, canvas_width), self.total_height))
        self.zoom_label.config(text=f"{int(self.zoom * 100)}%")
        self.total_label.config(text=f"/ {len(self.doc)}")
        self._update_search_label()
        self._refresh_toc()
        self._render_visible()

    def _schedule_visible(self):
        if self._visible_job is None:
            self._visible_job = self.after_idle(self._render_visible)

    def _render_visible(self):
        self._visible_job = None
        if not self.page_positions:
            return
        top = self.canvas.canvasy(0)
        height = max(self.canvas.winfo_height(), 1)
        bottom = top + height
        margin = height  # görünür alanın bir ekran altı/üstü önceden çizilir
        keep = set()
        for i, (page_top, page_bottom) in enumerate(self.page_positions):
            if page_bottom >= top - margin and page_top <= bottom + margin:
                keep.add(i)
            elif page_top > bottom + margin:
                break

        for i in list(self.page_images):
            if i not in keep:
                self.canvas.delete(f"page_{i}")
                del self.page_images[i]

        for i in sorted(keep):
            if i in self.page_images or i >= len(self.doc):
                continue
            page = self.doc[i]
            matrix = self._render_matrix(page, self.zoom)
            pix = tint_pixmap(page.get_pixmap(matrix=matrix, alpha=False), self.page_mode_var.get())
            img = tk.PhotoImage(data=pix.tobytes("ppm"))  # PNG'ye göre ~3 kat hızlı (kodlama adımı yok)
            self.page_images[i] = img
            x, y = self.page_x[i], self.page_positions[i][0]
            self.canvas.create_image(x, y, anchor="nw", image=img, tags=f"page_{i}")
            self._draw_hits(i, page, matrix, x, y)
            self._draw_widget_hits(i, page, matrix, x, y)
        if self._te:
            self._te_draw()

    def _draw_hits(self, index, page, matrix, x, y):
        """Arama eşleşmelerini belgeyi değiştirmeden canvas üzerinde işaretler."""
        rects = self.search_hits.get(index)
        if not rects:
            return
        origin = page.rect * matrix
        hits = self._flat_hits()
        current = hits[self._hit_index] if hits and self._hit_index < len(hits) else None
        for rect in rects:
            t = rect * page.rotation_matrix * matrix
            is_current = current is not None and current[0] == index and current[1] == rect
            self.canvas.create_rectangle(
                x + t.x0 - origin.x0, y + t.y0 - origin.y0,
                x + t.x1 - origin.x0, y + t.y1 - origin.y0,
                outline="#ef4444" if is_current else "#ffbf00", width=4 if is_current else 2,
                tags=f"page_{index}",
            )

    def _draw_widget_hits(self, index, page, matrix, x, y):
        rects = self.widget_hits.get(index, [])
        origin = page.rect * matrix
        for rect in rects:
            scaled = rect * page.rotation_matrix * matrix
            self.canvas.create_rectangle(
                x + scaled.x0 - origin.x0, y + scaled.y0 - origin.y0,
                x + scaled.x1 - origin.x0, y + scaled.y1 - origin.y0,
                outline="#16a34a", width=3, tags=f"page_{index}")
    def on_ctrl_wheel(self, e):
        factor = 1.15 if e.delta > 0 else 1 / 1.15
        self.zoom = min(max(self.zoom * factor, ZOOM_MIN), ZOOM_MAX)
        self.zoom_label.config(text=f"{int(self.zoom * 100)}%")
        if self._zoom_timer is not None:
            self.after_cancel(self._zoom_timer)
        self._zoom_timer = self.after(150, self.render_all_pages)
        return "break"

    def get_current_page_index(self):
        if not self.page_positions:
            return 0
        scroll_y = self.canvas.canvasy(0)
        for idx, (top, bottom) in enumerate(self.page_positions):
            if top <= scroll_y <= bottom or scroll_y < top:
                return idx
        return len(self.doc) - 1

    def on_scroll(self, first, last):
        self.vbar.set(first, last)  # kaydırma çubuğunu güncel tut
        if self.focus_get() is not self.page_entry:  # kullanıcı sayfa no yazarken ezme
            self.page_var.set(str(self.get_current_page_index() + 1))
        self._schedule_visible()

    def _step_page(self, delta, force=False):
        if not force and self.focus_get() is self.page_entry:
            return
        self.page_var.set(str(min(max(self.get_current_page_index() + 1 + delta, 1), len(self.doc))))
        self.goto_page()

    def search_step(self, direction):
        """Eşleşmeler arasında tek tek gezer (F3 / Shift+F3, ▲ ▼ düğmeleri); sondan başa sarar."""
        if not self._flat_hits():
            return
        self._goto_hit(self._hit_index + (1 if direction > 0 else -1))

    def set_zoom(self, z):
        self.zoom = min(max(z, ZOOM_MIN), ZOOM_MAX)
        self.render_all_pages()

    def fit_width(self):
        self.update_idletasks()
        if len(self.doc) == 0:
            return
        page_w = self._display_page_size(self.doc[0])[0]
        avail_w = max(self.canvas.winfo_width() - 60, 300)
        self.set_zoom(avail_w / page_w)

    def goto_page(self, _event=None):
        try:
            n = int(self.page_var.get()) - 1
        except ValueError:
            return
        if not self.page_positions:
            return
        n = min(max(n, 0), len(self.page_positions) - 1)
        target_y = self.page_positions[n][0]
        self.canvas.yview_moveto(target_y / max(self.total_height, 1))

    # ------------------------------------------------------------------
    # Düzenleme
    # ------------------------------------------------------------------
    def rotate_current_page(self):
        idx = self.get_current_page_index()
        if self.page_orientation:  # yön zorlaması açıkken döndürme görünmez kalırdı
            self.page_orientation = None
            self.portrait_var.set(False)
            self.landscape_var.set(False)
        page = self.doc[idx]
        self._record_edit()
        page.set_rotation((page.rotation + 90) % 360)
        self.search_hits = {}
        self.render_thumbnails()
        self.render_all_pages()

    def delete_current_page(self):
        if len(self.doc) <= 1:
            messagebox.showwarning("Uyarı", "PDF içerisindeki tek sayfa silinemez!")
            return
        idx = self.get_current_page_index()
        if messagebox.askyesno("Sayfa Sil", f"{idx + 1}. sayfayı silmek istediğinize emin misiniz?", parent=self):
            self._record_edit()
            self.doc.delete_page(idx)
            self.search_hits = {}
            self.render_thumbnails()
            self.render_all_pages()

    def save_pdf(self):
        """Ctrl+S. İlk kayıtta orijinalin üzerine yazma / numaralı kopya seçimi sorulur; seçim oturum boyunca geçerlidir."""
        if self.is_welcome:
            self.save_pdf_as()
            return
        if self.doc.is_encrypted or self._save_mode == "copy":
            self.save_pdf_copy()
            return
        if self._save_mode is None:
            answer = messagebox.askyesnocancel(
                "Kaydet",
                "Orijinal dosyanın üzerine yazılsın mı?\n\n"
                "Evet: üzerine yaz (ilk halin bir kopyası 'ad.yedek.pdf' olarak saklanır)\n"
                "Hayır: her kayıtta yeni kopya oluştur (_D1, _D2…)\n"
                "İptal: vazgeç", icon="question", parent=self)
            if answer is None:
                return
            self._save_mode = "overwrite" if answer else "copy"
            if not answer:
                self.save_pdf_copy()
                return
        self._overwrite_original()

    def _overwrite_original(self):
        """Düzenlenmiş belgeyi güvenli biçimde (geçici dosya + atomik değiştirme) orijinalin üzerine yazar."""
        path = self.pdf_path
        tmp = path + ".tmp"
        try:
            data = self.doc.tobytes(garbage=3, deflate=True)
            with open(tmp, "wb") as handle:
                handle.write(data)
            backup = os.path.splitext(path)[0] + ".yedek.pdf"
            if not os.path.exists(backup):
                shutil.copy2(path, backup)
            self.doc.close()
            try:
                for attempt in range(6):  # sürücü eşitleme/antivirüs taraması dosyayı anlık kilitleyebilir
                    try:
                        os.replace(tmp, path)
                        break
                    except PermissionError:
                        if attempt == 5:
                            raise
                        time.sleep(0.25)
                self.doc = fitz.open(path)
            except Exception:
                self.doc = fitz.open(stream=data, filetype="pdf")  # yazılamadı: düzenlemeler bellekte kalsın
                raise
        except Exception as exc:  # noqa: BLE001
            self._remove_quietly(tmp)
            log_exception("Kaydetme", exc)
            messagebox.showerror("Hata", f"Dosya kaydedilemedi (düzenlemeler bellekte duruyor):\n{exc}", parent=self)
            return
        self.is_edited = False
        self._clear_recovery()
        self.render_thumbnails()
        self.render_all_pages()
        self._toast(f"Kaydedildi: {os.path.basename(path)}", path)

    def save_pdf_copy(self):
        """Orijinal dosyayı bozmadan D1, D2 vb. adıyla aynı klasöre kaydeder."""
        if self.is_welcome:
            self.save_pdf_as()
            return
        save_path = generate_edited_filepath(self.pdf_path)
        try:
            self.doc.save(save_path)
            self.is_edited = False
            self._clear_recovery()
            self._toast(f"Kopya kaydedildi: {os.path.basename(save_path)}", save_path)
        except Exception as e:
            messagebox.showerror("Hata", f"Dosya kaydedilemedi:\n{e}", parent=self)

    def save_pdf_as(self):
        """Düzenlenmiş PDF'yi kullanıcının seçtiği ad ve konuma kaydeder."""
        save_path = filedialog.asksaveasfilename(
            title="PDF'yi Farklı Kaydet",
            defaultextension=".pdf",
            initialfile=os.path.basename(self.pdf_path),
            filetypes=[("PDF Dosyası", "*.pdf")],
        )
        if not save_path:
            return
        if self._is_same_file(save_path):
            messagebox.showwarning("Uyarı", "Açık olan dosyanın üzerine yazılamaz. Farklı bir ad seçin.", parent=self)
            return
        try:
            self.doc.save(save_path)
            self.is_edited = False
            self._clear_recovery()
            self._toast(f"Kaydedildi: {os.path.basename(save_path)}", save_path)
        except Exception as e:
            messagebox.showerror("Hata", f"Dosya kaydedilemedi:\n{e}")

    def action_open_new(self):
        file_path = filedialog.askopenfilename(filetypes=OPEN_FILETYPES)
        if file_path:
            try:
                self.open_pdf_path(file_path)
            except Exception as exc:
                messagebox.showerror("PDF Açma Hatası", f"PDF açılamadı:\n{exc}", parent=self)

    def action_open_book(self, kind=None):
        """E-kitap seçtirip açar; kind verilirse (EPUB, MOBI, AZW3, IBA, RTF) yalnızca o biçim listelenir."""
        if kind in BOOK_PATTERNS:
            title, types = f"{kind} dosyası seçin", [(f"{kind} dosyaları", BOOK_PATTERNS[kind])]
        else:
            title, types = "E-kitap seçin", BOOK_FILETYPES
        file_path = filedialog.askopenfilename(title=title, filetypes=types, parent=self)
        if file_path:
            self._open_then_path(file_path)

    def action_open_epub(self):
        self.action_open_book("EPUB")

    def action_extract_text(self):
        save_path = filedialog.asksaveasfilename(defaultextension=".txt", filetypes=[("Metin Dosyası", "*.txt")])
        if not save_path:
            return
        source, is_tmp = None, False
        try:
            source, is_tmp = self._working_pdf_path()
            PDFProcessor(source).extract_text(save_path)
            self._toast(f"Metin kaydedildi: {os.path.basename(save_path)}", save_path)
        except Exception as e:
            messagebox.showerror("Hata", f"Bir hata oluştu:\n{e}")
        finally:
            if is_tmp and source:
                self._remove_quietly(source)

    def action_convert_images(self):
        folder_path = filedialog.askdirectory(title="Görsellerin Kaydedileceği Klasörü Seçin")
        if not folder_path:
            return
        source, is_tmp = None, False
        try:
            source, is_tmp = self._working_pdf_path()
            PDFProcessor(source).convert_to_images(folder_path)
            self._toast("Sayfalar resim olarak kaydedildi", os.path.join(folder_path, "sayfa_1.png"))
        except Exception as e:
            messagebox.showerror("Hata", f"Bir hata oluştu:\n{e}")
        finally:
            if is_tmp and source:
                self._remove_quietly(source)

    def action_merge(self):
        files = filedialog.askopenfilenames(title="Birleştirilecek PDF'leri Seçin", filetypes=[("PDF Dosyaları", "*.pdf")])
        if files:
            save_path = filedialog.asksaveasfilename(defaultextension=".pdf", filetypes=[("PDF Dosyası", "*.pdf")])
            if save_path:
                try:
                    PDFProcessor.merge_pdfs(list(files), save_path)
                    self._toast(f"Birleştirildi: {os.path.basename(save_path)}", save_path)
                except Exception as e:
                    messagebox.showerror("Hata", f"Bir hata oluştu:\n{e}")

    def action_images_to_pdf(self, initial=None):
        """Resimleri sırayla tek PDF'e çevirir; sıra, sayfa boyutu ve çıktı adı küçük bir pencerede seçilir."""
        exts = " ".join("*" + e for e in PDFProcessor.IMAGE_EXTENSIONS)
        files = list(initial or filedialog.askopenfilenames(
            title="PDF'e çevrilecek resimleri seçin", filetypes=[("Resimler", exts), ("Tüm dosyalar", "*.*")], parent=self))
        if not files:
            return
        win = tk.Toplevel(self)
        win.title("Resimden PDF")
        win.transient(self)
        win.resizable(False, True)
        outer = ttk.Frame(win, padding=12)
        outer.pack(fill="both", expand=True)
        ttk.Label(outer, text="Resimler (sayfa sırası yukarıdan aşağıya):").pack(anchor="w")
        body = ttk.Frame(outer)
        body.pack(fill="both", expand=True, pady=6)
        box = tk.Listbox(body, width=60, height=12, selectmode="extended", activestyle="dotbox")
        bar = ttk.Scrollbar(body, command=box.yview)
        box.config(yscrollcommand=bar.set)
        box.pack(side="left", fill="both", expand=True)
        bar.pack(side="left", fill="y")
        paths = list(files)

        def refill(select=()):
            box.delete(0, "end")
            for p in paths:
                box.insert("end", os.path.basename(p))
            for i in select:
                box.selection_set(i)

        def move(step):
            picked = list(box.curselection())
            if not picked or (step < 0 and picked[0] == 0) or (step > 0 and picked[-1] == len(paths) - 1):
                return
            for i in (picked if step < 0 else reversed(picked)):
                paths[i], paths[i + step] = paths[i + step], paths[i]
            refill([i + step for i in picked])

        def remove():
            for i in reversed(box.curselection()):
                del paths[i]
            refill()

        def add():
            more = filedialog.askopenfilenames(title="Resim ekle", filetypes=[("Resimler", exts)], parent=win)
            paths.extend(more)
            refill()

        side = ttk.Frame(outer)
        side.pack(fill="x")
        for text, cmd in (("▲ Yukarı", lambda: move(-1)), ("▼ Aşağı", lambda: move(1)), ("Ekle…", add), ("Çıkar", remove)):
            ttk.Button(side, text=text, command=cmd).pack(side="left", padx=(0, 6))
        opts = ttk.Frame(outer)
        opts.pack(fill="x", pady=(10, 0))
        ttk.Label(opts, text="Sayfa boyutu:").pack(side="left")
        size_var = tk.StringVar(value="A4")
        ttk.Combobox(opts, textvariable=size_var, state="readonly", width=16,
                     values=list(PDFProcessor.IMAGE_PAGE_SIZES)).pack(side="left", padx=6)
        actions = ttk.Frame(outer)
        actions.pack(fill="x", pady=(12, 0))

        def run():
            if not paths:
                messagebox.showwarning("Uyarı", "Listede resim yok.", parent=win)
                return
            first = os.path.splitext(os.path.basename(paths[0]))[0]
            out = filedialog.asksaveasfilename(parent=win, title="PDF'i kaydet", defaultextension=".pdf",
                                               initialdir=os.path.dirname(paths[0]), initialfile=first + ".pdf",
                                               filetypes=[("PDF", "*.pdf")])
            if not out:
                return
            if self._is_same_file(out):
                messagebox.showwarning("Uyarı", "Açık olan dosyanın üzerine yazılamaz. Farklı bir ad seçin.", parent=win)
                return
            try:
                win.config(cursor="watch")
                win.update_idletasks()
                pages = PDFProcessor.images_to_pdf(paths, out, page_size=size_var.get())
            except Exception as exc:  # noqa: BLE001
                win.config(cursor="")
                messagebox.showerror("Hata", f"Bir hata oluştu:\n{exc}", parent=win)
                return
            win.destroy()
            self._toast(f"{pages} sayfalık PDF oluşturuldu: {os.path.basename(out)}", out)

        ttk.Button(actions, text="PDF Oluştur", command=run).pack(side="right")
        ttk.Button(actions, text="Vazgeç", command=win.destroy).pack(side="right", padx=6)
        refill()
        win.grab_set()

    def on_files_dropped(self, event):
        """Dosya gezgininden bırakılan PDF'yi açar; çoklu dosyada birleştirme sunar."""
        paths = list(self.tk.splitlist(event.data))
        pdfs = [p for p in paths if is_openable_document(p) and os.path.exists(p)]
        if not pdfs:
            pics = [p for p in paths if p.lower().endswith(PDFProcessor.IMAGE_EXTENSIONS) and os.path.exists(p)]
            if pics:
                self.action_images_to_pdf(initial=pics)
            return
        try:
            if len(pdfs) == 1:
                self.open_pdf_path(pdfs[0])
            elif messagebox.askyesno("PDF Dosyaları", f"{len(pdfs)} PDF bulundu. Birleştirilsin mi?", parent=self):
                out = filedialog.asksaveasfilename(defaultextension=".pdf", filetypes=[("PDF", "*.pdf")])
                if out:
                    PDFProcessor.merge_pdfs([resolve_document_path(p) for p in pdfs], out)
                    self.open_pdf_path(out)
        except Exception as exc:
            messagebox.showerror("Hata", str(exc), parent=self)

    def _resolve_with_progress(self, path):
        """E-kitap ise PDF'e çevirir; sürecek (önbellekte yok) ise ilerleme penceresi gösterir. PDF'te yolu aynen verir."""
        if not path or not is_book_file(path) or epub_cached_pdf(path):
            return resolve_document_path(path)
        result = {}

        def work():
            try:
                result["path"] = epub_to_pdf(path)
            except Exception as exc:  # noqa: BLE001
                result["error"] = exc

        dialog = tk.Toplevel(self)
        dialog.title("Kitap dönüştürülüyor")
        dialog.transient(self)
        dialog.resizable(False, False)
        ttk.Label(dialog, text=f"{os.path.basename(path)}\nsayfalara dönüştürülüyor, lütfen bekleyin…",
                  justify="center", padding=(24, 16, 24, 8)).pack()
        bar = ttk.Progressbar(dialog, mode="indeterminate", length=260)
        bar.pack(padx=24, pady=(0, 18))
        bar.start(12)
        dialog.protocol("WM_DELETE_WINDOW", lambda: None)
        dialog.update_idletasks()
        x = self.winfo_rootx() + (self.winfo_width() - dialog.winfo_width()) // 2
        y = self.winfo_rooty() + (self.winfo_height() - dialog.winfo_height()) // 3
        dialog.geometry(f"+{max(x, 0)}+{max(y, 0)}")
        dialog.grab_set()
        thread = threading.Thread(target=work, daemon=True)
        thread.start()
        try:
            while thread.is_alive():
                self.update()
                time.sleep(0.03)
        finally:
            bar.stop()
            dialog.grab_release()
            dialog.destroy()
        if "error" in result:
            log_exception("Kitap dönüştürme", result["error"])
            raise result["error"]
        return result["path"]

    def open_pdf_path(self, file_path, start_page=None):
        if not file_path or not os.path.exists(file_path):
            raise FileNotFoundError(f"PDF bulunamadı: {file_path}")
        if not self._confirm_discard():
            return
        self._save_position()
        file_path = self._resolve_with_progress(file_path)
        new_doc = fitz.open(file_path)
        password = None
        if new_doc.needs_pass:
            password = simpledialog.askstring("Şifre", "Bu PDF şifreli. Şifreyi girin:", show="*", parent=self)
            if password is None or not new_doc.authenticate(password):
                new_doc.close()
                raise ValueError("PDF şifresi yanlış veya girilmedi.")
        old_doc = self.doc
        self.doc = new_doc
        self.pdf_path = file_path
        add_recent_file(epub_source_of(file_path) or file_path)
        self._refresh_recent_menu()
        self._remove_quietly(recovery_path(file_path))
        old_doc.close()
        self.is_edited = False
        self._undo_stack.clear()
        self._redo_stack.clear()
        self.selected_pages.clear()
        self.widget_hits = {}
        self._password = password
        self._save_mode = None
        self._redact_scrub_asked = False
        self._nav_back.clear()
        self._nav_fwd.clear()
        self._toc_cache = None
        self.is_welcome = False
        self._refresh_menu_state()
        self.hide_welcome()
        self._search_none = False
        self.search_hits = {}
        self.page_orientation = None
        self.portrait_var.set(False)
        self.landscape_var.set(False)
        self.page_var.set("1")
        self.epub_source = epub_source_of(file_path)
        self.title(self._window_title())
        self.render_thumbnails()
        self.fit_width()
        self._restore_position(start_page)

    def action_search(self):
        query = simpledialog.askstring("PDF İçinde Ara", "Aranacak metin (boş bırakırsanız işaretler temizlenir):", parent=self)
        if query is None:
            return
        hits = {}
        if query:
            for i, page in enumerate(self.doc):
                rects = page.search_for(query)
                if rects:
                    hits[i] = rects
        self.search_hits = hits  # belge değiştirilmez, yalnızca ekranda işaretlenir
        self._hit_index = 0
        self._search_none = bool(query) and not hits
        if hits:
            self._goto_hit(0)  # ilk eşleşmeye git, geçerli eşleşmeyi vurgula
        else:
            self.render_all_pages()

    def _flat_hits(self):
        """Eşleşmeleri belge sırasıyla [(sayfa_indeksi, Rect), ...] olarak döndürür."""
        return [(page, rect) for page in sorted(self.search_hits) for rect in self.search_hits[page]]

    def _update_search_label(self):
        hits = self._flat_hits()
        menu_info = self.__dict__.get("_search_menu")
        if menu_info is not None:
            menu, labels = menu_info
            for text in labels:
                menu.entryconfigure(text, state="normal" if hits else "disabled")
        label = self.__dict__.get("search_label")
        if label is None:
            return
        for name in ("search_prev_button", "search_next_button"):
            button = self.__dict__.get(name)
            if button is not None:
                button.state(["!disabled"] if hits else ["disabled"])
        if hits:
            label.config(text=f"{min(self._hit_index, len(hits) - 1) + 1} / {len(hits)} eşleşme")
        elif getattr(self, "_search_none", False):
            label.config(text="Eşleşme yok")
        else:
            label.config(text="")

    def _goto_hit(self, position):
        """position'daki eşleşmeyi vurgular ve görünür alana kaydırır."""
        hits = self._flat_hits()
        if not hits:
            return
        self._hit_index = position % len(hits)
        page_index, rect = hits[self._hit_index]
        self.render_all_pages()  # geçerli eşleşme farklı renkle çizilir, sayaç güncellenir
        page = self.doc[page_index]
        matrix = self._render_matrix(page, self.zoom)
        origin = page.rect * matrix
        top = self.page_positions[page_index][0] + (rect * page.rotation_matrix * matrix).y0 - origin.y0
        self.canvas.yview_moveto(max(top - self.canvas.winfo_height() / 3, 0) / max(self.total_height, 1))
        self._render_visible()

    def action_split(self):
        folder = filedialog.askdirectory(title="Bölünen PDF'lerin klasörü")
        if not folder:
            return
        base = os.path.splitext(os.path.basename(self.pdf_path))[0]
        names = [os.path.join(folder, f"{base}_sayfa_{i + 1}.pdf") for i in range(len(self.doc))]
        existing = [n for n in names if os.path.exists(n)]
        if existing and not messagebox.askyesno(
                "Üzerine yazılsın mı?", f"{len(existing)} dosya zaten var ve üzerine yazılacak. Devam edilsin mi?", parent=self):
            return
        try:
            for i, name in enumerate(names):
                out = fitz.open()
                try:
                    out.insert_pdf(self.doc, from_page=i, to_page=i)
                    out.save(name, garbage=3, deflate=True)
                finally:
                    out.close()
            self._toast(f"{len(names)} sayfa ayrı PDF olarak kaydedildi", names[0])
        except Exception as e:
            messagebox.showerror("Hata", str(e), parent=self)

    def action_append_pages(self):
        files = filedialog.askopenfilenames(title="Eklenecek PDF'leri seçin", filetypes=[("PDF", "*.pdf")])
        if not files:
            return
        try:
            self._record_edit()
            for path in files:
                with fitz.open(path) as extra:
                    if extra.needs_pass:
                        raise ValueError(f"Şifreli PDF eklenemez: {os.path.basename(path)}")
                    self.doc.insert_pdf(extra)
            self.search_hits = {}
            self.render_thumbnails()
            self.render_all_pages()
        except Exception as e:
            messagebox.showerror("Hata", str(e), parent=self)

    def action_page_layout(self):
        win = tk.Toplevel(self)
        win.title("Sayfa düzeni")
        win.transient(self)
        values = {}
        options = (("Sayfalar", "Tümü"), ("Boyut", "Mevcut"), ("Genişlik", "595"), ("Yükseklik", "842"), ("Kenar boşluğu", "0"), ("Arka plan RGB", "255,255,255"))
        for row, (key, default) in enumerate(options):
            ttk.Label(win, text=key).grid(row=row, column=0, sticky="w", padx=10, pady=4)
            values[key] = tk.StringVar(value=default)
            if key == "Sayfalar":
                ttk.Combobox(win, textvariable=values[key], values=("Tümü", "Seçilenler"), state="readonly", width=18).grid(row=row, column=1, padx=10, pady=4)
            elif key == "Boyut":
                ttk.Combobox(win, textvariable=values[key], values=("Mevcut", "A4", "Letter", "Özel"), state="readonly", width=18).grid(row=row, column=1, padx=10, pady=4)
            else:
                ttk.Entry(win, textvariable=values[key], width=20).grid(row=row, column=1, padx=10, pady=4)
        def apply():
            try:
                indexes = sorted(self.selected_pages) if values["Sayfalar"].get() == "Seçilenler" and self.selected_pages else list(range(len(self.doc)))
                color = tuple(max(0, min(255, int(x.strip()))) / 255 for x in values["Arka plan RGB"].get().split(','))
                size = values["Boyut"].get()
                self._record_edit()
                for index in indexes:
                    page = self.doc[index]
                    if size == "A4":
                        width, height = 595, 842
                    elif size == "Letter":
                        width, height = 612, 792
                    elif size == "Özel":
                        width, height = float(values["Genişlik"].get()), float(values["Yükseklik"].get())
                    else:
                        width, height = page.rect.width, page.rect.height
                    margin = max(0, float(values["Kenar boşluğu"].get()))
                    page.set_mediabox(fitz.Rect(0, 0, width, height))
                    page.set_cropbox(fitz.Rect(margin, margin, max(margin + 1, width - margin), max(margin + 1, height - margin)))
                    page.draw_rect(page.rect, color=color, fill=color, overlay=False)
                win.destroy()
                self.render_thumbnails()
                self.render_all_pages()
            except (ValueError, TypeError) as exc:
                messagebox.showerror("Sayfa düzeni hatası", str(exc), parent=win)
        ttk.Button(win, text="Uygula", command=apply).grid(row=7, column=0, padx=10, pady=10)
        ttk.Button(win, text="İptal", command=win.destroy).grid(row=7, column=1, padx=10, pady=10)

    def action_headers_footers(self):
        win = tk.Toplevel(self)
        win.title("Üst bilgi, alt bilgi ve sayfa numarası")
        win.transient(self)
        win.resizable(False, False)
        fields = {}
        defaults = (("Aralık", "1-" + str(len(self.doc))), ("Üst bilgi", ""), ("Alt bilgi", "Sayfa {page}/{total}"), ("Yazı boyutu", "9"))
        for row, (key, default) in enumerate(defaults):
            ttk.Label(win, text=key).grid(row=row, column=0, sticky="w", padx=10, pady=5)
            fields[key] = tk.StringVar(value=default)
            ttk.Entry(win, textvariable=fields[key], width=42).grid(row=row, column=1, padx=10, pady=5)
        ttk.Label(win, text="Şablonlar: {page}, {total}, {filename}", foreground="#6b7280").grid(row=4, column=0, columnspan=2, padx=10, pady=5)
        def apply():
            try:
                pages = PDFProcessor.parse_page_range(fields["Aralık"].get(), len(self.doc))
                size = max(5, min(30, float(fields["Yazı boyutu"].get())))
                if not pages:
                    raise ValueError("Geçerli sayfa aralığı yok.")
                self._record_edit()
                total = len(self.doc)
                filename = os.path.basename(self.pdf_path)
                for number in pages:
                    page = self.doc[number - 1]
                    rect = page.rect
                    top = fields["Üst bilgi"].get().format(page=number, total=total, filename=filename)
                    bottom = fields["Alt bilgi"].get().format(page=number, total=total, filename=filename)
                    if top:
                        insert_text_tr(page, (36, 24), top, fontsize=size)
                    if bottom:
                        insert_text_tr(page, (36, rect.height - 24), bottom, fontsize=size)
                win.destroy()
                self.render_all_pages()
            except (ValueError, KeyError, IndexError) as exc:
                messagebox.showerror("Üst/alt bilgi hatası", str(exc), parent=win)
        ttk.Button(win, text="Uygula", command=apply).grid(row=5, column=0, padx=10, pady=10)
        ttk.Button(win, text="İptal", command=win.destroy).grid(row=5, column=1, padx=10, pady=10)

    def action_watermark(self):
        values = self._form_dialog("Filigran Ekle", [
            ("text", "Filigran metni", "GİZLİ", str),
            ("size", "Yazı boyutu", 56, float),
            ("opacity", "Saydamlık (% 5-100)", 25, float),
            ("angle", "Açı (derece)", 45, float),
        ])
        if not values or not values["text"]:
            return
        size = max(8.0, values["size"])
        opacity = min(max(values["opacity"], 5.0), 100.0) / 100
        self._record_edit()
        for page in self.doc:
            view = page.rect
            center = fitz.Point(view.width / 2, view.height / 2) * page.derotation_matrix
            angle = values["angle"] + page.rotation  # dönmüş sayfada filigran görünürde istenen açıda kalsın
            fit = math.hypot(view.width, view.height) * 0.9
            width = text_width_tr(values["text"], size)
            page_size = size * min(1.0, fit / width) if width else size
            width = text_width_tr(values["text"], page_size)
            radians = math.radians(angle)
            direction = (math.cos(radians), -math.sin(radians))
            down = (math.sin(radians), math.cos(radians))
            start = fitz.Point(center.x - direction[0] * width / 2 + down[0] * page_size * 0.35,
                               center.y - direction[1] * width / 2 + down[1] * page_size * 0.35)
            insert_text_tr(page, start, values["text"], fontsize=page_size, color=(0.6, 0.6, 0.72),
                           angle=angle, opacity=opacity)
        self.render_all_pages()

    def _ask_new_password(self, title):
        """Yeni parolayı iki kez ister; eşleşmezse None. Yanlış yazılan parola belgeyi kilitlememeli."""
        password = simpledialog.askstring(title, "Yeni şifre:", show="*", parent=self)
        if not password:
            return None
        again = simpledialog.askstring(title, "Şifreyi doğrulamak için tekrar girin:", show="*", parent=self)
        if again != password:
            messagebox.showwarning(title, "Şifreler eşleşmedi; işlem iptal edildi.", parent=self)
            return None
        return password

    def action_encrypt(self):
        password = self._ask_new_password("PDF Şifrele")
        if not password:
            return
        path = filedialog.asksaveasfilename(defaultextension=".pdf", filetypes=[("PDF", "*.pdf")])
        if not path:
            return
        if self._is_same_file(path):
            messagebox.showwarning("Uyarı", "Açık olan dosyanın üzerine yazılamaz. Farklı bir ad seçin.", parent=self)
            return
        try:
            # Sahip şifresi ayrı ve rastgele: aksi halde açma şifresini bilen izinleri de kaldırabilir.
            self.doc.save(path, encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw=secrets.token_urlsafe(16),
                          user_pw=password, permissions=fitz.PDF_PERM_PRINT)
            self._toast("Şifreli PDF oluşturuldu (yalnızca yazdırmaya izin verir)", path)
        except Exception as e:
            messagebox.showerror("Hata", str(e))

    def action_decrypt(self):
        if not self.doc.is_encrypted:
            messagebox.showinfo("Bilgi", "Açık PDF zaten şifresiz.", parent=self)
            return
        path = filedialog.asksaveasfilename(defaultextension=".pdf", filetypes=[("PDF", "*.pdf")])
        if not path:
            return
        if self._is_same_file(path):
            messagebox.showwarning("Uyarı", "Açık olan dosyanın üzerine yazılamaz. Farklı bir ad seçin.", parent=self)
            return
        try:
            # Belge açılırken şifre zaten doğrulandı; düzenlemeler de korunarak şifresiz yazılır.
            self.doc.save(path, encryption=fitz.PDF_ENCRYPT_NONE)
            self._toast("Şifresiz PDF oluşturuldu", path)
        except Exception as e:
            messagebox.showerror("Hata", str(e))

    # ------------------------------------------------------------------
    # OCR ve toplu metin çıkarma
    # ------------------------------------------------------------------
    def _snapshot(self):
        return self.doc.tobytes(garbage=4, deflate=True, clean=True)

    def _record_edit(self):
        self._undo_stack.append(self._snapshot())
        self._redo_stack.clear()
        # Her adım belgenin tamamını tuttuğundan yığın adım sayısı ve toplam bayt ile sınırlanır.
        while len(self._undo_stack) > UNDO_MAX_STEPS or (
                len(self._undo_stack) > 1 and sum(len(item) for item in self._undo_stack) > UNDO_MAX_BYTES):
            self._undo_stack.pop(0)
        self.is_edited = True

    def _swap_document(self, data):
        """Geri al/yinele: belgeyi anlık görüntüden değiştirir, görünümü ve küçük resimleri tazeler."""
        old = self.doc
        self.doc = fitz.open(stream=data, filetype="pdf")
        old.close()
        self.is_edited = True
        self.search_hits = {}
        self.selected_pages = {i for i in self.selected_pages if i < len(self.doc)}
        self.render_thumbnails()
        self.render_all_pages()

    def undo(self):
        if self._undo_stack:
            self._redo_stack.append(self._snapshot())
            self._swap_document(self._undo_stack.pop())

    def redo(self):
        if self._redo_stack:
            self._undo_stack.append(self._snapshot())
            self._swap_document(self._redo_stack.pop())

    def action_annotation(self, kind):
        """Geçerli sayfada aranan metni vurgular / altını / üstünü / dalgalı altını çizer (kind: MARKUP_KINDS)."""
        query = simpledialog.askstring("Metin seçimi", "İşaretlenecek metin:", parent=self)
        if not query:
            return
        page = self.doc[self.get_current_page_index()]
        rects = page.search_for(query)
        if not rects:
            self._toast("Bu sayfada aranan metin bulunamadı.")
            return
        self._record_edit()
        add_markup(self.doc[self.get_current_page_index()], rects, kind)
        self.render_all_pages()

    def action_remove_annotations(self):
        self._record_edit()
        for page in self.doc:
            annot = page.first_annot
            while annot:
                nxt = annot.next
                page.delete_annot(annot)
                annot = nxt
        self.render_all_pages()

    def action_add_bookmark(self):
        title = simpledialog.askstring("Yer imi", "Yer imi başlığı:", parent=self)
        if title:
            self._record_edit()
            toc = self.doc.get_toc()
            toc.append([1, title, self.get_current_page_index() + 1])
            self.doc.set_toc(toc)

    def action_extract_pages(self):
        value = simpledialog.askstring("Sayfa aralığı", "Örnek: 1-3, 5, 8-10", parent=self)
        if not value:
            return
        try:
            pages = PDFProcessor.parse_page_range(value, len(self.doc))
            if not pages:
                raise ValueError("Geçerli sayfa aralığı bulunamadı.")
            path = filedialog.asksaveasfilename(defaultextension='.pdf', filetypes=[('PDF', '*.pdf')])
            if path:
                self._export_from_working(lambda src: PDFProcessor.extract_pages(src, pages, path),
                                          f"{len(pages)} sayfa dışa aktarıldı.")
        except Exception as exc:
            messagebox.showerror("Dışa aktarma hatası", str(exc), parent=self)

    def action_check_document(self):
        try:
            lines, problem = PDFProcessor.check_document(self.doc)
        except Exception as exc:  # noqa: BLE001
            log_exception("PDF doğrulama", exc)
            messagebox.showerror("Hata", str(exc), parent=self)
            return
        title = "Sağlık denetimi: sorun bulundu" if problem else "Sağlık denetimi: sorun yok"
        tail = "\n\nİpucu: Araçlar ▸ Bozuk Dosyayı Onar ile temiz bir kopya kaydedebilirsiniz." if problem else ""
        messagebox.showinfo(title, "\n".join(lines) + tail, parent=self)

    def action_repair(self):
        path = filedialog.asksaveasfilename(title="Onarılmış kopyayı kaydet", defaultextension='.pdf',
                                            filetypes=[('PDF', '*.pdf')])
        if not path:
            return
        self._export_from_working(lambda src: PDFProcessor.repair_pdf(src, path), None,
                                  on_ok=lambda: self._toast("Yapısı yeniden yazılmış temiz kopya kaydedildi", path))

    def action_crop(self):
        margin_mm = simpledialog.askfloat("Kenarları kırp", "Her kenardan kırpılacak boşluk (mm):",
                                          initialvalue=10.0, minvalue=0.5, maxvalue=100.0, parent=self)
        if margin_mm is None:
            return
        indices = sorted(i for i in self.selected_pages if i < len(self.doc)) or list(range(len(self.doc)))
        scope = f"{len(indices)} seçili sayfa" if self.selected_pages else "tüm sayfalar"
        if not messagebox.askyesno("Kenarları kırp", f"{scope} için her kenardan {margin_mm:g} mm kırpılsın mı?\n"
                                   "(Ctrl+Z ile geri alınabilir.)", parent=self):
            return
        try:
            self._record_edit()
            PDFProcessor.crop_margins(self.doc, margin_mm / 25.4 * 72, indices)
        except Exception as exc:  # noqa: BLE001
            if self._undo_stack:
                self._swap_document(self._undo_stack.pop())  # yarım kalan kırpmayı geri al
            messagebox.showerror("Kırpma hatası", str(exc), parent=self)
            return
        self.search_hits = {}
        self.render_thumbnails()
        self.render_all_pages()
        self._toast(f"{len(indices)} sayfa kırpıldı")

    def action_crop_tool(self):
        """Kırp aracını seçer; kullanıcı sayfada tutulacak alanı sürükleyerek belirler."""
        if self.is_welcome:
            return
        self.tool_var.set("kirp")
        self._on_tool_change()
        self._toast("Kırp aracı seçildi: sayfada tutulacak alanı sürükleyerek seçin (✋ ile normal gezintiye dönersiniz)")

    def action_crop_reset(self):
        scope_selected = sorted(i for i in self.selected_pages if i < len(self.doc))
        indices = scope_selected or list(range(len(self.doc)))
        cropped = [i for i in indices if self.doc[i].cropbox != self.doc[i].mediabox]
        if not cropped:
            messagebox.showinfo("Kırpmayı kaldır", "Kırpılmış sayfa yok.", parent=self)
            return
        scope = f"{len(cropped)} sayfanın" if scope_selected else f"{len(cropped)} kırpılmış sayfanın"
        if not messagebox.askyesno("Kırpmayı kaldır", f"{scope} kırpması kaldırılsın mı?\n(Ctrl+Z ile geri alınabilir.)",
                                   parent=self):
            return
        self._apply_edit(lambda: PDFProcessor.reset_crop(self.doc, cropped), lambda n: f"{n} sayfanın kırpması kaldırıldı")

    def action_insert_blank(self):
        count = simpledialog.askinteger("Boş sayfa ekle", "Kaç boş sayfa eklensin?", initialvalue=1, minvalue=1,
                                        maxvalue=500, parent=self)
        if not count:
            return
        current = self.get_current_page_index()
        after = messagebox.askyesnocancel("Boş sayfa ekle", f"Evet: {current + 1}. sayfadan SONRA\n"
                                          f"Hayır: {current + 1}. sayfadan ÖNCE\nİptal: vazgeç", parent=self)
        if after is None:
            return
        position = current if after else current - 1
        self._apply_edit(lambda: PDFProcessor.insert_blank_pages(self.doc, position, count),
                         f"{count} boş sayfa eklendi")

    def action_delete_page_range(self):
        value = simpledialog.askstring("Sayfa aralığını sil", "Silinecek sayfalar (örn. 1-3, 5, 8-10):", parent=self)
        if not value:
            return
        try:
            numbers = PDFProcessor.parse_page_range(value, len(self.doc))
        except ValueError:
            messagebox.showerror("Hata", "Sayfa aralığı anlaşılamadı. Örnek: 1-3, 5, 8-10", parent=self)
            return
        if not numbers:
            messagebox.showinfo("Sayfa sil", "Belgede bu aralıkta sayfa yok.", parent=self)
            return
        if not messagebox.askyesno("Sayfa sil", f"{len(numbers)} sayfa silinsin mi? (Ctrl+Z ile geri alınabilir.)",
                                   parent=self):
            return
        self.selected_pages.clear()
        self._apply_edit(lambda: PDFProcessor.delete_page_numbers(self.doc, numbers), lambda n: f"{n} sayfa silindi")

    OFFICE_KINDS = {
        "docx": ("Word Belgesi", ".docx", "Word"),
        "xlsx": ("Excel Çalışma Kitabı", ".xlsx", "Excel"),
        "pptx": ("PowerPoint Sunusu", ".pptx", "PowerPoint"),
    }

    def action_save_office(self, kind):
        """Açık (düzenlenmiş) belgeyi Word/Excel/PowerPoint biçiminde yeni dosya olarak kaydeder."""
        if self.is_welcome:
            return
        label, ext, program = self.OFFICE_KINDS[kind]
        editable = True
        if kind in ("docx", "pptx"):
            answer = messagebox.askyesnocancel(
                f"{program} olarak kaydet",
                "Evet: düzenlenebilir (metin ve resimler ayrı ayrı; sayfa görünümü birebir olmayabilir)\n"
                "Hayır: her sayfa tek resim (görünüm birebir, metin düzenlenemez)\nİptal: vazgeç", parent=self)
            if answer is None:
                return
            editable = answer
        base = os.path.splitext(os.path.basename(self.pdf_path or "belge"))[0]
        path = filedialog.asksaveasfilename(title=f"{program} olarak kaydet", defaultextension=ext,
                                            initialfile=base + ext, filetypes=[(label, "*" + ext)], parent=self)
        if not path:
            return
        result = {}

        def work(source):
            if kind == "docx":
                PDFProcessor.export_docx(source, path, editable=editable)
            elif kind == "xlsx":
                result["tables"] = PDFProcessor.export_xlsx(source, path)
            else:
                PDFProcessor.export_pptx(source, path, editable=editable)

        def done():
            note = ""
            if kind == "xlsx":
                note = (f" ({result['tables']} tablo)" if result.get("tables")
                        else " (tablo bulunamadı; metin satır satır yazıldı)")
            self._toast(f"{program} dosyası kaydedildi: {os.path.basename(path)}{note}", path)

        self.config(cursor="watch")
        self.update_idletasks()
        try:
            self._export_from_working(work, None, on_ok=done)
        finally:
            self.config(cursor="")

    def action_booklet(self):
        path = filedialog.asksaveasfilename(title="Kitapçık PDF'ini kaydet", defaultextension='.pdf',
                                            filetypes=[('PDF', '*.pdf')])
        if not path:
            return
        self._export_from_working(lambda src: PDFProcessor.make_booklet(src, path), None,
                                  on_ok=lambda: self._toast("Kitapçık düzeni oluşturuldu (çift yüz yazdırıp ortadan katlayın)", path))

    def _apply_edit(self, func, done_text):
        """Belgede geri alınabilir bir düzenleme yapar; hata olursa önceki hâle döner. done_text str veya sonuç->str."""
        try:
            self._record_edit()
            result = func()
        except Exception as exc:  # noqa: BLE001
            log_exception("Düzenleme", exc)
            if self._undo_stack:
                self._swap_document(self._undo_stack.pop())
            messagebox.showerror("Hata", str(exc), parent=self)
            return None
        self.search_hits = {}
        self.render_thumbnails()
        self.render_all_pages()
        self._toast(done_text(result) if callable(done_text) else done_text)
        return result

    def action_sanitize(self):
        if not messagebox.askyesno("Gizli verileri temizle",
                                   "JavaScript, gömülü ve ekli dosyalar, küçük resimler, belge bilgileri ve XML meta verisi silinsin mi?\n"
                                   "Metin, görsel, bağlantı ve form alanları korunur. (Ctrl+Z ile geri alınabilir.)", parent=self):
            return
        self._apply_edit(lambda: PDFProcessor.sanitize(self.doc), "Gizli veriler temizlendi")

    def action_flatten(self):
        if not messagebox.askyesno("Düzleştir",
                                   "Açıklamalar ve form alanları sayfaya sabitlenecek; artık düzenlenemeyecek. Devam edilsin mi?\n"
                                   "(Ctrl+Z ile geri alınabilir.)", parent=self):
            return
        self._apply_edit(lambda: PDFProcessor.flatten(self.doc), "Açıklamalar ve formlar sayfaya düzleştirildi")

    def action_remove_blank(self):
        blank = PDFProcessor.find_blank_pages(self.doc)
        if not blank:
            messagebox.showinfo("Boş sayfalar", "Boş sayfa bulunamadı.", parent=self)
            return
        shown = ", ".join(str(i + 1) for i in blank[:20]) + (" …" if len(blank) > 20 else "")
        if not messagebox.askyesno("Boş sayfaları sil", f"{len(blank)} boş sayfa bulundu (sayfa {shown}). Silinsin mi?\n"
                                   "(Ctrl+Z ile geri alınabilir.)", parent=self):
            return
        self._apply_edit(lambda: PDFProcessor.remove_blank_pages(self.doc), lambda n: f"{n} boş sayfa silindi")

    def action_extract_images(self):
        if PDFProcessor.count_images(self.doc) == 0:
            messagebox.showinfo("Görseller", "Belgede görsel bulunamadı.", parent=self)
            return
        parent_dir = filedialog.askdirectory(title="Görsellerin yazılacağı klasörü seçin")
        if not parent_dir:
            return
        name = "gorseller" if self.is_welcome else os.path.splitext(os.path.basename(self.pdf_path))[0] + "_gorseller"
        folder = os.path.join(parent_dir, name)
        try:
            count = PDFProcessor.extract_images(self.doc, folder)
            self._toast(f"{count} görsel dışa aktarıldı", folder)
        except Exception as exc:  # noqa: BLE001
            log_exception("Görsel dışa aktarma", exc)
            messagebox.showerror("Hata", str(exc), parent=self)

    def action_remove_images(self):
        total = PDFProcessor.count_images(self.doc)
        if total == 0:
            messagebox.showinfo("Görseller", "Belgede görsel bulunamadı.", parent=self)
            return
        if not messagebox.askyesno("Görselleri sil", f"Belgedeki {total} görsel silinsin mi? Metin korunur.\n"
                                   "(Ctrl+Z ile geri alınabilir.)", parent=self):
            return
        self._apply_edit(lambda: PDFProcessor.remove_images(self.doc), lambda n: f"{n} görsel silindi")

    def action_permissions(self):
        allow_print = messagebox.askyesno("İzinler", "Yazdırmaya izin verilsin mi?", parent=self)
        allow_copy = messagebox.askyesno("İzinler", "Metin ve görsel kopyalamaya izin verilsin mi?", parent=self)
        allow_edit = messagebox.askyesno("İzinler", "Düzenlemeye (sayfa, açıklama, form) izin verilsin mi?", parent=self)
        path = filedialog.asksaveasfilename(title="İzinleri değiştirilmiş kopyayı kaydet", defaultextension=".pdf",
                                            filetypes=[("PDF", "*.pdf")])
        if not path:
            return
        if self._is_same_file(path):
            messagebox.showwarning("Uyarı", "Açık olan dosyanın üzerine yazılamaz. Farklı bir ad seçin.", parent=self)
            return
        try:
            PDFProcessor.save_with_permissions(self.doc, path, allow_print, allow_copy, allow_edit)
            self._toast("İzinleri değiştirilmiş kopya kaydedildi (açmak için parola gerekmez; izinleri kaldırmak için sahip parolası gerekir)", path)
        except Exception as exc:  # noqa: BLE001
            log_exception("İzin değiştirme", exc)
            messagebox.showerror("Hata", str(exc), parent=self)

    def action_stamp(self):
        values = self._form_dialog("Damga Ekle", [
            ("text", "Damga metni", "ONAYLANDI", str),
            ("size", "Yazı boyutu", 28, float),
            ("position", "Konum (1-9: 7 sol üst, 5 orta, 3 sağ alt)", 9, int),
            ("pages", "Sayfalar (örn. 1-3, 5; boş = tümü)", "", str),
        ])
        if not values or not values["text"]:
            return
        try:
            if values["pages"].strip():
                indices = [n - 1 for n in PDFProcessor.parse_page_range(values["pages"], len(self.doc))]
                if not indices:
                    raise ValueError("Geçerli sayfa aralığı bulunamadı.")
            else:
                indices = list(range(len(self.doc)))
            size = max(8.0, values["size"])
            self._apply_edit(lambda: PDFProcessor.stamp_pages(self.doc, values["text"], size, values["position"], indices),
                             f"{len(indices)} sayfaya damga eklendi")
        except ValueError as exc:
            messagebox.showerror("Damga", str(exc), parent=self)

    def action_nup(self):
        values = self._form_dialog("Çok Sayfalı Düzen", [("per", "Yaprak başına sayfa (2, 4, 6 veya 9)", 4, int)])
        if not values:
            return
        path = filedialog.asksaveasfilename(title="Düzeni kaydet", defaultextension=".pdf", filetypes=[("PDF", "*.pdf")])
        if not path:
            return
        self._export_from_working(lambda src: PDFProcessor.n_up(src, path, values["per"]), None,
                                  on_ok=lambda: self._toast(f"Yaprağa {values['per']} sayfa düzeni oluşturuldu", path))

    def action_single_page(self):
        path = filedialog.asksaveasfilename(title="Tek uzun sayfayı kaydet", defaultextension=".pdf", filetypes=[("PDF", "*.pdf")])
        if not path:
            return
        self._export_from_working(lambda src: PDFProcessor.single_long_page(src, path), None,
                                  on_ok=lambda: self._toast("Tüm sayfalar tek uzun sayfada birleştirildi", path))

    def action_redact(self):
        query = simpledialog.askstring("Kalıcı karartma", "Karartılacak metin (tüm sayfalarda aranır):", parent=self)
        if not query:
            return
        total = PDFProcessor.count_matches(self.doc, query)
        if total == 0:
            messagebox.showinfo("Kalıcı karartma", "Eşleşme bulunamadı.", parent=self)
            return
        if not messagebox.askyesno("Kalıcı karartma",
                                   f"{total} eşleşme siyah kutuyla kapatılacak ve altındaki metin belgeden silinecek.\n"
                                   "Kaydedilmeden önce Ctrl+Z ile geri alınabilir; kaydedilirse geri getirilemez. Devam edilsin mi?",
                                   parent=self):
            return
        self._apply_edit(lambda: PDFProcessor.redact_text(self.doc, query), lambda n: f"{n} eşleşme kalıcı olarak karartıldı")

    def action_edit_toc(self):
        win = tk.Toplevel(self)
        win.title("Yer İmlerini Düzenle")
        win.transient(self)
        win.grab_set()
        ttk.Label(win, text="Her satır: seviye | başlık | sayfa     (alt başlık için seviye 2, 3, …)").pack(anchor="w", padx=10, pady=(10, 4))
        box = tk.Text(win, width=72, height=18, wrap="none")
        box.pack(fill="both", expand=True, padx=10)
        box.insert("1.0", PDFProcessor.toc_to_text(self.doc.get_toc()))

        def save():
            try:
                toc = PDFProcessor.parse_toc_text(box.get("1.0", "end"), len(self.doc))
            except ValueError as exc:
                messagebox.showerror("Yer imleri", str(exc), parent=win)
                return
            self._record_edit()
            self.doc.set_toc(toc)
            win.destroy()
            self._toast(f"{len(toc)} yer imi kaydedildi")

        row = ttk.Frame(win)
        row.pack(fill="x", padx=10, pady=10)
        ttk.Button(row, text="Kaydet", command=save).pack(side="right")
        ttk.Button(row, text="Vazgeç", command=win.destroy).pack(side="right", padx=6)

    def action_compare(self):
        other = filedialog.askopenfilename(title="Karşılaştırılacak ikinci PDF", filetypes=[("PDF", "*.pdf")])
        if not other:
            return
        name_a = "(açılış sayfası)" if self.is_welcome else os.path.basename(self.pdf_path)
        try:
            with fitz.open(other) as second:
                if second.needs_pass:
                    raise ValueError("Şifreli PDF karşılaştırılamaz.")
                lines, different = PDFProcessor.compare_documents(self.doc, second, name_a, os.path.basename(other))
        except Exception as exc:  # noqa: BLE001
            log_exception("PDF karşılaştırma", exc)
            messagebox.showerror("Hata", str(exc), parent=self)
            return
        summary = "\n".join(lines[:25]) + ("\n…" if len(lines) > 25 else "")
        messagebox.showinfo("Karşılaştırma sonucu", summary, parent=self)
        if different and messagebox.askyesno("Rapor", "Tam raporu TXT olarak kaydetmek ister misiniz?", parent=self):
            path = filedialog.asksaveasfilename(title="Raporu kaydet", defaultextension=".txt", filetypes=[("Metin", "*.txt")])
            if path:
                with open(path, "w", encoding="utf-8") as handle:
                    handle.write("\n".join(lines) + "\n")
                self._toast("Karşılaştırma raporu kaydedildi", path)

    def action_auto_rotate(self):
        fixes = PDFProcessor.auto_rotate(self.doc, dry_run=True)
        if not fixes:
            messagebox.showinfo("Otomatik döndür", "Yan veya ters duran sayfa bulunamadı.", parent=self)
            return
        shown = ", ".join(str(i + 1) for i in sorted(fixes)[:20]) + (" …" if len(fixes) > 20 else "")
        if not messagebox.askyesno("Otomatik döndür", f"{len(fixes)} sayfanın metni yan/ters duruyor (sayfa {shown}). Döndürülsün mü?\n"
                                   "(Ctrl+Z ile geri alınabilir.)", parent=self):
            return
        self._apply_edit(lambda: PDFProcessor.auto_rotate(self.doc), lambda done: f"{len(done)} sayfa döndürüldü")

    def action_overlay(self):
        other = filedialog.askopenfilename(title="Üstüne bindirilecek PDF", filetypes=[("PDF", "*.pdf")])
        if not other:
            return
        path = filedialog.asksaveasfilename(title="Sonucu kaydet", defaultextension=".pdf", filetypes=[("PDF", "*.pdf")])
        if not path:
            return
        self._export_from_working(lambda src: PDFProcessor.overlay_pdf(src, other, path), None,
                                  on_ok=lambda: self._toast("PDF bindirildi", path))

    def action_attachments(self):
        win = tk.Toplevel(self)
        win.title("Gömülü Ekler")
        win.transient(self)
        win.grab_set()
        ttk.Label(win, text="Bu PDF'in içine gömülü dosyalar:").pack(anchor="w", padx=10, pady=(10, 4))
        listbox = tk.Listbox(win, width=60, height=10, exportselection=False)
        listbox.pack(fill="both", expand=True, padx=10)

        def refresh():
            listbox.delete(0, "end")
            for name, size in PDFProcessor.list_attachments(self.doc):
                listbox.insert("end", f"{name}   ({size:,} bayt)")

        def selected_name():
            picked = listbox.curselection()
            if not picked:
                messagebox.showinfo("Ekler", "Önce listeden bir dosya seçin.", parent=win)
                return None
            return PDFProcessor.list_attachments(self.doc)[picked[0]][0]

        def add():
            for file in filedialog.askopenfilenames(title="Eklenecek dosyalar", parent=win):
                self._record_edit()
                PDFProcessor.add_attachment(self.doc, file)
            refresh()

        def export():
            name = selected_name()
            if name:
                out = filedialog.asksaveasfilename(initialfile=name, parent=win)
                if out:
                    PDFProcessor.extract_attachment(self.doc, name, out)
                    self._toast("Ek dışa aktarıldı", out)

        def remove():
            name = selected_name()
            if name and messagebox.askyesno("Ek sil", f"'{name}' PDF'ten silinsin mi? (Ctrl+Z ile geri alınabilir.)", parent=win):
                self._record_edit()
                PDFProcessor.remove_attachment(self.doc, name)
                refresh()

        refresh()
        row = ttk.Frame(win)
        row.pack(fill="x", padx=10, pady=10)
        ttk.Button(row, text="Kapat", command=win.destroy).pack(side="right")
        ttk.Button(row, text="Sil", command=remove).pack(side="right", padx=6)
        ttk.Button(row, text="Dışa aktar…", command=export).pack(side="right")
        ttk.Button(row, text="Ekle…", command=add).pack(side="right", padx=6)

    def _menu_commands(self):
        """Menü çubuğundaki tüm çalıştırılabilir komutlar: [(yol, menü, indeks)]."""
        found = []

        def walk(menu, trail):
            for index in range(menu.index("end") + 1 if menu.index("end") is not None else 0):
                kind = menu.type(index)
                if kind == "cascade":
                    walk(self.nametowidget(menu.entrycget(index, "menu")), trail + [menu.entrycget(index, "label")])
                elif kind in ("command", "checkbutton") and str(menu.entrycget(index, "state")) != "disabled":
                    found.append((" ▸ ".join(trail + [menu.entrycget(index, "label")]), menu, index))

        bar = self.nametowidget(self.cget("menu"))
        for index in range(bar.index("end") + 1):
            if bar.type(index) == "cascade":
                walk(self.nametowidget(bar.entrycget(index, "menu")), [bar.entrycget(index, "label")])
        return found

    def action_command_palette(self):
        commands = self._menu_commands()
        win = tk.Toplevel(self)
        win.title("Komut Ara")
        win.transient(self)
        win.geometry("560x340")
        entry = ttk.Entry(win)
        entry.pack(fill="x", padx=10, pady=(10, 4))
        listbox = tk.Listbox(win, activestyle="dotbox", exportselection=False)
        listbox.pack(fill="both", expand=True, padx=10, pady=(0, 10))
        shown = []

        def refill(*_):
            shown[:] = [c for c in commands if turkish_fold(entry.get()) in turkish_fold(c[0])]
            listbox.delete(0, "end")
            for item in shown:
                listbox.insert("end", item[0])
            if shown:
                listbox.selection_set(0)

        def run(_event=None):
            picked = listbox.curselection()
            if picked:
                _path, menu, index = shown[picked[0]]
                win.destroy()
                menu.invoke(index)

        def move(step):
            picked = listbox.curselection()
            target = min(max((picked[0] if picked else -1) + step, 0), max(len(shown) - 1, 0))
            listbox.selection_clear(0, "end")
            listbox.selection_set(target)
            listbox.see(target)
            return "break"

        entry.bind("<KeyRelease>", refill)
        entry.bind("<Return>", run)
        entry.bind("<Down>", lambda _e: move(1))
        entry.bind("<Up>", lambda _e: move(-1))
        listbox.bind("<Double-1>", run)
        win.bind("<Escape>", lambda _e: win.destroy())
        refill()
        entry.focus_set()

    def action_compress(self):
        path = filedialog.asksaveasfilename(defaultextension='.pdf', filetypes=[('PDF', '*.pdf')])
        if not path:
            return
        before = os.path.getsize(self.pdf_path) if os.path.exists(self.pdf_path) else 0

        def finish():
            after = os.path.getsize(path)
            note = f" ({before:,} → {after:,} bayt)" if before else ""
            self._toast(f"Sıkıştırıldı{note}", path)

        self._export_from_working(lambda src: PDFProcessor.compress(src, path), None, on_ok=finish)

    def action_metadata(self):
        meta = dict(self.doc.metadata)
        for key, label in [('title', 'Başlık'), ('author', 'Yazar'), ('subject', 'Konu'), ('keywords', 'Anahtar kelimeler')]:
            value = simpledialog.askstring('Metadata', f'{label}:', initialvalue=meta.get(key, '') or '', parent=self)
            if value is not None:
                meta[key] = value
        self._record_edit()
        self.doc.set_metadata(meta)
        self._toast("Belge bilgileri güncellendi")

    def action_contact_sheet(self):
        path = filedialog.asksaveasfilename(defaultextension='.png', filetypes=[('PNG', '*.png')])
        if not path:
            return
        self._export_from_working(lambda src: PDFProcessor.create_contact_sheet(src, path),
                                  "İletişim paftası oluşturuldu.")

    def action_epub_settings(self):
        """E-kitap okuma ayarları: yazı boyutu ve sayfa boyutu. Açık kitap varsa yeni ayarla, aynı yerden yeniden düzenlenir."""
        settings = load_settings()
        dialog = tk.Toplevel(self)
        dialog.title("E-Kitap Okuma Ayarları")
        dialog.transient(self)
        dialog.resizable(False, False)
        body = ttk.Frame(dialog, padding=16)
        body.pack()
        old_font = settings.get("epub_font") if settings.get("epub_font") in EPUB_FONT_SIZES else "Orta"
        old_page = settings.get("epub_page") if settings.get("epub_page") in EPUB_PAGE_SIZES else "Kitap"
        font_var = tk.StringVar(value=old_font)
        page_var = tk.StringVar(value=old_page)
        ttk.Label(body, text="Yazı boyutu:").grid(row=0, column=0, sticky="w", pady=4)
        ttk.Combobox(body, textvariable=font_var, values=list(EPUB_FONT_SIZES), state="readonly", width=16).grid(
            row=0, column=1, padx=(10, 0))
        ttk.Label(body, text="Sayfa boyutu:").grid(row=1, column=0, sticky="w", pady=4)
        ttk.Combobox(body, textvariable=page_var, values=list(EPUB_PAGE_SIZES), state="readonly", width=16).grid(
            row=1, column=1, padx=(10, 0))
        note = "Ayarlar sonraki e-kitaplara da uygulanır."
        if getattr(self, "epub_source", None):
            note = "Açık kitap yeni ayarlarla yeniden düzenlenir; okuduğunuz yerden devam edersiniz."
        ttk.Label(body, text=note, wraplength=300).grid(row=2, column=0, columnspan=2, pady=(10, 0))
        buttons = ttk.Frame(body)
        buttons.grid(row=3, column=0, columnspan=2, pady=(14, 0), sticky="e")

        def apply():
            changed = (font_var.get(), page_var.get()) != (old_font, old_page)
            save_setting("epub_font", font_var.get())
            save_setting("epub_page", page_var.get())
            dialog.destroy()
            if changed and getattr(self, "epub_source", None):
                self._reflow_current_epub()

        ttk.Button(buttons, text="Tamam", command=apply).pack(side="left", padx=4)
        ttk.Button(buttons, text="İptal", command=dialog.destroy).pack(side="left")
        dialog.grab_set()

    def _reflow_current_epub(self):
        source = self.epub_source
        if not source or not os.path.exists(source):
            return
        if not self._confirm_discard():
            return
        try:
            index = self.get_current_page_index()
            count = len(self.doc)
            data = self.doc.tobytes()  # dosya kilitli kalmasın: belgeyi belleğe al, sonra yeniden üret
            self.doc.close()
            self.doc = fitz.open("pdf", data)
            self.is_edited = False
            new_path = self._resolve_with_progress(source)
            with fitz.open(new_path) as probe:
                new_count = len(probe)
            self.open_pdf_path(source, start_page=round(index / max(count, 1) * new_count))
        except Exception as exc:  # noqa: BLE001
            messagebox.showerror("E-Kitap Hatası", f"E-kitap yeniden düzenlenemedi:\n{exc}", parent=self)

    def action_document_info(self):
        name = "(açılış sayfası)" if self.is_welcome else os.path.basename(self.pdf_path)
        info = [f"Dosya: {name}", f"Sayfa: {len(self.doc)}"]
        if getattr(self, "epub_source", None):
            info.append(f"Kaynak kitap: {os.path.basename(self.epub_source)}")
        if not self.is_welcome and os.path.exists(self.pdf_path):
            info.append(f"Boyut: {os.path.getsize(self.pdf_path):,} bayt")
        info.append(f"PDF sürümü: {self.doc.metadata.get('format', '')}")
        info.append(f"Şifreli: {'evet' if self.doc.is_encrypted else 'hayır'}")
        info.extend(f"{key}: {value}" for key, value in self.doc.metadata.items() if value and key != "format")
        messagebox.showinfo("Belge bilgileri", "\n".join(info), parent=self)

    def action_shortcuts(self):
        messagebox.showinfo("Klavye kısayolları", "\n".join((
            "Ctrl+O  PDF aç              Ctrl+S  Kaydet",
            "Ctrl+Shift+S  Farklı kaydet   Ctrl+P  Yazdır",
            "Ctrl+Z / Ctrl+Y  Geri al / Yinele",
            "Ctrl+F  Ara    F3 / Shift+F3  Sonraki / önceki eşleşme",
            "Ctrl++ / Ctrl+-  Yakınlaştır / uzaklaştır   Ctrl+0  Genişliğe sığdır",
            "Ctrl+Fare tekerleği  Yakınlaştırma",
            "PageUp / PageDown  Sayfa değiştir",
            "Alt+← / Alt+→  Bağlantıdan önceki / sonraki konuma dön    F4  İçindekiler",
            "F11  Tam ekran    Esc  Tam ekrandan çık",
        )), parent=self)

    def action_set_default_app(self):
        """Uygulamayı kullanıcı için PDF işleyicisi olarak kaydeder ve Windows'un Varsayılan uygulamalar ekranını açar.
        Windows 10/11 varsayılanı programdan sessizce değiştirmeye izin vermez; son onay kullanıcıdadır."""
        if not getattr(sys, "frozen", False):
            messagebox.showinfo("Varsayılan program", "Bu özellik yalnızca kurulu (EXE) sürümde çalışır.", parent=self)
            return
        if not messagebox.askyesno(
                "Varsayılan program",
                f"{APP_NAME}, bu kullanıcı için PDF dosyalarını açabilen program olarak kaydedilecek ve Windows'un "
                "«Varsayılan uygulamalar» ekranı açılacak.\n\nBu ekranda .pdf için Baba PDF'i seçmeniz gerekir "
                "(Windows bunu sessizce yapmaya izin vermez). Devam edilsin mi?", parent=self):
            return
        try:
            register_pdf_handler(os.path.abspath(sys.executable))
            os.startfile("ms-settings:defaultapps")
        except Exception as exc:  # noqa: BLE001
            messagebox.showerror("Varsayılan program", f"Kayıt yapılamadı:\n{exc}", parent=self)

    def action_check_update(self):
        url = update_manifest_url()
        if not url:
            messagebox.showinfo(
                "Güncelleme", "Güncelleme adresi tanımlı değil.\n\nBelirlemek için BABA_PDF_UPDATE_URL ortam değişkenini "
                "veya APPDATA\\BabaPDF\\guncelleme_adresi.txt dosyasını (tek satır, https:// adresi) kullanın.\n"
                "Denetim yalnızca bu menüden, sizin isteğinizle yapılır; adrese hiçbir belge ya da kişisel veri gönderilmez.",
                parent=self)
            return
        try:
            newer, manifest = check_for_update(url)
        except Exception as exc:  # noqa: BLE001
            messagebox.showerror("Güncelleme", f"Güncelleme denetlenemedi:\n{exc}", parent=self)
            return
        if not newer:
            messagebox.showinfo("Güncelleme", f"Güncel sürümü kullanıyorsunuz ({APP_VERSION}).", parent=self)
            return
        link = str(manifest.get("url", ""))
        text = f"Yeni sürüm var: {manifest.get('version')}\n{manifest.get('notes', '')}"
        if not link.lower().startswith("https://"):
            messagebox.showinfo("Güncelleme", text, parent=self)
        elif messagebox.askyesno("Güncelleme", text + "\n\nİndirme sayfası açılsın mı?", parent=self):
            import webbrowser
            webbrowser.open(link)

    def action_license(self):
        t = self.theme
        win = tk.Toplevel(self)
        win.title("Lisans ve Kaynak Kodu")
        win.transient(self)
        win.geometry("760x560")
        notice = ""
        for name in ("NOTICE.txt", "LICENSE.txt"):
            try:
                with open(resource_path(name), encoding="utf-8") as handle:
                    notice += handle.read() + "\n\n" + "=" * 70 + "\n\n"
            except OSError:
                notice += f"({name} bulunamadı)\n\n"
        source_dir = os.path.join(os.path.dirname(os.path.abspath(sys.executable)), "kaynak")
        where = source_dir if os.path.isdir(source_dir) else "(kurulum paketindeki «kaynak» klasörü)"
        header = f"Bu program GNU Affero Genel Kamu Lisansı v3 (AGPL-3.0) altında sunulur. Tam kaynak kodu:\n{where}"
        tk.Label(win, text=header, bg=t["bg"], fg=t["text"], font=(FONT_UI, 9), justify="left", anchor="w").pack(fill="x", padx=12, pady=(10, 4))
        frame = tk.Frame(win, bg=t["bg"])
        frame.pack(fill="both", expand=True, padx=12, pady=4)
        box = tk.Text(frame, wrap="word", font=("Consolas", 9), bg=t["entry"], fg=t["text"], relief="flat")
        bar = ttk.Scrollbar(frame, command=box.yview)
        box.configure(yscrollcommand=bar.set)
        bar.pack(side="right", fill="y")
        box.pack(side="left", fill="both", expand=True)
        box.insert("1.0", notice)
        box.configure(state="disabled")
        buttons = tk.Frame(win, bg=t["bg"])
        buttons.pack(fill="x", padx=12, pady=8)
        if os.path.isdir(source_dir):
            ttk.Button(buttons, text="Kaynak klasörünü aç", command=lambda: reveal_in_folder(source_dir)).pack(side="left")
        ttk.Button(buttons, text="Kapat", style="Save.TButton", command=win.destroy).pack(side="right")

    def action_about(self):
        t = self.theme
        win = tk.Toplevel(self)
        win.title(f"{APP_NAME} Hakkında")
        win.transient(self)
        win.resizable(False, False)
        card = tk.Frame(win, bg=t["lavender"], padx=28, pady=20)
        card.pack(padx=16, pady=16)
        tk.Label(card, text="📄", bg=t["lavender"], font=("Segoe UI Emoji", 34)).pack()
        tk.Label(card, text=APP_NAME, bg=t["lavender"], fg=t["text"], font=(FONT_UI, 20, "bold")).pack()
        version_text = f"Sürüm {APP_VERSION} · {APP_DATE}"
        if LAST_STARTUP_SECONDS:
            version_text += f"\nSon açılış süresi: {LAST_STARTUP_SECONDS:.1f} sn"
        tk.Label(card, text=version_text, bg=t["lavender"], fg=t["muted"],
                 font=(FONT_UI, 10)).pack(pady=(0, 10))
        libs = (f"PyMuPDF {getattr(fitz, 'VersionBind', '?')}  (AGPL-3.0)\n"
                "Tesseract OCR  (Apache-2.0)\n"
                f"{APP_NAME} AGPL-3.0 ile sunulur; kaynak kodu kurulumla gelir (Yardım ▸ Lisans).\n"
                "Belgeleriniz bilgisayarınızdan çıkmaz; hiçbir veri gönderilmez.")
        tk.Label(card, text=libs, bg=t["lavender"], fg=t["text"], font=(FONT_UI, 9), justify="center").pack()
        ttk.Button(card, text="Tamam", style="Save.TButton", command=win.destroy).pack(pady=(14, 0))
        win.grab_set()

    def action_fill_form(self):
        fields = []
        for page_index, page in enumerate(self.doc):
            widget = page.first_widget
            while widget:
                fields.append((page_index, widget))
                widget = widget.next
        if not fields:
            messagebox.showinfo("Form", "Bu PDF içinde doldurulabilir form alanı bulunamadı.", parent=self)
            return
        for page_index, widget in fields:
            current = widget.field_value or ""
            value = simpledialog.askstring("Form alanı", f"{widget.field_name} (sayfa {page_index + 1}):", initialvalue=str(current), parent=self)
            if value is not None:
                self._record_edit()
                widget.field_value = value
                widget.update()
        self.render_all_pages()

    def _clear_widget_hits(self):
        if self.widget_hits:
            self.widget_hits = {}
            try:
                self.render_all_pages()
            except tk.TclError:
                pass

    def action_form_editor(self):
        fields = []
        for page_index, page in enumerate(self.doc):
            widget = page.first_widget
            while widget:
                fields.append((page_index, widget))
                widget = widget.next
        if not fields:
            messagebox.showinfo("Form Editörü", "Bu PDF içinde form alanı bulunamadı.", parent=self)
            return
        win = tk.Toplevel(self)
        win.title("Form Editörü")
        win.geometry("760x420")
        win.transient(self)
        columns = ("page", "name", "type", "value")
        tree = ttk.Treeview(win, columns=columns, show="headings", selectmode="browse")
        for col, title, width in (("page", "Sayfa", 70), ("name", "Alan", 220), ("type", "Tür", 100), ("value", "Değer", 320)):
            tree.heading(col, text=title)
            tree.column(col, width=width)
        tree.pack(fill="both", expand=True, padx=8, pady=8)
        for index, (page_index, widget) in enumerate(fields):
            tree.insert("", "end", iid=str(index), values=(page_index + 1, widget.field_name or "", widget.field_type_string, widget.field_value or ""))
        def edit_selected():
            selected = tree.selection()
            if not selected:
                return
            index = int(selected[0])
            page_index, widget = fields[index]
            value = simpledialog.askstring("Form alanı", widget.field_name or "Değer", initialvalue=str(widget.field_value or ""), parent=win)
            if value is not None:
                self._record_edit()
                widget.field_value = value
                widget.update()
                tree.item(selected[0], values=(page_index + 1, widget.field_name or "", widget.field_type_string, value))
        def focus_field(_event=None):
            selected = tree.selection()
            if not selected:
                return
            page_index, widget = fields[int(selected[0])]
            self.widget_hits = {page_index: [widget.rect]}
            self.page_var.set(str(page_index + 1))
            self.render_all_pages()
            self.goto_page()
        tree.bind("<<TreeviewSelect>>", focus_field)
        tree.bind("<Double-1>", lambda _e: edit_selected())
        win.bind("<Destroy>", lambda e: self._clear_widget_hits() if e.widget is win else None)
        ttk.Button(win, text="Seçili alanı düzenle", command=edit_selected).pack(side="left", padx=8, pady=(0, 8))
        ttk.Button(win, text="Kapat", command=win.destroy).pack(side="right", padx=8, pady=(0, 8))

    def action_add_signature(self):
        path = filedialog.askopenfilename(title="İmza görseli seçin", filetypes=[("Görseller", "*.png *.jpg *.jpeg")])
        if not path:
            return
        win = tk.Toplevel(self)
        win.title("İmza yerleştir")
        win.transient(self)
        page_var = tk.IntVar(value=self.get_current_page_index() + 1)
        x_var = tk.DoubleVar(value=50)
        y_var = tk.DoubleVar(value=50)
        width_var = tk.DoubleVar(value=180)
        height_var = tk.DoubleVar(value=80)
        try:
            from PIL import Image, ImageTk
            preview_image = Image.open(path).convert("RGBA")
            preview_image.thumbnail((240, 120))
            preview_photo = ImageTk.PhotoImage(preview_image)
            ttk.Label(win, image=preview_photo, text="").grid(row=0, column=2, rowspan=5, padx=12)
            win._preview_photo = preview_photo
        except Exception:
            ttk.Label(win, text="Önizleme yüklenemedi").grid(row=0, column=2, rowspan=5, padx=12)
        ttk.Label(win, text="PDF koordinatları sol üstten başlar.", foreground="#6b7280").grid(row=6, column=0, columnspan=3, padx=10, pady=4)
        def place():
            try:
                page = self.doc[page_var.get() - 1]
                rect = fitz.Rect(x_var.get(), y_var.get(), x_var.get() + width_var.get(), y_var.get() + height_var.get())
                rect = rect & page.rect
                if rect.is_empty:
                    raise ValueError("İmza alanı sayfa dışında.")
                self._record_edit()
                page.insert_image(rect, filename=path, keep_proportion=True, overlay=True)
                win.destroy()
                self.render_all_pages()
            except (ValueError, IndexError) as exc:
                messagebox.showerror("İmza yerleştirme", str(exc), parent=win)
        ttk.Button(win, text="Yerleştir", command=place).grid(row=5, column=0, padx=10, pady=10)
        ttk.Button(win, text="İptal", command=win.destroy).grid(row=5, column=1, padx=10, pady=10)

    def action_open_error_log(self):
        path = error_log_path()
        if not os.path.exists(path):
            with open(path, "w", encoding="utf-8") as log:
                log.write("Baba PDF hata günlüğü\n")
        try:
            if sys.platform == "win32":
                os.startfile(path)
            elif sys.platform == "darwin":
                subprocess.run(["open", path], check=False)
            else:
                subprocess.run(["xdg-open", path], check=False)
        except Exception as exc:
            show_runtime_error(self, "Hata günlüğü açma", exc)

    def action_batch_extract(self):
        files = filedialog.askopenfilenames(title="PDF dosyalarını seçin", filetypes=[("PDF", "*.pdf")])
        folder = filedialog.askdirectory(title="TXT çıktı klasörü") if files else None
        if not files or not folder:
            return
        failed = []
        for path in files:
            out = os.path.join(folder, os.path.splitext(os.path.basename(path))[0] + ".txt")
            try:
                PDFProcessor(path).extract_text(out)
            except Exception as e:
                failed.append(f"{os.path.basename(path)}: {e}")
        done = len(files) - len(failed)
        if failed:
            messagebox.showwarning("Tamamlandı (hatalarla)", f"{done} PDF işlendi, {len(failed)} başarısız:\n" + "\n".join(failed[:10]))
        else:
            self._toast(f"{done} PDF işlendi")

    def action_ocr(self):
        try:
            pytesseract, Image = load_ocr()
            lang = pick_ocr_langs(pytesseract, ["tur", "eng"])
        except RuntimeError as exc:
            messagebox.showerror("OCR Kullanılamıyor", str(exc))
            return
        path = filedialog.asksaveasfilename(defaultextension=".txt", filetypes=[("Metin", "*.txt")])
        if not path:
            return
        try:
            source_path, is_tmp = self._working_pdf_path()
        except Exception as exc:
            messagebox.showerror("OCR Hatası", str(exc), parent=self)
            return

        def task(progress):
            try:
                with fitz.open(source_path) as doc, open(path, "w", encoding="utf-8") as out:
                    total = len(doc)
                    for i, page in enumerate(doc):
                        pix = page.get_pixmap(dpi=200)
                        image = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                        out.write(f"--- Sayfa {i + 1} ---\n{pytesseract.image_to_string(image, lang=lang)}\n\n")
                        progress(i + 1, total, "OCR sayfa")
            finally:
                if is_tmp:
                    self._remove_quietly(source_path)

        def done(_result, error):
            if error is not None:
                messagebox.showerror("OCR Hatası", str(error), parent=self)
            else:
                self._toast("OCR metni kaydedildi", path)

        if not self._run_background(task, done, "OCR hazırlanıyor"):
            if is_tmp:
                self._remove_quietly(source_path)

    def action_print(self):
        source, is_tmp = None, False
        try:
            source, is_tmp = self._working_pdf_path()
            if sys.platform == "win32":
                os.startfile(source, "print")
            elif sys.platform == "darwin":
                subprocess.run(["open", source], check=False)
            else:
                subprocess.run(["lp", source], check=False)
            if is_tmp:  # yazdırma programı dosyayı okuyana dek bekle, sonra düzenlenmiş kopyayı diskten sil
                self.after(180000, lambda p=source: self._remove_quietly(p))
            self._toast("Yazdırma işlemi başlatıldı")
        except Exception as e:
            if is_tmp and source:
                self._remove_quietly(source)
            messagebox.showerror(
                "Yazdırma Hatası",
                f"{e}\n\nWindows'ta bu işlem için varsayılan bir PDF programı gerekir. "
                "Olmazsa PDF'yi kaydedip başka bir programda yazdırabilirsiniz.", parent=self)

    def on_wheel(self, e):
        self.canvas.yview_scroll(int(-e.delta / 120) * 3, "units")


if __name__ == "__main__":
    try:  # Kurulum/kaldırma programı (Inno AppMutex) çalışan uygulamayı bu adla tanır
        _app_mutex = ctypes.windll.kernel32.CreateMutexW(None, False, "BabaPDFCalisiyorMutex")
    except Exception:  # noqa: BLE001
        _app_mutex = None
    _arg = ""
    if len(sys.argv) > 1 and is_openable_document(sys.argv[1]):
        _arg = os.path.normcase(os.path.abspath(sys.argv[1]))
    if not claim_startup(_arg):
        sys.exit(0)  # başka bir kopya şu an açılıyor: ikinci pencere açma
    pdf_path = None
    welcome_start = False
    sweep_old_temp_files()
    if len(sys.argv) > 1 and is_openable_document(sys.argv[1]) and os.path.exists(sys.argv[1]):
        try:
            pdf_path = resolve_document_path(sys.argv[1])
        except Exception as exc:  # noqa: BLE001
            log_exception("Kitap dönüştürme", exc)
            _r = tk.Tk()
            _r.withdraw()
            messagebox.showerror("Kitap açılamadı", str(exc))
            _r.destroy()
            sys.exit(1)
    else:
        # EXE'ye çift tıklanınca dosya seçme penceresi açılmaz; boş çalışma alanı açılır.
        # _v2: eski sürümün ğ/ı eksik (Helvetica) önbellek dosyası yeniden kullanılmasın.
        welcome_start = True
        pdf_path = os.path.join(tempfile.gettempdir(), "baba_pdf_empty_workspace_v2.pdf")
        if not os.path.exists(pdf_path):
            starter = fitz.open()
            page = starter.new_page()
            insert_text_tr(page, (72, 100), "PDF Düzenleme Merkezi", fontsize=20, color=(0.2, 0.2, 0.2))
            insert_text_tr(page, (72, 135), "Dosya menüsünden Yeni PDF Aç seçeneğini kullanın.", fontsize=11, color=(0.35, 0.35, 0.35))
            starter.save(pdf_path)
            starter.close()

    try:
        app = PDFViewer(pdf_path, welcome=welcome_start)
        app.is_welcome = welcome_start
        app._refresh_menu_state()
        if not welcome_start:
            app._restore_position()
        recovery_file = None if welcome_start else find_recovery_for(pdf_path)
        if recovery_file and messagebox.askyesno(
                "Kurtarma bulundu",
                "Bu PDF için daha yeni bir kurtarma kaydı bulundu. Geri yüklensin mi?",
                parent=app):
            app.restore_recovery(recovery_file)
        elif recovery_file:
            app._remove_quietly(recovery_file)
    except Exception as exc:
        log_exception("Başlangıç hatası", exc)
        root = tk.Tk()
        root.withdraw()
        messagebox.showerror("PDF açılamadı", f"{exc}\n\nAyrıntılar:\n{error_log_path()}")
        root.destroy()
        sys.exit(1)
    def _startup_done(_event=None):
        app.unbind("<Map>")
        mark_startup_ready()

    app.bind("<Map>", _startup_done)
    try:
        app.mainloop()
    except Exception as exc:
        log_exception("Ana döngü hatası", exc)
        try:
            messagebox.showerror("Uygulama hatası", f"{exc}\n\nAyrıntılar:\n{error_log_path()}", parent=app)
        except Exception:
            pass
        sys.exit(1)





