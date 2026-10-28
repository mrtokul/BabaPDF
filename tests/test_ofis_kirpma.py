"""Kırpma, sayfa ekleme/silme ve Office dışa aktarma çekirdek işlevleri için testler (arayüz açılmaz)."""
import os
import sys

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

pymupdf = pytest.importorskip("pymupdf")
import pdf_gui as g  # noqa: E402

P = g.PDFProcessor


def _doc(pages=3):
    doc = pymupdf.open()
    for n in range(pages):
        page = doc.new_page()
        page.insert_text((72, 100), f"Baslik {n + 1}", fontsize=20)
        page.insert_text((72, 140), "Normal metin satiri", fontsize=11)
    return doc


def test_alana_kirpma_ve_kaldirma():
    doc = _doc()
    P.crop_to_rect(doc, [0], pymupdf.Rect(50, 50, 250, 250))
    assert round(doc[0].rect.width) == 200 and round(doc[0].rect.height) == 200
    assert doc[1].rect.width > 500  # diğer sayfa etkilenmez
    assert P.reset_crop(doc, [0, 1, 2]) == 1
    assert doc[0].cropbox == doc[0].mediabox


def test_kirpilmis_sayfa_tekrar_kirpilinca_kaymaz():
    doc = _doc(1)
    P.crop_to_rect(doc, [0], pymupdf.Rect(100, 100, 400, 500))
    P.crop_to_rect(doc, [0], pymupdf.Rect(10, 10, 100, 100))   # görünen uzayda
    assert doc[0].cropbox == pymupdf.Rect(110, 110, 200, 200)


def test_donuk_sayfada_kirpma():
    doc = _doc(1)
    doc[0].set_rotation(90)
    P.crop_to_rect(doc, [0], pymupdf.Rect(0, 0, 200, 100))  # ekranda sol üst 200x100
    assert round(doc[0].rect.width) == 200 and round(doc[0].rect.height) == 100


def test_cok_kucuk_alan_reddedilir():
    with pytest.raises(ValueError):
        P.crop_to_rect(_doc(1), [0], pymupdf.Rect(10, 10, 15, 15))


def test_kenar_kirpma_iki_kez_yapilabilir():
    doc = _doc(1)
    P.crop_margins(doc, 20, [0])
    P.crop_margins(doc, 20, [0])
    assert round(doc[0].rect.width) == round(595.28 - 80)


def test_bos_sayfa_ekleme_ve_aralik_silme():
    doc = _doc(3)
    P.insert_blank_pages(doc, 0, 2)
    assert len(doc) == 5 and doc[1].get_text().strip() == "" and "Baslik 2" in doc[3].get_text()
    P.insert_blank_pages(doc, -1, 1)
    assert doc[0].get_text().strip() == "" and len(doc) == 6
    assert P.delete_page_numbers(doc, P.parse_page_range("1-3", len(doc))) == 3
    assert len(doc) == 3
    with pytest.raises(ValueError):
        P.delete_page_numbers(doc, [1, 2, 3])
    with pytest.raises(ValueError):
        P.insert_blank_pages(doc, 0, 0)


@pytest.fixture
def pdf_path(tmp_path):
    doc = _doc(2)
    page = doc[0]
    # tablo benzeri çizgiler + metin
    for x in (72, 172, 272):
        page.draw_line((x, 200), (x, 260))
    for y in (200, 230, 260):
        page.draw_line((72, y), (272, y))
    for r, row in enumerate((("Ad", "Deger"), ("Elma", "5"))):
        for c, cell in enumerate(row):
            page.insert_text((80 + c * 100, 220 + r * 30), cell, fontsize=10)
    path = tmp_path / "k.pdf"
    doc.save(path)
    return str(path)


def test_word_aktarimi(pdf_path, tmp_path):
    from docx import Document
    out = str(tmp_path / "k.docx")
    P.export_docx(pdf_path, out, editable=True)
    text = "\n".join(p.text for p in Document(out).paragraphs)
    assert "Baslik 1" in text and "Baslik 2" in text
    out2 = str(tmp_path / "k2.docx")
    P.export_docx(pdf_path, out2, editable=False)
    assert len(Document(out2).inline_shapes) == 2  # her sayfa bir resim


def test_excel_aktarimi_tablo_ve_metin(pdf_path, tmp_path):
    from openpyxl import load_workbook
    out = str(tmp_path / "k.xlsx")
    tables = P.export_xlsx(pdf_path, out)
    book = load_workbook(out)
    if tables:
        assert any(name.startswith("S1-T") for name in book.sheetnames)
    else:
        assert book.sheetnames == ["Metin"]
    plain = pymupdf.open()
    plain.new_page().insert_text((72, 100), "sadece metin")
    plain_path = str(tmp_path / "p.pdf")
    plain.save(plain_path)
    out2 = str(tmp_path / "p.xlsx")
    assert P.export_xlsx(plain_path, out2) == 0
    sheet = load_workbook(out2)["Metin"]
    assert [c.value for c in sheet[2]] == [1, "sadece metin"]


def test_powerpoint_aktarimi(pdf_path, tmp_path):
    from pptx import Presentation
    for editable in (True, False):
        out = str(tmp_path / f"k{editable}.pptx")
        P.export_pptx(pdf_path, out, editable=editable)
        prs = Presentation(out)
        assert len(prs.slides) == 2
        assert "Baslik 1" in prs.slides[0].notes_slide.notes_text_frame.text
    editable_prs = Presentation(str(tmp_path / "kTrue.pptx"))
    assert any(sh.has_text_frame and "Baslik 1" in sh.text_frame.text for sh in editable_prs.slides[0].shapes)


def test_gecersiz_kontrol_karakterleri_temizlenir():
    assert P._clean_text("a\x00b\x0cc\ud800d") == "abcd"
    assert P._clean_text(None) == ""


@pytest.fixture
def rapor_pdf(tmp_path):
    """Başlık, madde imli liste, çizgili tablo (sayılı, sağa yaslı) ve iki sütunlu metin içeren sayfa."""
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 80), "Rapor", fontsize=22, fontname="hebo", color=(0.8, 0, 0))
    page.insert_text((72, 120), "•", fontsize=11)
    page.insert_text((90, 120), "Birinci madde", fontsize=11)
    page.insert_text((72, 140), "•", fontsize=11)
    page.insert_text((90, 140), "Ikinci madde", fontsize=11)
    for x in (72, 172, 272):
        page.draw_line((x, 200), (x, 260))
    for y in (200, 230, 260):
        page.draw_line((72, y), (272, y))
    for r, (a, b) in enumerate((("Ad", "Deger"), ("Elma", "1.250,50"))):
        page.insert_text((80, 220 + r * 30), a, fontsize=10)
        page.insert_text((266 - pymupdf.get_text_length(b, "helv", 10), 220 + r * 30), b, fontsize=10)
    metin = "Lorem ipsum dolor sit amet consectetur " * 12
    page.insert_textbox(pymupdf.Rect(72, 320, 290, 500), metin, fontsize=10)
    page.insert_textbox(pymupdf.Rect(310, 320, 528, 500), metin, fontsize=10)
    path = tmp_path / "r.pdf"
    doc.save(path)
    return str(path)


def test_word_duzeni_korunur(rapor_pdf, tmp_path):
    from docx import Document
    out = str(tmp_path / "r.docx")
    P.export_docx(rapor_pdf, out)
    d = Document(out)
    assert d.sections[0].page_width.pt == pytest.approx(595, abs=1)  # PDF sayfa boyutu (A4)
    assert len(d.tables) >= 2  # gerçek tablo + iki sütunlu düzen tablosu
    cells = [c.text for t in d.tables for r in t.rows for c in r.cells]
    assert "Elma" in cells and "1.250,50" in cells
    texts = [p.text for p in d.paragraphs]
    assert any(t[:1] in ("•", "·") and "Birinci madde" in t for t in texts)  # madde imi aynı satırda
    title = next(r for p in d.paragraphs for r in p.runs if r.text == "Rapor")
    assert title.bold and title.font.size.pt == 22 and title.font.color.rgb is not None
    assert not any("" in t for t in texts)


def test_excel_sayilar_sayi_olur(rapor_pdf, tmp_path):
    from openpyxl import load_workbook
    out = str(tmp_path / "r.xlsx")
    assert P.export_xlsx(rapor_pdf, out) >= 1
    sheet = load_workbook(out).worksheets[0]
    assert sheet["A2"].value == "Elma"
    assert sheet["B2"].value == pytest.approx(1250.5)
    assert sheet["B2"].alignment.horizontal == "right"


def test_powerpoint_tablo_hucreleri_ayri_kutu(rapor_pdf, tmp_path):
    from pptx import Presentation
    out = str(tmp_path / "r.pptx")
    P.export_pptx(rapor_pdf, out)
    boxes = [s.text_frame.text for s in Presentation(out).slides[0].shapes if s.has_text_frame]
    assert "Elma" in boxes and "1.250,50" in boxes


def test_word_karmasik_sayfa_sabit_konumlu(tmp_path):
    """3 sütunlu, çok resimli broşür sayfası: tek sayfa kalır, yazılar sayfada sabit konumlu çerçeve olur."""
    import zipfile
    doc = pymupdf.open()
    page = doc.new_page()
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 40, 40), False)
    pix.set_rect(pix.irect, (200, 120, 40))
    for i in range(3):
        x = 40 + i * 180
        page.insert_textbox(pymupdf.Rect(x, 60, x + 160, 400), "Lorem ipsum dolor sit amet consectetur " * 9, fontsize=9)
        page.insert_image(pymupdf.Rect(x, 420, x + 60, 480), pixmap=pix)
    page.insert_image(pymupdf.Rect(40, 500, 100, 560), pixmap=pix)
    src = str(tmp_path / "b.pdf")
    doc.save(src)
    out = str(tmp_path / "b.docx")
    P.export_docx(src, out)
    xml = zipfile.ZipFile(out).read("word/document.xml").decode()
    assert "w:framePr" in xml and "behindDoc" in xml
    with pymupdf.open(src) as d:
        assert len(d) == 1
